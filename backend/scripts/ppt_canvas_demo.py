"""新方案原型：模板只当"底板"（页眉、红条、校徽），内容由程序按固定版式排进安全区。

用法（在 backend 目录）：python scripts/ppt_canvas_demo.py <模板.pptx> <输出.pptx>
默认用模板 53（f52834f22b5a），四种底板页：封面 1、转场 3、内容 29、结尾 31。
这里的文字是手写的"计划"示例；正式版由模型按 plan schema 产出。
"""

from __future__ import annotations

import math
import pathlib
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1]))
from services.pptx_splice import _clone_slide, _delete_slides  # noqa: E402

ACCENT = RGBColor(0x9B, 0x1B, 0x30)
ACCENT_SOFT = RGBColor(0xF6, 0xEC, 0xEE)
DARK = RGBColor(0x26, 0x26, 0x26)
GRAY = RGBColor(0x59, 0x59, 0x59)
MUTED = RGBColor(0x8C, 0x8C, 0x8C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

SAFE = (0.74, 1.45, 12.6, 6.75)  # 内容页安全区：页眉线以下、页脚以上


# ───────────── 量字与自适应字号 ─────────────

def _em(text: str) -> float:
    return sum(1.0 if ord(c) > 0x2E80 else 0.55 for c in text)


def fit_pt(text: str, w_in: float, h_in: float, max_pt: int, min_pt: int, spacing: float = 1.3) -> int | None:
    """在 [min_pt, max_pt] 里找能放下的最大字号；放不下返回 None（交给"改短/拆页"）。"""
    for pt in range(max_pt, min_pt - 1, -1):
        per_line = max(1, math.floor(w_in * 72 * 0.94 / pt))
        lines = sum(max(1, math.ceil(_em(part) / per_line)) for part in text.split("\n"))
        if lines * pt * spacing <= h_in * 72:
            return pt
    return None


def text_box(slide, x, y, w, h, text, *, max_pt=18, min_pt=12, color=DARK, bold=False,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.3):
    pt = fit_pt(text, w, h, max_pt, min_pt, spacing)
    if pt is None:
        raise ValueError(f"放不下：{text[:20]}…（{w:.1f}x{h:.1f}in）")
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, Inches(0.03))
    for i, line in enumerate(text.split("\n")):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        para.line_spacing = spacing
        run = para.add_run()
        run.text = line
        run.font.size = Pt(pt)
        run.font.bold = bold
        run.font.color.rgb = color
    return box


def rect(slide, x, y, w, h, fill, *, shape=MSO_SHAPE.RECTANGLE, line=None):
    item = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    item.fill.solid()
    item.fill.fore_color.rgb = fill
    if line is None:
        item.line.fill.background()
    else:
        item.line.color.rgb = line
        item.line.width = Pt(1)
    item.shadow.inherit = False
    if item.has_text_frame:
        item.text_frame.text = ""
    return item


def badge(slide, x, y, d, label, *, fill=ACCENT, pt=16):
    circle = rect(slide, x, y, d, d, fill, shape=MSO_SHAPE.OVAL)
    tf = circle.text_frame
    tf.margin_left = tf.margin_right = Inches(0)
    para = tf.paragraphs[0]
    para.alignment = PP_ALIGN.CENTER
    run = para.add_run()
    run.text = label
    run.font.size = Pt(pt)
    run.font.bold = True
    run.font.color.rgb = WHITE
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE


# ───────────── 底板处理 ─────────────

def _inch(v) -> float:
    return Emu(v or 0).inches


def _texts(slide):
    return [s for s in slide.shapes if s.has_text_frame]


def _set_text(shape, text: str) -> None:
    """换字但保留第一段第一个 run 的样式。"""
    tf = shape.text_frame
    first = tf.paragraphs[0]
    runs = first.runs
    if runs:
        runs[0].text = text
        for extra in runs[1:]:
            extra._r.getparent().remove(extra._r)
    else:
        first.text = text
    for extra in tf.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)


def _drop(shape) -> None:
    shape._element.getparent().remove(shape._element)


def _widen(shape, right_limit: float) -> None:
    shape.width = Inches(max(_inch(shape.width), right_limit - _inch(shape.left)))


def _left(shape) -> None:
    for para in shape.text_frame.paragraphs:
        para.alignment = PP_ALIGN.LEFT


def _overlaps(a, b) -> bool:
    ax, ay, aw, ah = (_inch(v) for v in (a.left, a.top, a.width, a.height))
    bx, by, bw, bh = (_inch(v) for v in (b.left, b.top, b.width, b.height))
    return ax < bx + bw and bx < ax + aw and ay < by + bh and by < ay + ah


def _drop_lines_over(slide, boxes) -> None:
    """模板的装饰细线压到文字上就删掉（碰撞检查的最简版）。"""
    for shape in list(slide.shapes):
        if shape.shape_type == 9 and any(_overlaps(shape, box) for box in boxes):
            _drop(shape)


def _notes(slide, text: str) -> None:
    slide.notes_slide.notes_text_frame.text = text


def content_canvas(prs, src, title: str, page_no: int):
    """复制内容页底板：清掉安全区里的示例内容，只留页眉/校徽/页脚。"""
    slide = _clone_slide(prs, src)
    x0, y0, x1, y1 = SAFE
    for shape in list(slide.shapes):
        cy = _inch(shape.top) + _inch(shape.height) / 2
        if y0 - 0.1 < cy < y1 + 0.15 and _inch(shape.height) > 0.6:
            _drop(shape)
    for shape in _texts(slide):
        text = shape.text_frame.text.strip()
        if text == "请输入你的标题":
            _widen(shape, 10.1)
            _set_text(shape, title)
            shape.text_frame.word_wrap = False
        elif "/" in text and text[:4].isdigit():
            _set_text(shape, "")
        elif text.isdigit():
            _set_text(shape, str(page_no))
    return slide


# ───────────── 版式 ─────────────

def layout_toc(slide, items):
    x0, y0, x1, _ = SAFE
    gap = (x1 - x0) / len(items)
    for i, (label, desc) in enumerate(items):
        cx = x0 + gap * i + gap / 2
        badge(slide, cx - 0.5, y0 + 1.0, 1.0, f"{i + 1:02d}", pt=22)
        text_box(slide, cx - gap / 2 + 0.1, y0 + 2.25, gap - 0.2, 0.55, label,
                 max_pt=22, min_pt=16, bold=True, align=PP_ALIGN.CENTER)
        rect(slide, cx - 0.3, y0 + 2.9, 0.6, 0.05, ACCENT)
        text_box(slide, cx - gap / 2 + 0.15, y0 + 3.1, gap - 0.3, 0.8, desc,
                 max_pt=14, min_pt=12, color=MUTED, align=PP_ALIGN.CENTER)


def layout_cards(slide, cards, takeaway):
    x0, y0, x1, _ = SAFE
    gap = 0.35
    w = (x1 - x0 - gap * (len(cards) - 1)) / len(cards)
    for i, (key, desc) in enumerate(cards):
        x = x0 + i * (w + gap)
        rect(slide, x, y0 + 0.2, w, 3.55, ACCENT_SOFT)
        rect(slide, x, y0 + 0.2, w, 0.08, ACCENT)
        badge(slide, x + 0.3, y0 + 0.55, 0.62, f"{i + 1:02d}", pt=14)
        text_box(slide, x + 0.3, y0 + 1.35, w - 0.6, 0.6, key, max_pt=22, min_pt=18, color=ACCENT, bold=True)
        text_box(slide, x + 0.3, y0 + 2.0, w - 0.6, 1.6, desc, max_pt=16, min_pt=13, color=GRAY)
    rect(slide, x0, y0 + 4.15, x1 - x0, 0.85, ACCENT)
    text_box(slide, x0 + 0.35, y0 + 4.15, x1 - x0 - 0.7, 0.85, takeaway,
             max_pt=18, min_pt=14, color=WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE)


def layout_compare(slide, left, right):
    x0, y0, x1, _ = SAFE
    w = (x1 - x0 - 0.9) / 2
    for col, (head, points, fill) in enumerate([(*left, GRAY), (*right, ACCENT)]):
        x = x0 + col * (w + 0.9)
        rect(slide, x, y0 + 0.2, w, 0.7, fill)
        text_box(slide, x + 0.3, y0 + 0.2, w - 0.6, 0.7, head, max_pt=20, min_pt=16,
                 color=WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        rect(slide, x, y0 + 0.9, w, 4.1, WHITE, line=RGBColor(0xD9, 0xD9, 0xD9))
        for j, point in enumerate(points):
            y = y0 + 1.2 + j * 1.25
            rect(slide, x + 0.35, y + 0.14, 0.14, 0.14, fill)
            text_box(slide, x + 0.65, y, w - 1.0, 1.05, point, max_pt=17, min_pt=14)
    mid = x0 + w + 0.45
    badge(slide, mid - 0.32, y0 + 2.6, 0.64, "→", pt=20)


def layout_timeline(slide, nodes):
    x0, y0, x1, _ = SAFE
    line_y = y0 + 2.0
    rect(slide, x0, line_y, x1 - x0, 0.05, ACCENT)
    gap = (x1 - x0) / len(nodes)
    for i, (year, head, desc) in enumerate(nodes):
        cx = x0 + gap * i + gap / 2
        text_box(slide, cx - gap / 2, line_y - 1.0, gap, 0.7, year, max_pt=28, min_pt=22,
                 color=ACCENT, bold=True, align=PP_ALIGN.CENTER)
        rect(slide, cx - 0.14, line_y - 0.115, 0.28, 0.28, ACCENT, shape=MSO_SHAPE.OVAL)
        text_box(slide, cx - gap / 2 + 0.1, line_y + 0.45, gap - 0.2, 0.5, head, max_pt=18, min_pt=15,
                 bold=True, align=PP_ALIGN.CENTER)
        text_box(slide, cx - gap / 2 + 0.15, line_y + 1.05, gap - 0.3, 1.5, desc, max_pt=14, min_pt=12,
                 color=GRAY, align=PP_ALIGN.CENTER)


def layout_quote(slide, quote, source, note):
    x0, y0, x1, _ = SAFE
    rect(slide, x0 + 1.2, y0 + 1.15, 0.08, 1.4, ACCENT)
    text_box(slide, x0 + 1.2, y0 + 1.05, x1 - x0 - 2.4, 1.6, quote, max_pt=30, min_pt=24,
             bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    rect(slide, (x0 + x1) / 2 - 0.4, y0 + 2.9, 0.8, 0.05, ACCENT)
    text_box(slide, x0 + 1.2, y0 + 3.1, x1 - x0 - 2.4, 0.45, source, max_pt=16, min_pt=13,
             color=MUTED, align=PP_ALIGN.CENTER)
    rect(slide, x0 + 1.2, y0 + 3.9, x1 - x0 - 2.4, 1.0, ACCENT_SOFT)
    text_box(slide, x0 + 1.5, y0 + 3.9, x1 - x0 - 3.0, 1.0, note, max_pt=16, min_pt=13,
             color=GRAY, anchor=MSO_ANCHOR.MIDDLE)


def layout_numbered(slide, rows):
    x0, y0, x1, _ = SAFE
    h = 1.35
    for i, (head, desc) in enumerate(rows):
        y = y0 + 0.2 + i * (h + 0.22)
        rect(slide, x0, y, 1.2, h, ACCENT)
        text_box(slide, x0, y, 1.2, h, f"{i + 1:02d}", max_pt=32, min_pt=26, color=WHITE, bold=True,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        rect(slide, x0 + 1.2, y, x1 - x0 - 1.2, h, ACCENT_SOFT)
        text_box(slide, x0 + 1.55, y + 0.15, 3.2, h - 0.3, head, max_pt=20, min_pt=16, color=ACCENT,
                 bold=True, anchor=MSO_ANCHOR.MIDDLE)
        text_box(slide, x0 + 4.9, y + 0.12, x1 - x0 - 5.2, h - 0.24, desc, max_pt=16, min_pt=13,
                 color=GRAY, anchor=MSO_ANCHOR.MIDDLE)


# ───────────── 组装 ─────────────

def build(template: pathlib.Path, out: pathlib.Path) -> None:
    prs = Presentation(str(template))
    originals = list(prs.slides)
    cover_src, section_src, content_src, end_src = originals[0], originals[2], originals[28], originals[30]
    page = 0

    def next_page():
        nonlocal page
        page += 1
        return page

    # 封面
    slide = _clone_slide(prs, cover_src)
    next_page()
    for shape in _texts(slide):
        text = shape.text_frame.text.strip()
        if _inch(shape.top) < 0:
            _drop(shape)
        elif text == "中国人民大学模板":
            _widen(shape, 8.0)
            _set_text(shape, "国家与土地产权关系的变化")
            shape.text_frame.paragraphs[0].runs[0].font.size = Pt(40)
        elif text.startswith("Thesis"):
            _set_text(shape, "中国农村改革的经济制度变迁史回顾")
        elif "答辩学生" in text:
            _set_text(shape, ">>  汇报人：stu01")
        elif "指导老师" in text:
            _set_text(shape, ">>  课程：中国经济史专题")
    _notes(slide, "开场：今天用制度变迁的视角，回顾改革中国家与土地产权关系怎样变化。")

    # 目录
    slide = content_canvas(prs, content_src, "汇报框架", next_page())
    layout_toc(slide, [
        ("理论框架", "国家与产权的悖论"),
        ("改革前产权", "集体化的由来与特征"),
        ("公社内变迁", "收益、费用与政策退却"),
        ("改革与产权", "包产到户与乡镇企业"),
        ("结论评论", "国家如何保护有效产权"),
    ])

    # 转场
    slide = _clone_slide(prs, section_src)
    next_page()
    subs = ["所有权悖论", "诺斯难题与新假设"]
    filled = []
    for shape in _texts(slide):
        text = shape.text_frame.text.strip()
        if text == "研究现状":
            _widen(shape, 7.3)
            _set_text(shape, "理论框架")
            _left(shape)
            filled.append(shape)
        elif text in ("摘要", "引言"):
            _widen(shape, 6.2)
            _set_text(shape, subs.pop(0))
            _left(shape)
            filled.append(shape)
    _drop_lines_over(slide, filled)

    # 卡片：三个要点 + 一句结论
    slide = content_canvas(prs, content_src, "所有权悖论：产权离不开国家，也容易被国家侵蚀", next_page())
    layout_cards(slide, [
        ("国家必要", "产权的界定与执行，离不开国家的强制力和法律保护。"),
        ("国家自利", "国家追求租金最大化，可能抬高索取，甚至直接剥夺产权。"),
        ("所有权残缺", "Demsetz：国家或他人可以限制所有者对资产的部分权利。"),
    ], "悖论：没有国家，产权无从保护；国家一旦介入，又容易造成所有权残缺。")
    _notes(slide, "三张卡片依次讲：必要性、自利性、残缺的概念；最后落到悖论。")

    # 对比
    slide = content_canvas(prs, content_src, "改革前后：谁掌握土地的剩余权", next_page())
    layout_compare(
        slide,
        ("人民公社体制", [
            "国家用指令计划、统购统销控制农村经济",
            "集体承受控制结果，却没有完整的所有权",
            "监督计量困难，激励不足，效率长期低下",
        ]),
        ("家庭承包制", [
            "农民自发包产到户，国家逐步认可",
            "国家与农户订立新合约，农户获得剩余权",
            "激励恢复，农业产出与收入明显增长",
        ]),
    )

    # 时间线
    slide = content_canvas(prs, content_src, "国家控制费用的两次高峰与政策调整", next_page())
    layout_timeline(slide, [
        ("1952", "合作化起步", "国家开始集中农业剩余"),
        ("1958", "人民公社", "集体化取代农民私有权"),
        ("1960", "第一次退却", "承认副业，以队为基础"),
        ("1978", "包产到户", "费用高峰下自发试点"),
        ("1982", "承包合法化", "订立新的产权合约"),
    ])

    # 引用
    slide = content_canvas(prs, content_src, "农民表达不满的方式", next_page())
    layout_quote(
        slide,
        "农民留在体制内，并非出于对集体的忠诚，\n而是因为别无选择。",
        "—— 文中对公社制下农民处境的概括",
        "既没有“退出权”，也无权自由“叫喊”，农民只能用怠工表达不满，集体效率随之下降。",
    )

    # 结论
    slide = content_canvas(prs, content_src, "结论：国家如何保护有效率的产权", next_page())
    layout_numbered(slide, [
        ("产权保护是增长关键", "长期经济增长取决于国家能否保护有效率的产权制度。"),
        ("保护不会自动出现", "国家通常不会主动提供保护，需要来自社会的持续推动。"),
        ("社会与国家的均势", "农户与新兴产权代理人参与博弈，与国家达成互利交易。"),
    ])

    # 结尾
    slide = _clone_slide(prs, end_src)
    next_page()
    for shape in _texts(slide):
        text = shape.text_frame.text.strip()
        if "答辩学生" in text:
            _set_text(shape, ">>  汇报人：stu01")
        elif "指导老师" in text:
            _set_text(shape, ">>  欢迎批评指正")

    _delete_slides(prs, len(originals))
    prs.save(str(out))


if __name__ == "__main__":
    here = pathlib.Path(__file__).resolve().parents[1]
    default_tpl = here / "data/workspaces/c46ff668-f8c3-46d9-ba62-7f2c271816d4/ppt_templates/f52834f22b5a.pptx"
    tpl = pathlib.Path(sys.argv[1]) if len(sys.argv) > 1 else default_tpl
    dst = pathlib.Path(sys.argv[2]) if len(sys.argv) > 2 else pathlib.Path("ppt_canvas_demo.pptx")
    build(tpl, dst)
    print("saved", dst)
