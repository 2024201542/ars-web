"""用现在的出片方式跑叙事评测，写出指标，不改用户文件夹。

用法（在 backend 目录）：
  D:\\Miniconda3\\python.exe scripts\\ppt\\eval_story.py --matrix quick|full --run v0
  D:\\Miniconda3\\python.exe scripts\\ppt\\eval_story.py --compare v0 v1
"""

from __future__ import annotations

import argparse
import asyncio
import json
import math
import shutil
import sys
import time
from pathlib import Path

BACKEND = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(BACKEND))

import importlib.util

_judge_spec = importlib.util.spec_from_file_location("eval_judge", Path(__file__).with_name("eval_judge.py"))
_judge = importlib.util.module_from_spec(_judge_spec)
_judge_spec.loader.exec_module(_judge)
JUDGE_SYSTEM = _judge.JUDGE_SYSTEM
outline_for_judge = _judge.outline_for_judge
parse_judge = _judge.parse_judge

ROOT = Path(r"C:\Users\28473\Desktop\学术写作工具测试")
USER_ID = "c46ff668-f8c3-46d9-ba62-7f2c271816d4"
STRUCTURAL = {"cover", "toc", "section", "end"}
BRIDGE_WORDS = ("上一页", "前面", "接着", "但是", "但", "然而", "进一步", "例如", "为什么", "回到", "所以", "因此")
RESEARCH_TYPES = {"research", "paper_report"}

CASES = [
    {"id": "M1", "kind": "research", "pages": 12, "files": [ROOT / "记忆方法论文写作" / "物理学习中的记忆机制与教学干预（完整稿）.txt"],
     "guidance": "根据这篇论文做一个研究汇报 PPT，做成 12 页"},
    {"id": "M2", "kind": "lecture", "pages": 12, "files": [ROOT / "学习方法从认知规律到可执行策略.内容稿.md"],
     "guidance": "做一个给学生讲学习方法的课件，做成 12 页"},
    {"id": "M3", "kind": "paper_report", "pages": 12, "files": [],
     "guidance": "组会要讲这篇论文，做成 12 页"},
    {"id": "M4", "kind": "research", "pages": 16, "files": [],
     "guidance": "做一个中国近代女性形象研究的 PPT，做成 16 页",
     "material": "中国近代女性形象研究"},
]
TEMPLATES = [
    {"id": "K1", "template_id": "4d7e82ced452"},
    {"id": "K2", "template_id": "7acd6eb053b2"},
    {"id": "K3", "template_id": "builtin-academic-red"},
]
QUICK = {("M1", "K1"), ("M2", "K1"), ("M3", "K2")}


def read_material(case: dict) -> str:
    if case.get("material"):
        return case["material"]
    files = list(case.get("files") or [])
    if case["id"] == "M3":
        outline = ROOT / "周其仁_中国农村改革_汇报大纲.txt"
        pdfs = list((ROOT / "土地改革论文写作").glob("6 *.pdf"))
        files = [outline] + pdfs[:1]
    parts = []
    for path in files:
        if not path.is_file():
            parts.append(f"（找不到文件 {path.name}）")
            continue
        if path.suffix.lower() == ".pdf":
            import fitz
            doc = fitz.open(path)
            text = "\n".join(page.get_text() for page in doc)
            doc.close()
        else:
            text = path.read_text(encoding="utf-8")
        parts.append(text[:20000])
    return "\n\n".join(parts)


def _content_slides(slides: list[dict]) -> list[tuple[int, dict]]:
    return [(i, s) for i, s in enumerate(slides, 1) if (s or {}).get("layout") not in STRUCTURAL]


SAMPLE_WORDS = ("关键词标题", "此处添加", "在这里输入", "请在此", "请输入", "点击输入", "添加标题", "标题文字")


def count_samples(deck: Path) -> int:
    """成品里还留着模板示例句的文本框个数。组合里的也算。"""
    from pptx import Presentation
    from services.ppt import ppt_compose
    total = 0
    for slide in Presentation(str(deck)).slides:
        for shape in ppt_compose._iter_shapes(slide):
            if not getattr(shape, "has_text_frame", False):
                continue
            text = shape.text_frame.text or ""
            if any(word in text for word in SAMPLE_WORDS):
                total += 1
    return total


def _repeat_content(slides: list) -> int:
    """相邻两张内容页用了同一张模板页的次数。章节页把前后断开。"""
    prev = None
    repeat = 0
    for slide in slides or []:
        layout = (slide or {}).get("layout")
        if layout in STRUCTURAL:
            prev = None
            continue
        number = (slide or {}).get("template_page")
        try:
            current = int(number) if number else None
        except (TypeError, ValueError):
            current = None
        if current and prev == current:
            repeat += 1
        prev = current
    return repeat


def measure(plan: dict, issues: list[dict], template_pages: list[int], requested: int,
            kind: str, judge: dict | None, base_pages: dict | None = None,
            story_issues: list | None = None) -> dict:
    """从计划和自检结果算出一行指标。不看正文。"""
    slides = plan.get("slides") or []
    total = len(slides)
    section_n = sum(1 for s in slides if (s or {}).get("layout") == "section")
    content = _content_slides(slides)
    notes_ok = 0
    bridge_ok = 0
    bridge_need = 0
    seen_section = False
    chapter_first = True
    for index, slide in enumerate(slides, 1):
        layout = (slide or {}).get("layout")
        if layout == "section":
            seen_section = True
            chapter_first = True
            continue
        if layout in STRUCTURAL:
            continue
        notes = str((slide or {}).get("notes") or "")
        if len(notes) >= 80:
            notes_ok += 1
        if seen_section and chapter_first:
            chapter_first = False
            continue
        if content and not seen_section and index == content[0][0]:
            continue
        bridge_need += 1
        slots = (slide or {}).get("slots")
        lead = str((slide or {}).get("lead") or (slots.get("lead") if isinstance(slots, dict) else "") or "")
        head = notes[:40]
        if lead.strip() or any(word in head for word in BRIDGE_WORDS):
            bridge_ok += 1
    kinds = {}
    for item in issues or []:
        kinds[item.get("kind")] = kinds.get(item.get("kind"), 0) + 1
    used = set(int(n) for n in (template_pages or []) if n)
    for layout, key in (("cover", "cover"), ("end", "end"), ("section", "section")):
        if any((s or {}).get("layout") == layout for s in slides):
            number = (base_pages or {}).get(key)
            if number:
                used.add(int(number))
    judge = judge or {}
    scores = judge.get("scores") or {}
    first_content = content[0][0] if content else 0
    window = max(math.ceil(total * 0.4) if total else 1, (first_content + 1) if first_content else 1)
    end_index = next((i for i, s in enumerate(slides, 1) if (s or {}).get("layout") == "end"), total)
    has_roles = any(str((slide or {}).get("role") or "") for slide in slides)
    if has_roles:
        first_role = str((content[0][1] if content else {}).get("role") or "")
        opening = first_role in {"background", "hook", "concept"}
        if kind in RESEARCH_TYPES:
            opening = opening and any(
                str((slide or {}).get("role") or "") == "question" and index <= window
                for index, slide in enumerate(slides, 1))
        prev = slides[end_index - 2] if end_index >= 2 else {}
        closing = str((prev or {}).get("role") or "") in {"conclusion", "summary", "takeaway"} or (
            (prev or {}).get("layout") == "summary")
    else:
        background = set(judge.get("background_pages") or [])
        question = set(judge.get("question_pages") or [])
        summary = set(judge.get("summary_pages") or [])
        opening = first_content in background and (
            kind not in RESEARCH_TYPES or any(page <= window for page in question))
        closing = bool(summary) and (end_index - 1 in summary or end_index in summary)
    story = [item for item in (story_issues or []) if isinstance(item, dict)]
    if any(item.get("code") == "S_SUMMARY_NOT_ANSWER" for item in story):
        closing = False
    return {
        "pages": total,
        "delta": total - requested,
        "section_ratio": round(section_n / total, 2) if total else 0,
        "content_ratio": round(len(content) / total, 2) if total else 0,
        "opening": opening,
        "first_content": first_content,
        "closing": closing,
        "notes_ratio": round(notes_ok / len(content), 2) if content else 0,
        "bridge_ratio": round(bridge_ok / bridge_need, 2) if bridge_need else 0,
        "template_pages": len(used),
        "repeat_content": _repeat_content(slides),
        "requested": requested,
        "offslide": kinds.get("offslide", 0),
        "overlap": kinds.get("overlap", 0),
        "overflow": kinds.get("overflow", 0),
        "junk": kinds.get("junk", 0),
        "story_errors": sum(1 for item in story if item.get("level") == "error"),
        "story_warns": sum(1 for item in story if item.get("level") == "warn"),
        "judge_mean": judge.get("mean") or 0,
        "scores": scores,
        "structure": "".join((_mark((s or {}).get("layout")) for s in slides)),
    }


def _mark(layout: str) -> str:
    return {"cover": "C", "toc": "T", "section": "S", "end": "E"}.get(layout or "", "X")


def contact_sheet(images: list[Path], dest: Path) -> None:
    from PIL import Image
    opened = [Image.open(path).convert("RGB") for path in images if path.is_file()]
    if not opened:
        return
    width = 320
    thumbs = []
    for image in opened:
        height = max(1, int(image.height * width / image.width))
        thumbs.append(image.resize((width, height)))
    cols = 4
    row_h = max(image.height for image in thumbs)
    rows = (len(thumbs) + cols - 1) // cols
    sheet = Image.new("RGB", (cols * width, rows * row_h), "white")
    for index, image in enumerate(thumbs):
        sheet.paste(image, ((index % cols) * width, (index // cols) * row_h))
    dest.parent.mkdir(parents=True, exist_ok=True)
    sheet.save(dest)


async def pick_judge(writer_name: str) -> tuple[str, bool]:
    from providers import MODELS
    from services.ppt.pptx_outline import resolve_model
    writer_id = ""
    for model in MODELS:
        if model["name"] == writer_name or model["id"] == writer_name:
            writer_id = model["id"]
            break
    for model in MODELS:
        if model["id"] == writer_id:
            continue
        entry, key, _base, err = await resolve_model(USER_ID, model["id"])
        if entry and key and not err:
            return model["id"], False
    return writer_id or writer_name, True


async def judge_plan(plan: dict, model_id: str) -> tuple[dict, str]:
    from services.ppt import ppt_plan
    from services.ppt.pptx_outline import _extract_json
    text, name, err = await ppt_plan._ask_text(
        USER_ID, JUDGE_SYSTEM, outline_for_judge(plan), model_id, max_tokens=800)
    if err or not text:
        return {}, err or "评委没有返回"
    data = _extract_json(text) or {}
    parsed = parse_judge(data if isinstance(data, dict) else {})
    parsed["model"] = name
    return parsed, ""


def write_report(rows: list[dict], dest: Path, judge_same: bool) -> None:
    headers = ["案例", "页数", "偏差", "章节占比", "内容占比", "开场", "结尾", "讲稿", "承接",
               "模板页", "连用", "出界", "叠字", "字太长", "示例字", "示例句", "按格子", "审稿错", "审稿警",
               "评委", "秒", "调用", "结构"]
    lines = [
        "# 叙事评测 " + dest.parent.name,
        "",
        "口径：计划页上带叙事角色时，开场看第一张内容页是不是背景、钩子或概念；"
        "研究类还要求前 40% 里有研究问题页。结尾看结束页前一张是不是总结，"
        "总结没有回扣研究问题时记为没通过。没有角色时仍由评委看标题。"
        "评委和写稿" + ("是同一个模型。" if judge_same else "不是同一个模型。"),
        "M4 只有题目，没有先写内容稿。现行对话在材料不足 800 字时会先写内容稿，这里按评测计划直接出计划。",
        "",
        "| " + " | ".join(headers) + " |",
        "| " + " | ".join("---" for _ in headers) + " |",
    ]
    for row in rows:
        m = row.get("metrics") or {}
        lines.append("| " + " | ".join([
            row["case"],
            str(m.get("pages", "")),
            str(m.get("delta", "")),
            str(m.get("section_ratio", "")),
            str(m.get("content_ratio", "")),
            "是" if m.get("opening") else "否",
            "是" if m.get("closing") else "否",
            str(m.get("notes_ratio", "")),
            str(m.get("bridge_ratio", "")),
            str(m.get("template_pages", "")),
            str(m.get("repeat_content", "")),
            str(m.get("offslide", "")),
            str(m.get("overlap", "")),
            str(m.get("overflow", "")),
            str(m.get("junk", "")),
            str(m.get("samples", "")),
            f"{m.get('slot_written', '')}/{m.get('content_pages', '')}" if "slot_written" in m else "",
            str(m.get("story_errors", "")),
            str(m.get("story_warns", "")),
            str(m.get("judge_mean", "")),
            str(row.get("seconds", "")),
            str(row.get("calls", "")),
            str(m.get("structure", "")),
        ]) + " |")
    if rows:
        def avg(key):
            values = [row["metrics"][key] for row in rows if row.get("metrics") and key in row["metrics"]]
            if not values or isinstance(values[0], str):
                return ""
            return round(sum(values) / len(values), 2)
        lines.append("")
        lines.append(
            f"平均：页数偏差 {avg('delta')}，章节占比 {avg('section_ratio')}，内容占比 {avg('content_ratio')}，"
            f"讲稿 {avg('notes_ratio')}，承接 {avg('bridge_ratio')}，评委 {avg('judge_mean')}，"
            f"出界 {avg('offslide')}，叠字 {avg('overlap')}，字太长 {avg('overflow')}，"
            f"示例字 {avg('junk')}，审稿错 {avg('story_errors')}，审稿警 {avg('story_warns')}，"
            f"内容页连用 {avg('repeat_content')}。"
        )
        lines.append(
            f"开场通过 {sum(1 for row in rows if (row.get('metrics') or {}).get('opening'))}/{len(rows)}，"
            f"结尾通过 {sum(1 for row in rows if (row.get('metrics') or {}).get('closing'))}/{len(rows)}。"
        )
    failed = [row for row in rows if row.get("error")]
    if failed:
        lines.append("")
        lines.append("失败：" + "、".join(f"{row['case']}（{row['error'][:40]}）" for row in failed))
    dest.write_text("\n".join(lines) + "\n", encoding="utf-8")


async def run_case(case: dict, template_case: dict, out_dir: Path, judge_model: str, calls: dict) -> dict:
    from services.ppt import ppt_base, ppt_check, ppt_compose, ppt_library, ppt_plan
    name = f"{case['id']}-{template_case['id']}"
    folder = out_dir / name
    folder.mkdir(parents=True, exist_ok=True)
    started = time.time()
    row = {"case": name, "seconds": 0, "calls": 0, "error": ""}
    before = calls["n"]
    try:
        material = read_material(case)
        plan, model_name, err = await ppt_plan.make_plan(
            USER_ID, material, pages=case["pages"], guidance=case["guidance"],
            template_id=template_case["template_id"])
        row["model"] = model_name
        if err or not (plan or {}).get("slides"):
            row["error"] = err or "没有计划"
            return row
        (folder / "plan.json").write_text(json.dumps(plan, ensure_ascii=False, indent=2), encoding="utf-8")
        template = ppt_library.find(USER_ID, template_case["template_id"])
        if not template or not template.get("file_path"):
            row["error"] = "找不到模板"
            return row
        from services.ppt import ppt_pages
        page_doc = ppt_pages.load_pages(template) or ppt_pages.empty_pages(template, template)
        base = ppt_base.base_for(template, page_doc)
        deck = folder / "deck.pptx"
        result = ppt_compose.compose(
            base, plan, template["file_path"], deck, None, template, pages=page_doc)
        issues = ppt_check.check_deck(deck, base)
        judged, judge_err = await judge_plan(plan, judge_model)
        if judge_err:
            row["judge_error"] = judge_err
        row["metrics"] = measure(
            plan, issues, result.get("template_pages") or [], case["pages"], case["kind"],
            judged, (base or {}).get("pages"), plan.get("story_issues") or [])
        row["metrics"]["samples"] = count_samples(deck)
        content_slides = _content_slides(plan.get("slides") or [])
        row["metrics"]["slot_written"] = sum(
            1 for _i, slide in content_slides
            if isinstance(slide.get("slots"), dict) and any(str(v or "").strip() for v in slide["slots"].values()))
        row["metrics"]["content_pages"] = len(content_slides)
        from services.ppt import pptx_render
        pages = list(range(1, len(plan.get("slides") or []) + 1))
        images = pptx_render.render_pages(deck, pages, dpi=72) if pptx_render.available() else {}
        ordered = [images[n] for n in pages if n in images]
        if ordered:
            contact_sheet(ordered, folder / "sheet.png")
        if not ordered:
            row["render_error"] = pptx_render.last_error() if pptx_render.available() else "没有渲染器"
    except Exception as exc:
        row["error"] = f"{type(exc).__name__}: {exc}"
    finally:
        row["seconds"] = round(time.time() - started, 1)
        row["calls"] = calls["n"] - before
    return row


def _load_rows(name: str) -> list:
    path = Path(r"d:\academic-research-skills-main\.tmp\eval") / name / "report.json"
    data = json.loads(path.read_text(encoding="utf-8"))
    if isinstance(data, dict):
        return list(data.get("rows") or [])
    return list(data)


def _mean(rows: list, key: str, on_row: bool = False):
    values = []
    for row in rows:
        source = row if on_row else (row.get("metrics") or {})
        value = source.get(key)
        if isinstance(value, bool) or not isinstance(value, (int, float)):
            continue
        values.append(value)
    if not values:
        return None
    return round(sum(values) / len(values), 2)


def _count_flag(rows: list, key: str) -> str:
    got = sum(1 for row in rows if (row.get("metrics") or {}).get(key))
    return f"{got}/{len(rows)}" if rows else "0/0"


def compare_runs(left: str, right: str) -> int:
    """打印两轮指标和右轮没达标的项。"""
    left_rows = _load_rows(left)
    right_rows = _load_rows(right)
    labels = [
        ("delta", "页数偏差", False),
        ("section_ratio", "章节占比", False),
        ("content_ratio", "内容占比", False),
        ("notes_ratio", "讲稿", False),
        ("bridge_ratio", "承接", False),
        ("template_pages", "模板页", False),
        ("repeat_content", "内容页连用", False),
        ("offslide", "出界", False),
        ("overlap", "叠字", False),
        ("overflow", "字太长", False),
        ("junk", "示例字", False),
        ("samples", "示例句", False),
        ("slot_written", "按格子写的内容页", False),
        ("story_errors", "审稿错", False),
        ("story_warns", "审稿警", False),
        ("judge_mean", "评委", False),
        ("seconds", "秒", True),
        ("calls", "调用", True),
    ]
    lines = [f"# {left} 对比 {right}", ""]
    lines.append(f"| 指标 | {left} | {right} | 变化 |")
    lines.append("| --- | --- | --- | --- |")
    for key, label, on_row in labels:
        a = _mean(left_rows, key, on_row)
        b = _mean(right_rows, key, on_row)
        change = "" if a is None or b is None else round(b - a, 2)
        lines.append(f"| {label} | {a if a is not None else '—'} | {b if b is not None else '—'} | {change if change != '' else '—'} |")
    lines.append(f"| 开场 | {_count_flag(left_rows, 'opening')} | {_count_flag(right_rows, 'opening')} | |")
    lines.append(f"| 结尾 | {_count_flag(left_rows, 'closing')} | {_count_flag(right_rows, 'closing')} | |")
    misses = []
    for row in right_rows:
        m = row.get("metrics") or {}
        name = row.get("case") or ""
        if row.get("error"):
            misses.append(f"{name}：出片失败（{row['error'][:40]}）")
            continue
        if not m:
            misses.append(f"{name}：没有指标")
            continue
        asked = m.get("requested") or (16 if str(name).startswith("M4") else 12)
        need_pages = 10 if int(asked) >= 16 else 8
        checks = [
            (abs(m.get("delta") or 0) > 1, f"页数偏差 {m.get('delta')}，目标 ±1"),
            ((m.get("section_ratio") or 0) > 0.25, f"章节占比 {m.get('section_ratio')}，目标 ≤0.25"),
            ((m.get("content_ratio") or 0) < 0.60, f"内容占比 {m.get('content_ratio')}，目标 ≥0.60"),
            (not m.get("opening"), "开场没通过"),
            (not m.get("closing"), "结尾没回扣"),
            ((m.get("story_errors") or 0) > 0, f"审稿错误 {m.get('story_errors')}，目标 0"),
            ((m.get("notes_ratio") or 0) < 1, f"讲稿 {m.get('notes_ratio')}，目标 1"),
            ((m.get("bridge_ratio") or 0) < 0.90, f"承接 {m.get('bridge_ratio')}，目标 ≥0.90"),
            ((m.get("template_pages") or 0) < need_pages, f"模板页 {m.get('template_pages')}，目标 ≥{need_pages}"),
            ((m.get("repeat_content") or 0) > 0, f"内容页连用 {m.get('repeat_content')}，目标 0"),
            ((m.get("offslide") or 0) > 0, f"出界 {m.get('offslide')}"),
            ((m.get("overlap") or 0) > 0, f"叠字 {m.get('overlap')}"),
            ((m.get("junk") or 0) > 0, f"示例字 {m.get('junk')}"),
            ((m.get("samples") or 0) > 0, f"示例句 {m.get('samples')} 处"),
        ]
        for failed, text in checks:
            if failed:
                misses.append(f"{name}：{text}")
    warn_avg = _mean(right_rows, "story_warns")
    judge_avg = _mean(right_rows, "judge_mean")
    overflow_avg = _mean(right_rows, "overflow")
    if warn_avg is not None and warn_avg > 2:
        misses.append(f"审稿警告平均 {warn_avg}，目标 ≤2")
    if judge_avg is not None and judge_avg < 4:
        misses.append(f"评委平均 {judge_avg}，目标 ≥4")
    if overflow_avg is not None and overflow_avg > 1:
        misses.append(f"字太长平均 {overflow_avg}，目标 ≤1")
    left_seconds = _mean(left_rows, "seconds", True)
    right_seconds = _mean(right_rows, "seconds", True)
    if left_seconds and right_seconds and right_seconds > left_seconds * 2:
        lines.append("")
        lines.append(f"{right} 平均 {right_seconds} 秒，{left} 平均 {left_seconds} 秒，慢了一倍以上。")
    lines.append("")
    if misses:
        lines.append("没达标：")
        lines.extend(f"- {item}" for item in misses)
    else:
        lines.append("指标都在目标里。")
    text = "\n".join(lines) + "\n"
    print(text)
    out = Path(r"d:\academic-research-skills-main\.tmp\eval") / right / "compare.md"
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(text, encoding="utf-8")
    shot = Path(r"d:\academic-research-skills-main\开发笔记\截图\ppt-叙事") / right
    shot.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(out, shot / "compare.md")
    return 0


async def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--matrix", choices=("quick", "full"), default="quick")
    parser.add_argument("--run", default="v0")
    parser.add_argument("--model", default="")
    parser.add_argument("--judge", default="")
    parser.add_argument("--compare", nargs=2, metavar=("LEFT", "RIGHT"))
    parser.add_argument("--only", default="", help="逗号分隔的案例名，例如 M2-K2。补跑时不删已有结果")
    args = parser.parse_args()
    if args.compare:
        raise SystemExit(compare_runs(args.compare[0], args.compare[1]))
    import logging
    logging.disable(logging.CRITICAL)
    from services.ppt import ppt_plan
    calls = {"n": 0}
    ask_json = ppt_plan._ask_json
    ask_text = ppt_plan._ask_text

    async def count_json(*a, **k):
        calls["n"] += 1
        return await ask_json(*a, **k)

    async def count_text(*a, **k):
        calls["n"] += 1
        return await ask_text(*a, **k)

    ppt_plan._ask_json = count_json
    ppt_plan._ask_text = count_text

    out_dir = Path(r"d:\academic-research-skills-main\.tmp\eval") / args.run
    only = {part.strip() for part in (args.only or "").split(",") if part.strip()}
    if out_dir.exists() and not only:
        shutil.rmtree(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    pairs = [(case, tpl) for case in CASES for tpl in TEMPLATES
             if args.matrix == "full" or (case["id"], tpl["id"]) in QUICK]
    if only:
        pairs = [pair for pair in pairs if f"{pair[0]['id']}-{pair[1]['id']}" in only]
    probe, _model, err = await ppt_plan._ask_text(USER_ID, "只回复 OK", "OK", args.model, max_tokens=8)
    if err:
        print("模型不可用：", err)
        sys.stdout.flush()
        raise SystemExit(1)
    judge_model = args.judge
    same = True
    if not judge_model:
        judge_model, same = await pick_judge(_model or args.model)
    previous = []
    if only and (out_dir / "report.json").is_file():
        loaded = json.loads((out_dir / "report.json").read_text(encoding="utf-8"))
        if isinstance(loaded, list):
            previous = [item for item in loaded if isinstance(item, dict)]
    rows = []
    snapshot = rows
    for index, (case, tpl) in enumerate(pairs, 1):
        print(f"[{index}/{len(pairs)}] {case['id']}-{tpl['id']}", flush=True)
        row = await run_case(case, tpl, out_dir, judge_model, calls)
        rows.append(row)
        snapshot = rows
        if previous:
            merged = {item.get("case"): item for item in previous}
            order = [item.get("case") for item in previous]
            for item in rows:
                merged[item.get("case")] = item
            snapshot = [merged[name] for name in order if name in merged]
            seen = set(order)
            snapshot.extend(item for item in merged.values() if item.get("case") not in seen)
        (out_dir / "report.json").write_text(json.dumps(snapshot, ensure_ascii=False, indent=2), encoding="utf-8")
        print(" ", row.get("error") or row.get("metrics", {}).get("structure"), row.get("seconds"), "s", flush=True)
    if previous:
        rows = snapshot
    else:
        (out_dir / "report.json").write_text(json.dumps(rows, ensure_ascii=False, indent=2), encoding="utf-8")
    report = out_dir / "report.md"
    write_report(rows, report, same)
    shot = Path(r"d:\academic-research-skills-main\开发笔记\截图\ppt-叙事") / args.run
    shot.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(report, shot / "report.md")
    for sheet in out_dir.glob("*/sheet.png"):
        shutil.copyfile(sheet, shot / f"{sheet.parent.name}.png")
    print("报告已写入", report)
    sys.stdout.flush()


if __name__ == "__main__":
    import os
    code = 0
    try:
        asyncio.run(main())
    except SystemExit as exc:
        code = exc.code if isinstance(exc.code, int) else 1
    except Exception as exc:
        print(type(exc).__name__, exc)
        code = 1
    sys.stdout.flush()
    os._exit(code)
