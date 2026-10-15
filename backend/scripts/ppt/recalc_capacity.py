"""按这一行还空着的宽度，重算页说明里标题类格子的字数上限。

只改 max_chars、max_lines、grow_w。不调看图模型，不改角色和 user_edited。
用法（在 backend 目录）：
  D:\\Miniconda3\\python.exe scripts\\ppt\\recalc_capacity.py
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

BACKEND = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(BACKEND))

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from services.ppt import ppt_base, ppt_pages

BUILTIN = BACKEND / "assets" / "ppt_templates" / "builtin"
WATCH = {
    3: "section_title",
    11: "title",
    13: "title",
    14: "title",
    26: "title",
}
FLOORS = {3: 8, 11: 12, 13: 12, 14: 12, 26: 8}


def _boxes(shapes, page_w: float, top: bool = True) -> list[dict]:
    from services.ppt import ppt_check

    found = []
    for shape in shapes:
        box = None
        try:
            box = {
                "x": ppt_base._inches(shape.left),
                "y": ppt_base._inches(shape.top),
                "w": ppt_base._inches(shape.width),
                "h": ppt_base._inches(shape.height),
            }
        except Exception:
            box = None
        if box and box["w"] > 0 and box["h"] > 0:
            if top and ppt_check.is_blocking_graphic(shape, page_w):
                box["graphic"] = True
            found.append(box)
        if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.GROUP:
            try:
                found.extend(_boxes(shape.shapes, page_w, top=False))
            except Exception:
                pass
    return found


def obstacles(pptx_path: Path) -> dict[int, list]:
    if not pptx_path.is_file():
        return {}
    prs = Presentation(str(pptx_path))
    page_w = ppt_base._inches(prs.slide_width)
    return {index: _boxes(slide.shapes, page_w) for index, slide in enumerate(prs.slides, 1)}


def _pair(pages_path: Path) -> tuple[Path, Path, Path]:
    stem = pages_path.name.replace(".pages.json", "")
    return pages_path, pages_path.with_name(stem + ".analysis.json"), pages_path.with_name(stem + ".pptx")


def _targets() -> list[Path]:
    found = sorted(BUILTIN.glob("*.pages.json"))
    roots = [BACKEND / "data", BACKEND / "assets"]
    for root in roots:
        if root.is_dir():
            found.extend(root.rglob("4d7e82ced452.pages.json"))
            found.extend(root.rglob("7acd6eb053b2.pages.json"))
    unique = []
    seen = set()
    for path in found:
        key = str(path.resolve())
        if key not in seen:
            seen.add(key)
            unique.append(path)
    return unique


def _slot_line(page: dict, name: str) -> str:
    for slot in page.get("slots") or []:
        if slot.get("name") == name:
            return f"{slot.get('max_chars')} grow={slot.get('grow_w')}"
    return "没有这个格子"


def recalc_file(pages_path: Path) -> list[str]:
    pages_path, analysis_path, pptx_path = _pair(pages_path)
    if not analysis_path.is_file():
        print(f"{pages_path.name} 没有分析文件，跳过", flush=True)
        return []
    pages = json.loads(pages_path.read_text(encoding="utf-8"))
    analysis = json.loads(analysis_path.read_text(encoding="utf-8"))
    before = {}
    for page in pages.get("pages") or []:
        number = page.get("page")
        if number in WATCH:
            before[number] = (
                page.get("role"),
                page.get("user_edited"),
                _slot_line(page, WATCH[number]),
            )
    ppt_pages.recalc_capacities(pages, analysis, obstacles(pptx_path))
    problems = []
    stem = pages_path.name.replace(".pages.json", "")
    if stem == "4d7e82ced452":
        by_page = {page.get("page"): page for page in pages.get("pages") or []}
        for number, name in WATCH.items():
            page = by_page.get(number) or {}
            slot = next((item for item in page.get("slots") or [] if item.get("name") == name), None)
            chars = int((slot or {}).get("max_chars") or 0)
            role_before, edited_before, old = before.get(number, ("", None, ""))
            print(
                f"模板-55 第 {number} 页 {name} {old} -> {chars} grow={(slot or {}).get('grow_w')} "
                f"角色 {page.get('role')}",
                flush=True,
            )
            if page.get("role") != role_before or page.get("user_edited") != edited_before:
                problems.append(f"第 {number} 页角色或 user_edited 变了")
            if chars < FLOORS[number]:
                problems.append(f"第 {number} 页 {name} 上限 {chars}，不到 {FLOORS[number]}")
    else:
        grown = 0
        for page in pages.get("pages") or []:
            grown += sum(1 for slot in page.get("slots") or [] if slot.get("grow_w"))
        print(f"{stem} 标题拉宽 {grown} 格", flush=True)
    tmp = pages_path.with_name(pages_path.name + ".tmp")
    tmp.write_text(json.dumps(pages, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    tmp.replace(pages_path)
    return problems


def main() -> int:
    targets = _targets()
    if not any(path.name.startswith("4d7e82ced452") for path in targets):
        print("没有找到模板-55 的页说明", flush=True)
        return 1
    problems = []
    for path in targets:
        problems.extend(recalc_file(path))
    for item in problems:
        print(item, flush=True)
    return 1 if problems else 0


if __name__ == "__main__":
    import os
    code = main()
    sys.stdout.flush()
    os._exit(code)
