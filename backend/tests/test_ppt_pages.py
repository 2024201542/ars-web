"""页说明的读写、校验、合并和导入。"""

import json
import pathlib
import shutil

import pytest
from fastapi import HTTPException
from pptx import Presentation
from pptx.util import Inches, Pt

from routers.template_library import PagesPatch, put_template_pages
from services.ppt import ppt_library, ppt_pages, ppt_template


@pytest.fixture(autouse=True)
def reset_config(monkeypatch):
    sandbox = pathlib.Path(__file__).resolve().parents[3] / ".tmp" / "pytest-env"
    monkeypatch.setenv("ARS_WEB_DB_PATH", str(sandbox / "test_ars.db"))
    monkeypatch.setenv("ARS_SKILLS_PATH", str(sandbox / "fake_skills"))
    monkeypatch.setenv("ARS_WORKSPACE_DIR", str(sandbox / "workspaces"))


@pytest.fixture()
def workdir():
    folder = pathlib.Path(__file__).resolve().parents[3] / ".tmp" / "ppt_pages_test"
    shutil.rmtree(folder, ignore_errors=True)
    folder.mkdir(parents=True, exist_ok=True)
    return folder


def _textbox(slide, x, y, w, h, text, pt):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = text
    run.font.size = Pt(pt)
    return box


def _five_page_template(path: pathlib.Path) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    cover = prs.slides.add_slide(blank)
    _textbox(cover, 1.0, 2.4, 8.0, 0.8, "汇报题目写在这里", 32)

    toc = prs.slides.add_slide(blank)
    _textbox(toc, 1.0, 0.4, 6.0, 0.6, "汇报框架", 28)
    _textbox(toc, 1.2, 1.6, 4.0, 0.4, "研究背景", 16)

    cards = prs.slides.add_slide(blank)
    _textbox(cards, 0.6, 0.3, 8.0, 0.5, "三个并列的发现", 24)
    _textbox(cards, 0.4, 2.2, 3.6, 1.2, "第一点说明", 14)
    _textbox(cards, 4.6, 2.2, 3.6, 1.2, "第二点说明", 14)
    _textbox(cards, 8.8, 2.2, 3.6, 1.2, "第三点说明", 14)

    two = prs.slides.add_slide(blank)
    _textbox(two, 0.6, 0.3, 8.0, 0.5, "左右两种做法", 24)
    _textbox(two, 0.4, 2.0, 5.5, 1.4, "左边的做法", 14)
    _textbox(two, 7.0, 2.0, 5.5, 1.4, "右边的做法", 14)

    end = prs.slides.add_slide(blank)
    _textbox(end, 3.0, 3.0, 6.0, 0.8, "谢谢各位", 32)
    prs.save(str(path))


def _template(path: pathlib.Path, analysis: dict) -> dict:
    return {
        "id": analysis["id"],
        "name": analysis["name"],
        "source": "user",
        "file_path": str(path),
        "specimens": analysis["specimens"],
    }


def test_empty_pages_guesses_roles(workdir):
    path = workdir / "five.pptx"
    _five_page_template(path)
    analysis = ppt_template.analyze_pptx(path, tid="five", name="五页", source="user")
    data = ppt_pages.empty_pages(_template(path, analysis), analysis)
    roles = [page["role"] for page in data["pages"]]
    forms = [page["form"] for page in data["pages"]]
    assert roles == ["cover", "toc", "content", "content", "end"]
    assert forms[2] == "cards"
    assert data["pages"][2]["repeat"] == 3
    assert any(slot["name"] == "title" for slot in data["pages"][2]["slots"])
    assert data["parsed"] is False


def test_validate_pages_reports_chinese_errors():
    data = {"pages": [
        {"page": 1, "role": "nope", "form": "cards", "narrative_fit": [], "slots": []},
        {"page": 1, "role": "content", "form": "cards", "narrative_fit": [],
         "slots": [{"name": "item0.head"}]},
        {"page": 99, "role": "content", "form": "cards", "narrative_fit": [], "slots": []},
    ]}
    errors = ppt_pages.validate_pages(data, 5)
    assert any("不在词表里" in item for item in errors)
    assert any("item0.head" in item for item in errors)
    assert any("99" in item for item in errors)
    clean = {"pages": [
        {"page": 1, "role": "content", "form": "cards", "narrative_fit": ["finding"],
         "slots": [{"name": "item1.head"}]},
    ]}
    assert ppt_pages.validate_pages(clean, 5) == []


def test_merge_keeps_user_role():
    old = {"pages": [
        {"page": 3, "role": "summary", "user_edited": True, "slots": [{"frame": 1, "name": "title"}]},
        {"page": 4, "role": "content", "user_edited": False},
    ]}
    new = {"pages": [
        {"page": 3, "role": "content", "user_edited": False, "slots": [{"frame": 1, "name": "body"}]},
        {"page": 4, "role": "content", "user_edited": False},
    ]}
    merged = ppt_pages.merge_user_edits(old, new)
    assert merged["pages"][0]["role"] == "summary"
    assert merged["pages"][0]["slots"][0]["name"] == "title"
    assert merged["pages"][1]["role"] == "content"


def test_import_warns_when_sha_differs_and_skips_overflow():
    result = ppt_pages.import_document(
        {"template_sha1": "aaaa", "pages": [{"page": 2, "role": "summary"}, {"page": 25}]},
        current_sha="bbbb", page_count=10)
    assert result["warnings"]
    assert result["skipped"] == [25]
    assert [page["page"] for page in result["imported"]] == [2]


def test_builtin_pages_cannot_be_saved(workdir):
    path = workdir / "builtin.pptx"
    path.write_bytes(b"not-a-real-pptx")
    with pytest.raises(PermissionError):
        ppt_pages.save_pages({"source": "builtin", "file_path": str(path)}, {"pages": []})


@pytest.mark.asyncio
async def test_builtin_put_returns_400(monkeypatch, workdir):
    path = workdir / "builtin.pptx"
    Presentation().save(str(path))
    item = {"id": "builtin-x", "source": "builtin", "file_path": str(path), "name": "内置", "specimens": []}
    monkeypatch.setattr(ppt_library, "find", lambda user_id, template_id: item)
    before = path.read_bytes()
    with pytest.raises(HTTPException) as caught:
        await put_template_pages("builtin-x", PagesPatch(pages=[{"page": 1, "role": "summary"}]), "user")
    assert caught.value.status_code == 400
    assert path.read_bytes() == before


def _narrow_title(x=0.8, w=0.6):
    return {"i": 0, "x": x, "y": 0.4, "w": w, "h": 0.45, "size": 18}


def test_narrow_title_grows_when_the_row_is_empty():
    frame = _narrow_title()
    slot = ppt_pages._make_slot("title", 0, frame, True, [frame], 13.333)
    assert slot["max_chars"] >= 12
    assert slot["grow_w"] > frame["w"]


def test_title_stays_when_the_next_shape_touches():
    frame = _narrow_title()
    neighbor = {"i": 1, "x": 1.4, "y": 0.4, "w": 2.0, "h": 0.45, "size": 18}
    blocked = ppt_pages._make_slot("title", 0, frame, True, [frame, neighbor], 13.333)
    raw = ppt_pages._make_slot("title", 0, frame, True)
    assert blocked["max_chars"] == raw["max_chars"]
    assert "grow_w" not in blocked


def test_shape_on_another_row_does_not_block_the_title():
    frame = _narrow_title()
    below = {"i": 1, "x": 0.8, "y": 1.6, "w": 4.0, "h": 1.2, "size": 14}
    slot = ppt_pages._make_slot("title", 0, frame, True, [frame, below], 13.333)
    assert slot["max_chars"] >= 12
    assert slot["grow_w"] > frame["w"]


def test_large_font_title_still_fits_a_sentence():
    # 章节名、结尾标题的框已经顶到页边，48pt / 72pt 按原字号只能写 5 个字。
    section = {"i": 1, "x": 7.77, "y": 2.95, "w": 5.38, "h": 0.91, "size": 48}
    ending = {"i": 2, "x": 0.4, "y": 2.53, "w": 8.57, "h": 1.31, "size": 72}
    assert ppt_pages._make_slot("section_title", 1, section, True)["max_chars"] >= 8
    assert ppt_pages._make_slot("title", 2, ending, True)["max_chars"] >= 8
    boxed = {"i": 0, "x": 0.8, "y": 0.4, "w": 0.6, "h": 0.45, "size": 48}
    neighbor = {"i": 1, "x": 1.4, "y": 0.4, "w": 2.0, "h": 0.45, "size": 18}
    blocked = ppt_pages._make_slot("title", 0, boxed, True, [boxed, neighbor], 13.333)
    assert blocked["max_chars"] < 12


def test_body_slot_does_not_grow():
    frame = {"i": 0, "x": 0.8, "y": 1.5, "w": 2.0, "h": 1.2, "size": 14}
    slot = ppt_pages._make_slot("item1.body", 0, frame, True, [frame], 13.333)
    assert "grow_w" not in slot
    chars, lines, _pt = ppt_pages._slot_capacity(frame)
    assert slot["max_chars"] == chars
    assert slot["max_lines"] == lines


def test_duplicate_slot_name_is_renamed_and_warned():
    page = {
        "page": 8, "role": "content", "form": "compare", "narrative_fit": [],
        "parse_warnings": ["叙事角色 analysis 不在词表里"],
        "slots": [
            {"name": "item4.body", "frame": 14, "required": True},
            {"name": "item2.body", "frame": 11, "required": True},
            {"name": "item4.body", "frame": 14, "required": True},
        ],
    }
    added = ppt_pages.dedupe_slot_names(page)
    assert [slot["name"] for slot in page["slots"]] == ["item4.body", "item2.body", "item4.body_2"]
    assert added == ["格子名 item4.body 重复，已改成 item4.body_2"]
    assert "叙事角色 analysis 不在词表里" in page["parse_warnings"]
    assert added[0] in page["parse_warnings"]
    assert ppt_pages.dedupe_slot_names(page) == []
    errors = ppt_pages.validate_pages({"pages": [page]}, 8)
    assert errors == []


def test_recalc_keeps_the_user_role():
    pages = {"pages": [{
        "page": 1, "role": "summary", "user_edited": True, "description": "用户改过",
        "slots": [{"name": "title", "frame": 0, "max_chars": 2, "max_lines": 1, "required": True}],
    }]}
    analysis = {"slide_width": 13.333, "specimens": [{
        "index": 0,
        "frames": [_narrow_title()],
    }]}
    ppt_pages.recalc_capacities(pages, analysis)
    page = pages["pages"][0]
    assert page["role"] == "summary"
    assert page["user_edited"] is True
    assert page["description"] == "用户改过"
    assert page["slots"][0]["max_chars"] >= 12
    assert page["slots"][0]["grow_w"] > 0.6


def test_markdown_and_catalog():
    data = {"template_name": "样例", "deck": {}, "pages": [
        {"page": 5, "role": "content", "form": "cards", "usable": True,
         "narrative_fit": ["finding"], "description": "三张卡片",
         "slots": [{"name": "title", "max_chars": 18}]},
        {"page": 6, "role": "content", "form": "text", "usable": False,
         "narrative_fit": [], "description": "不用", "slots": []},
    ]}
    text = ppt_pages.to_markdown(data)
    assert "第 5 页" in text
    assert "并列卡片" in text
    catalog = ppt_pages.catalog_for_prompt(data, limit_chars=3000)
    assert "第5页" in catalog
    assert "第6页" not in catalog
    assert len(ppt_pages.catalog_for_prompt(data, limit_chars=8)) <= 8


@pytest.mark.asyncio
async def test_delete_template_removes_pages_file(monkeypatch, workdir):
    folder = workdir / "ppt_templates"
    folder.mkdir()
    monkeypatch.setattr(ppt_library, "user_dir", lambda user_id: workdir)
    (folder / "abc.pptx").write_bytes(b"pptx")
    (folder / "abc.analysis.json").write_text("{}", encoding="utf-8")
    (folder / "abc.pages.json").write_text("{}", encoding="utf-8")

    from routers.template_library import delete_template
    result = await delete_template("abc", "user")
    assert result["ok"] is True
    assert not (folder / "abc.pages.json").exists()
    assert not (folder / "abc.pptx").exists()


def test_closing_title_keeps_a_short_template_heading():
    assert ppt_pages.closing_title("感谢您的观看") == "感谢您的观看"
    assert ppt_pages.closing_title("谢谢") == "谢谢"
    assert ppt_pages.closing_title("物理学习记忆机制与干预") == "谢谢"
    assert ppt_pages.closing_title("") == "谢谢"


def test_cover_title_shrinks_to_stop_before_an_overlapping_graphic():
    frame = {"i": 2, "x": -0.59, "y": 2.53, "w": 8.57, "h": 1.31, "size": 72}
    arc = {"i": 9, "x": 7.85, "y": -1.56, "w": 2.5, "h": 10.51, "graphic": True}
    bleed = {"i": 8, "x": 0.05, "y": 0.0, "w": 13.23, "h": 7.5, "graphic": True}
    open_slot = ppt_pages._make_slot("title", 2, frame, True, [frame], 13.333, clip_graphics=True)
    clipped = ppt_pages._make_slot(
        "title", 2, frame, True, [frame, arc, bleed], 13.333, clip_graphics=True)
    assert open_slot["max_chars"] >= 12
    assert clipped["max_chars"] <= 6
    assert "grow_w" not in clipped
