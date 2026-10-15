"""逐页看图解析。看图调用全部打桩，不连网。"""

import asyncio
import json
import pathlib
import shutil

import pytest
from fastapi import HTTPException
from PIL import Image

from routers.template_library import (
    ParseRequest,
    start_template_parse,
    template_parse_status,
)
from services.ppt import ppt_jobs, ppt_library, ppt_pages, ppt_template, ppt_vision


@pytest.fixture(autouse=True)
def reset_config(monkeypatch):
    sandbox = pathlib.Path(__file__).resolve().parents[3] / ".tmp" / "pytest-env"
    monkeypatch.setenv("ARS_WEB_DB_PATH", str(sandbox / "test_ars.db"))
    monkeypatch.setenv("ARS_SKILLS_PATH", str(sandbox / "fake_skills"))
    monkeypatch.setenv("ARS_WORKSPACE_DIR", str(sandbox / "workspaces"))
    ppt_jobs._jobs.clear()
    ppt_jobs._tasks.clear()


@pytest.fixture()
def workdir(monkeypatch):
    folder = pathlib.Path(__file__).resolve().parents[3] / ".tmp" / "ppt_vision_parse_test"
    shutil.rmtree(folder, ignore_errors=True)
    folder.mkdir(parents=True, exist_ok=True)
    monkeypatch.setattr(ppt_vision, "CACHE_DIR", folder / "vision")
    return folder


def _analysis(tags: list[str]) -> dict:
    specs = []
    for index, tag in enumerate(tags):
        specs.append({
            "index": index,
            "tag": tag,
            "frames": [
                {"i": 0, "shape_id": 10 + index, "x": 1, "y": 0.4, "w": 6, "h": 0.8,
                 "size": 28, "text": "示例英文"},
                {"i": 1, "shape_id": 20 + index, "x": 1, "y": 1.6, "w": 3, "h": 0.5,
                 "size": 16, "text": "卡片"},
            ],
            "images": [],
        })
    return {"specimens": specs, "slide_width": 13.333, "slide_height": 7.5}


def _template(folder: pathlib.Path, tags: list[str]) -> dict:
    path = folder / "tpl.pptx"
    path.write_bytes(b"pptx-bytes")
    item = {
        "id": "tpl",
        "source": "user",
        "file_path": str(path),
        "name": "测试模板",
    }
    item.update(_analysis(tags))
    return item


def test_mark_frames_keeps_size_and_draws_red(workdir):
    png = workdir / "page-1.png"
    Image.new("RGB", (200, 100), "white").save(png)
    out = ppt_vision.mark_frames(
        png,
        [{"i": 0, "x": 0.5, "y": 0.4, "w": 2, "h": 0.8}],
        10, 5,
    )
    with Image.open(out) as marked:
        assert marked.size == (200, 100)
        raw = marked.convert("RGB").tobytes()
    red = sum(1 for i in range(0, len(raw), 3) if raw[i] > 200 and raw[i + 1] < 40 and raw[i + 2] < 40)
    assert red > 0


def test_apply_page_parse_writes_slots_and_title_index():
    analysis = _analysis(["cards-3"])
    pages = ppt_pages.empty_pages({"id": "t", "name": "t"}, analysis)
    result = {
        "role": "content",
        "form": "cards",
        "frames": [
            {"id": "F0", "slot": "title"},
            {"id": "F1", "slot": "item1.head"},
        ],
    }
    ppt_pages.apply_page_parse(pages, analysis, 1, result)
    names = [slot["name"] for slot in pages["pages"][0]["slots"]]
    assert "title" in names
    assert "item1.head" in names
    assert analysis["specimens"][0]["title_i"] == 0
    assert pages["pages"][0]["vision_ok"] is True


def test_invalid_role_falls_back_and_missing_is_only_a_warning():
    analysis = _analysis(["cover"])
    pages = ppt_pages.empty_pages({"id": "t", "name": "t"}, analysis)
    ppt_pages.apply_page_parse(pages, analysis, 1, {
        "role": "foo",
        "form": "none",
        "frames": [],
        "missing": [{"slot": "body", "box": [0.1, 0.2, 0.3, 0.1]}],
    })
    page = pages["pages"][0]
    assert page["role"] == "cover"
    assert any("不在词表里" in item for item in page["parse_warnings"])
    assert any("图片上的字" in item for item in page["parse_warnings"])
    assert page["slots"] == []


def test_group_text_comes_after_top_level_frames():
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    class _Text:
        def __init__(self, value):
            self.text = value

    class _Shape:
        def __init__(self, kind, text="", children=None, sid=0):
            self.shape_type = kind
            self.shape_id = sid
            self.shapes = children or []
            self.text_frame = _Text(text)
            self.is_placeholder = False

        @property
        def has_text_frame(self):
            return self.shape_type != MSO_SHAPE_TYPE.GROUP

    class _Slide:
        def __init__(self, shapes):
            self.shapes = shapes

    title = _Shape(MSO_SHAPE_TYPE.TEXT_BOX, "标题", sid=1)
    inside = _Shape(MSO_SHAPE_TYPE.TEXT_BOX, "课题背景", sid=13)
    group = _Shape(MSO_SHAPE_TYPE.GROUP, children=[inside], sid=6)
    found = ppt_template.iter_text_shapes(_Slide([group, title]))
    assert [shape.shape_id for shape in found] == [1, 13]


def test_unknown_frame_id_is_dropped():
    analysis = _analysis(["cover"])
    pages = ppt_pages.empty_pages({"id": "t", "name": "t"}, analysis)
    ppt_pages.apply_page_parse(pages, analysis, 1, {
        "role": "cover",
        "form": "none",
        "frames": [{"id": "F9", "slot": "title"}],
    })
    page = pages["pages"][0]
    assert page["slots"] == []
    assert any("F9" in item for item in page["parse_warnings"])
    assert analysis["specimens"][0].get("title_i") != 9


def test_decor_clear_goes_to_decor():
    analysis = _analysis(["content"])
    pages = ppt_pages.empty_pages({"id": "t", "name": "t"}, analysis)
    ppt_pages.apply_page_parse(pages, analysis, 1, {
        "role": "content",
        "form": "text",
        "frames": [{"id": "F0", "slot": "decor_clear"}],
    })
    page = pages["pages"][0]
    assert page["decor"][0]["action"] == "clear"
    assert page["decor"][0]["frame"] == 0
    assert all(slot["name"] != "decor_clear" for slot in page["slots"])


def test_slot_capacity_title_box_is_in_range():
    # 和 ppt_template._capacity_chars 比：那边不乘 0.85，行距用 1.45。
    # 6 英寸 × 0.8 英寸、28pt 时两边都是 1 行、每行 15 字。
    # 本函数再乘 0.85，大约 12 字，和旧函数差 20%，不到 30%。标题再乘 0.9 发生在写入 title 格子时。
    frame = {"w": 6, "h": 0.8, "size": 28}
    chars, lines = ppt_pages.slot_capacity(frame, 28)
    per_line = int((6 * 72) / 28)
    assert 1 <= lines <= 2
    assert 14 <= per_line <= 16
    assert 8 <= chars <= 20
    old = ppt_template._capacity_chars(frame)
    assert old > 0
    assert abs(chars - old) / old <= 0.30


@pytest.mark.asyncio
async def test_key_error_on_first_page_stops_the_job(workdir, monkeypatch):
    calls = []

    async def fake(user_id, analysis, template_path, page, model, force=False):
        calls.append(page)
        if page == 1:
            raise RuntimeError("密钥无效")
        return {"ok": True, "result": {"role": "content", "form": "text", "frames": []}}

    monkeypatch.setattr(ppt_vision, "parse_page", fake)
    item = _template(workdir, ["cover", "content"])
    progress = {}
    await ppt_vision.run_parse("user", item, "qwen3-vl-plus", [1, 2], False, progress)
    assert progress["status"] == "failed"
    assert calls == [1]
    assert "密钥" in progress["errors"][0]["error"]


@pytest.mark.asyncio
async def test_timeout_retries_once_then_continues(workdir, monkeypatch):
    calls = []

    async def fake(user_id, analysis, template_path, page, model, force=False):
        calls.append(page)
        if page == 3:
            return {"ok": False, "error": "TimeoutError: 超时"}
        return {"ok": True, "result": {"role": "content", "form": "bullets", "frames": []}}

    monkeypatch.setattr(ppt_vision, "parse_page", fake)
    item = _template(workdir, ["cover", "toc", "section", "content"])
    progress = {}
    await ppt_vision.run_parse("user", item, "qwen3-vl-plus", [1, 2, 3, 4], False, progress)
    assert calls.count(3) == 2
    assert 4 in calls
    assert progress["done"] == progress["total"] == 4
    assert progress["errors"][0]["page"] == 3
    saved = json.loads((workdir / "tpl.pages.json").read_text(encoding="utf-8"))
    assert saved["parse"]["done"] == 4

    ppt_jobs._jobs.clear()
    monkeypatch.setattr(ppt_library, "find", lambda user_id, template_id: item)
    view = await template_parse_status("tpl", "user")
    assert view["done"] == 4
    assert view["current_page"] == 4


@pytest.mark.asyncio
async def test_cache_hit_does_not_call_vision(workdir, monkeypatch):
    path = workdir / "cached.pptx"
    path.write_bytes(b"cached-pptx")
    calls = []

    async def boom(user_id, image_path, model_id, prompt=ppt_vision.PROMPT, max_tokens=900):
        calls.append(1)
        return {}, "不应调用"

    monkeypatch.setattr(ppt_vision, "_vision_call", boom)
    cache = ppt_vision._parse_cache_path(path, 1)
    cache.write_text(json.dumps({"result": {"role": "content", "form": "text", "frames": []}}),
                     encoding="utf-8")
    out = await ppt_vision.parse_page(
        "user", _analysis(["content"]), path, 1, "qwen3-vl-plus",
    )
    assert calls == []
    assert out["ok"] is True
    assert out["cached"] is True
    assert out["result"]["role"] == "content"


@pytest.mark.asyncio
async def test_running_job_returns_the_same_id():
    gate = asyncio.Event()

    async def factory(progress):
        progress["total"] = 1
        await gate.wait()
        progress["status"] = "done"

    first = ppt_jobs.start("parse:user:tpl", factory)
    second = ppt_jobs.start("parse:user:tpl", factory)
    assert first == second
    gate.set()
    await ppt_jobs.wait("parse:user:tpl")
    assert ppt_jobs.status("parse:user:tpl")["status"] == "done"


@pytest.mark.asyncio
async def test_builtin_parse_returns_400_and_file_status_survives(monkeypatch, workdir):
    builtin = {"id": "builtin-x", "source": "builtin", "file_path": str(workdir / "b.pptx"),
               "name": "内置", "specimens": [{"index": 0, "tag": "cover", "frames": []}]}
    (workdir / "b.pptx").write_bytes(b"pptx")
    monkeypatch.setattr(ppt_library, "find", lambda user_id, template_id: builtin)
    with pytest.raises(HTTPException) as caught:
        await start_template_parse("builtin-x", ParseRequest(), "user")
    assert caught.value.status_code == 400

    user = _template(workdir, ["cover"])
    user["id"] = "file-only"
    user["file_path"] = str(workdir / "file-only.pptx")
    pathlib.Path(user["file_path"]).write_bytes(b"pptx")
    doc = ppt_pages.empty_pages(user, user)
    doc["parse"] = {
        "status": "running", "done": 12, "total": 26, "current_page": 13,
        "errors": [{"page": 7, "error": "超时"}],
    }
    ppt_pages.save_pages(user, doc)
    monkeypatch.setattr(ppt_library, "find", lambda user_id, template_id: user)
    view = await template_parse_status("file-only", "user")
    assert view["status"] == "running"
    assert view["done"] == 12
    assert view["current_page"] == 13
    assert view["total"] == 26


def test_parse_prompt_keeps_toc_index_and_blocks_empty_section():
    assert ppt_vision.PROMPT_VERSION >= 4
    assert "-01-" in ppt_vision.PARSE_PROMPT
    assert "不能是 section" in ppt_vision.PARSE_PROMPT


def test_toc_index_is_kept_and_section_without_a_title_becomes_content():
    analysis = _analysis(["cards-3"])
    analysis["specimens"][0]["frames"][0]["text"] = "-01-"
    analysis["specimens"][0]["frames"][1]["text"] = "您的内容打在这里"
    pages = ppt_pages.empty_pages({"id": "t", "name": "t"}, analysis)
    ppt_pages.apply_page_parse(pages, analysis, 1, {
        "role": "section",
        "form": "bullets",
        "frames": [
            {"id": "F0", "slot": "item1.head"},
            {"id": "F1", "slot": "item1.body"},
        ],
    })
    page = pages["pages"][0]
    assert page["role"] == "content"
    assert page["decor"][0]["action"] == "keep"
    assert page["decor"][0]["text"] == "-01-"
    assert any("章节名" in item for item in page["parse_warnings"])


def test_finish_deck_leaves_summary_empty_when_the_model_stalls(monkeypatch):
    from services.ppt import ppt_deck_parse

    async def hang(*_args, **_kwargs):
        await asyncio.sleep(2)

    monkeypatch.setattr(ppt_deck_parse, "deck_story", hang)
    monkeypatch.setattr(ppt_vision, "SUMMARY_TIMEOUT", 0.05)
    pages = {"pages": [{"page": 1, "role": "cover", "form": "text"}]}
    asyncio.run(ppt_vision._finish_deck("user", pages, {"specimens": []}, ""))
    assert pages["deck"]["summary"] == ""
