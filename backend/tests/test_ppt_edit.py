"""修改意见拆成操作之后，只改点到的页。"""

import copy
import json
import pathlib

import pytest

from services.ppt import ppt_edit


@pytest.fixture(autouse=True)
def reset_config(monkeypatch):
    sandbox = pathlib.Path(__file__).resolve().parents[3] / ".tmp" / "pytest-env"
    monkeypatch.setenv("ARS_WEB_DB_PATH", str(sandbox / "test_ars.db"))
    monkeypatch.setenv("ARS_SKILLS_PATH", str(sandbox / "fake_skills"))
    monkeypatch.setenv("ARS_WORKSPACE_DIR", str(sandbox / "workspaces"))


def _plan():
    slides = []
    for index in range(1, 9):
        slides.append({
            "layout": "bullets",
            "role": "finding",
            "title": f"第{index}页标题",
            "lead": f"承接{index}",
            "notes": f"第{index}页第一句。第{index}页第二句。",
            "slots": {"title": f"第{index}页标题", "body": "字" * 20},
        })
    slides[0]["layout"] = "cover"
    slides[0]["role"] = "cover"
    slides[1]["layout"] = "toc"
    slides[1]["role"] = "toc"
    slides[1]["items"] = [{"label": "旧章"}]
    return {"title": "测试", "slides": slides, "storyline": {"deck_type": "lecture", "slots": []}}


def _dump(slide):
    return json.dumps(slide, ensure_ascii=False, sort_keys=True)


async def test_shorten_only_changes_page_three():
    plan = _plan()
    before = copy.deepcopy(plan["slides"])
    updated, changes = await ppt_edit.apply_ops(plan, [{"op": "shorten", "page": 3}])
    assert updated["slides"][2]["slots"]["body"] != before[2]["slots"]["body"]
    assert len(updated["slides"][2]["slots"]["body"]) < len(before[2]["slots"]["body"])
    assert updated["slides"][2]["title"] == before[2]["title"]
    for index, slide in enumerate(updated["slides"]):
        if index == 2:
            continue
        assert _dump(slide) == _dump(before[index])
    assert changes[0]["page"] == 3


async def test_add_background_page_updates_toc_and_template():
    plan = _plan()
    catalog = {"pages": [{"page": 4, "role": "content", "form": "bullets", "usable": True,
                          "slots": [{"name": "title"}, {"name": "body"}]}]}
    updated, changes = await ppt_edit.apply_ops(
        plan, [{"op": "add_page", "after": 2, "role": "background", "instruction": "背景介绍"}],
        catalog=catalog)
    assert len(updated["slides"]) == 9
    added = updated["slides"][2]
    assert added["role"] == "background"
    assert added["template_page"] == 4
    named = [slide for slide in updated["slides"]
             if slide.get("role") in {"background", "section"} or slide.get("layout") == "section"]
    toc = next(slide for slide in updated["slides"] if slide.get("layout") == "toc")
    assert len(toc["items"]) == len(named)
    assert any("新增第 3 页" in item["what"] for item in changes)
    assert ppt_edit.changed_pages_of(plan, updated, changes) == [3, 2]


async def test_bridge_only_changes_lead_and_first_sentence():
    plan = _plan()
    before = copy.deepcopy(plan["slides"])
    updated, _changes = await ppt_edit.apply_ops(plan, [{"op": "add_bridge", "page": 6}])
    slide = updated["slides"][5]
    assert slide["lead"] != before[5]["lead"]
    assert slide["notes"].startswith("接着上一页。")
    assert slide["notes"].endswith("第6页第二句。")
    assert slide["title"] == before[5]["title"]
    assert slide["slots"] == before[5]["slots"]
    for index, item in enumerate(updated["slides"]):
        if index == 5:
            continue
        assert _dump(item) == _dump(before[index])


async def test_delete_then_shorten_uses_original_page_numbers():
    plan = _plan()
    updated, _changes = await ppt_edit.apply_ops(plan, [
        {"op": "delete_page", "page": 8},
        {"op": "shorten", "page": 3},
    ])
    titles = [slide["title"] for slide in updated["slides"]]
    assert "第8页标题" not in titles
    page3 = next(slide for slide in updated["slides"] if slide["title"] == "第3页标题")
    page4 = next(slide for slide in updated["slides"] if slide["title"] == "第4页标题")
    assert len(page3["slots"]["body"]) < 20
    assert page4["slots"]["body"] == "字" * 20


async def test_bad_json_falls_back_to_shorten(monkeypatch):
    async def fake_text(*_args, **_kwargs):
        return "不是 JSON", "stub", ""

    monkeypatch.setattr("services.ppt.ppt_plan._ask_text", fake_text)
    plan = _plan()
    planned = await ppt_edit.plan_ops("第3页文字太挤", plan, None, "m", "u")
    assert planned["ops"] == [{"op": "shorten", "page": 3}]
    updated, _changes = await ppt_edit.apply_ops(plan, planned["ops"])
    assert len(updated["slides"][2]["slots"]["body"]) < 20
    assert updated["slides"][3]["slots"]["body"] == "字" * 20


def _deck(path, title):
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()
    blank = prs.slide_layouts[6]
    for index in range(1, 4):
        slide = prs.slides.add_slide(blank)
        box = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(8), Inches(0.6))
        box.name = f"ars:s{index}:title"
        box.text_frame.text = title if index == 3 else f"页{index}"
        body = slide.shapes.add_textbox(Inches(0.5), Inches(1.4), Inches(8), Inches(0.6))
        body.name = f"ars:s{index}:body"
        body.text_frame.text = "字" * 20
        for shape in (box, body):
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(20)
    prs.save(str(path))


async def test_hand_edit_title_stays_locked_until_named():
    folder = pathlib.Path(__file__).resolve().parents[3] / ".tmp" / "pytest-env" / "e5"
    folder.mkdir(parents=True, exist_ok=True)
    deck = folder / "hand.pptx"
    _deck(deck, "手改标题")
    plan = _plan()
    plan["slides"][2]["slots"]["title"] = "原标题"
    plan["slides"][2]["title"] = "原标题"
    assert ppt_edit.sync_manual_edits(deck, plan) is True
    assert plan["slides"][2]["slots"]["title"] == "手改标题"
    assert "title" in plan["slides"][2]["locked_slots"]
    ppt_edit.stamp_plan_mark(deck, "e5/plan.json", "site")
    assert ppt_edit.read_plan_mark(deck)["path"] == "e5/plan.json"
    from pptx import Presentation
    opened = Presentation(str(deck))
    ppt_edit._save_keeping_mark(opened, deck)
    assert ppt_edit.read_plan_mark(deck)["source"] == "site"

    updated, _changes = await ppt_edit.apply_ops(plan, [{"op": "apply_all"}])
    assert updated["slides"][2]["slots"]["title"] == "手改标题"
    assert updated["slides"][2]["title"] == "手改标题"
    assert len(updated["slides"][2]["slots"]["body"]) < 20
    assert "title" in updated["slides"][2]["locked_slots"]

    renamed, _changes = await ppt_edit.apply_ops(
        updated, [{"op": "rewrite", "page": 3, "instruction": "把标题改成点名之后的标题"}])
    assert renamed["slides"][2]["slots"]["title"] == "点名之后的标题"
    assert renamed["slides"][2]["title"] == "点名之后的标题"
    assert "title" not in (renamed["slides"][2].get("locked_slots") or [])


def test_plan_mark_falls_back_to_sibling_name():
    folder = pathlib.Path(__file__).resolve().parents[3] / ".tmp" / "pytest-env" / "e5"
    folder.mkdir(parents=True, exist_ok=True)
    deck = folder / "sibling.pptx"
    _deck(deck, "原标题")
    sibling = folder / "sibling.ppt计划.json"
    sibling.write_text("{}", encoding="utf-8")
    found = ppt_edit.find_plan_for_deck(deck)
    assert found["absolute"] == str(sibling)
    assert found["path"] == "sibling.ppt计划.json"


def test_unnamed_box_is_marked_when_text_changes():
    from pptx import Presentation
    from pptx.util import Inches

    folder = pathlib.Path(__file__).resolve().parents[3] / ".tmp" / "pytest-env" / "e5"
    folder.mkdir(parents=True, exist_ok=True)
    deck = folder / "plain.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(8), Inches(0.6))
    box.name = "TextBox 1"
    box.text_frame.text = "旧字"
    prs.save(str(deck))
    before = ppt_edit.slide_snapshot(deck)
    box.text_frame.text = "新字"
    prs.save(str(deck))
    plan = {"slides": [{"layout": "bullets", "title": "旧字", "slots": {}}]}
    assert ppt_edit.sync_manual_edits(deck, plan, before) is True
    assert plan["slides"][0]["manual_notes"]
    assert "第 1 页有手改内容" in ppt_edit.manual_loss_note(plan)


def test_missing_output_tells_user_to_compose():
    text = ppt_edit.result_reply({"slides": []}, "已改第 3 页", [{"page": 3, "what": "缩短了这一页的文字"}])
    assert "点生成 PPT" in text


async def test_fix_layout_still_overlapping_shortens():
    plan = _plan()
    updated, changes = await ppt_edit.apply_ops(plan, [{"op": "fix_layout", "page": 3}])
    assert len(changes) == 2
    assert "出界" in changes[0]["what"] or "叠字" in changes[0]["what"]
    assert "缩短" in changes[1]["what"]
    assert len(updated["slides"][2]["slots"]["body"]) < 20
    assert updated["slides"][0]["slots"]["body"] == "字" * 20
