"""底板档案：自动猜页、清空安全区、残留示例字、存档往返。"""

import pathlib
import shutil

import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

from services.ppt import ppt_base


@pytest.fixture(autouse=True)
def reset_config(monkeypatch):
    """覆盖 conftest 的同名 fixture：不用 pytest 的 tmp_path（沙箱里 0o700 进不去）。"""
    sandbox = pathlib.Path(__file__).resolve().parents[3] / ".tmp" / "pytest-env"
    monkeypatch.setenv("ARS_WEB_DB_PATH", str(sandbox / "test_ars.db"))
    monkeypatch.setenv("ARS_SKILLS_PATH", str(sandbox / "fake_skills"))
    monkeypatch.setenv("ARS_WORKSPACE_DIR", str(sandbox / "workspaces"))


@pytest.fixture()
def workdir():
    folder = pathlib.Path(__file__).resolve().parents[3] / ".tmp" / "ppt_base_test"
    shutil.rmtree(folder, ignore_errors=True)
    folder.mkdir(parents=True, exist_ok=True)
    return folder


def _textbox(slide, x, y, w, h, text, pt):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = text
    run.font.size = Pt(pt)
    return box


def _rect(slide, x, y, w, h, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y),
                                   Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(color)
    shape.line.fill.background()
    return shape


def _make_template(path: pathlib.Path) -> pathlib.Path:
    """造一份 4 页小模板：封面 / 转场 / 内容（带示例字和页脚）/ 结尾。"""
    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    accent = "9B1B30"

    cover = prs.slides.add_slide(blank)
    _rect(cover, 0.2, 1.0, 0.5, 2.0, accent)
    _textbox(cover, 1.6, 2.4, 8.0, 0.8, "中国人民大学模板", 32)
    _textbox(cover, 1.6, 3.6, 8.0, 0.5,
             "Thesis defense template for graduation project", 14)
    _textbox(cover, 1.6, 4.6, 6.0, 0.4, "答辩学生：苗猫大丸子", 14)
    _textbox(cover, 1.6, 5.0, 6.0, 0.4, "指导老师：苗猫大丸子", 14)

    section = prs.slides.add_slide(blank)
    _textbox(section, 1.0, 2.4, 1.4, 1.0, "01", 44)
    _textbox(section, 2.4, 2.6, 5.0, 0.6, "研究现状", 24)
    _textbox(section, 1.2, 3.8, 4.0, 0.4, "摘要", 12)
    _textbox(section, 1.2, 4.2, 4.0, 0.4, "引言", 12)

    content = prs.slides.add_slide(blank)
    _rect(content, 11.5, 0.3, 1.2, 0.5, accent)          # 右上角校徽位
    _textbox(content, 0.74, 0.3, 6.0, 0.6, "请输入你的标题", 20)
    _rect(content, 0.74, 1.13, 12.0, 0.03, accent)       # 页眉横线
    _textbox(content, 1.5, 2.0, 8.0, 1.5, "这是一段模板自带的示例正文。", 16)
    _rect(content, 3.0, 3.5, 2.0, 1.2, accent)           # 安全区里的装饰块
    _textbox(content, 0.74, 6.95, 2.0, 0.3, "2025/5/23", 10)
    _textbox(content, 9.5, 6.95, 1.0, 0.3, "5", 10)

    end = prs.slides.add_slide(blank)
    _textbox(end, 1.6, 2.5, 6.0, 0.8, "谢谢", 30)
    _textbox(end, 1.6, 4.6, 6.0, 0.4, "答辩学生：苗猫大丸子", 14)
    _textbox(end, 1.6, 5.0, 6.0, 0.4, "指导老师：苗猫大丸子", 14)

    prs.save(str(path))
    return path


def _big_shapes_in_safe(slide, safe, min_height=0.6):
    """安全区里"高过门槛且中心落在区内"的形状。"""
    _, y0, _, y1 = safe
    out = []
    for shape in slide.shapes:
        top = ppt_base._inches(shape.top)
        height = ppt_base._inches(shape.height)
        if y0 - 0.1 < top + height / 2 < y1 + 0.15 and height > min_height:
            out.append(shape)
    return out


def test_guess_base_finds_four_pages(workdir):
    path = _make_template(workdir / "demo.pptx")
    base = ppt_base.guess_base(path)
    assert base["pages"] == {"cover": 1, "section": 2, "content": 3, "end": 4}
    assert base["confirmed"] is False
    assert base["slide_size"] == [13.33, 7.5]


def test_guess_base_safe_area_and_colors(workdir):
    path = _make_template(workdir / "demo.pptx")
    base = ppt_base.guess_base(path)
    x0, y0, x1, y1 = base["safe_area"]
    assert x0 == 0.7 and x1 == 12.63
    assert y0 == 1.46      # 页眉最低处（横线底 1.16）+ 0.3
    assert y1 == 6.75      # 页脚最高处（6.95）− 0.2
    assert base["colors"]["accent"] == "#9B1B30"
    # 按计划"accent 与白 1:9 混"算出来的值（demo 里硬编码的是 #F6ECEE，相当于 0.917）
    assert base["colors"]["accent_soft"] == "#F5E8EA"


def test_guess_base_slots_keep_template_text(workdir):
    path = _make_template(workdir / "demo.pptx")
    slots = ppt_base.guess_base(path)["slots"]
    assert slots["cover"]["title"] == "中国人民大学模板"
    assert slots["cover"]["subtitle"] == "Thesis defense template for graduation project"
    assert slots["cover"]["info"] == ["答辩学生：苗猫大丸子", "指导老师：苗猫大丸子"]
    assert slots["section"]["index"] == "01"
    assert slots["section"]["title"] == "研究现状"
    assert slots["section"]["subs"] == ["摘要", "引言"]
    assert slots["content"]["title"] == "请输入你的标题"
    assert slots["content"]["title_right_limit"] == 10.1


def test_content_canvas_clears_safe_area(workdir):
    path = _make_template(workdir / "demo.pptx")
    base = ppt_base.guess_base(path)
    prs = Presentation(str(path))
    src = list(prs.slides)[base["pages"]["content"] - 1]
    slide = ppt_base.content_canvas(prs, src, base, "国家与土地产权关系的变化", 7)

    assert _big_shapes_in_safe(slide, base["safe_area"]) == []
    texts = {s.text_frame.text.strip() for s in slide.shapes if s.has_text_frame}
    assert "国家与土地产权关系的变化" in texts      # 标题换成了我们的
    assert "请输入你的标题" not in texts
    assert "7" in texts                            # 页码改成实际页号
    assert "5" not in texts
    assert "" in texts                             # 日期被清空
    assert "2025/5/23" not in texts


@pytest.mark.parametrize("text, junk", [
    ("请输入你的标题", True),
    ("请在此输入文字", True),
    ("点击添加文字", True),
    ("Add your text here", True),
    ("Type Your Title", True),
    ("Click here", True),
    ("此部分内容", True),
    ("您的内容打在这里", True),
    ("苗猫大丸子", True),
    ("HSVmode 配色", True),
    ("XXX", True),
    ("Thesis defense template for graduation project", True),
    ("2025/5/23", True),
    ("关键词标题", True),
    ("请在这里输入你的正文", True),
    ("5", True),
    ("12", True),
    ("中国人民大学模板", False),
    ("国家与土地产权关系的变化", False),
    ("一、问题意识：所有权悖论", False),
    ("", False),
])
def test_is_junk_text(text, junk):
    assert ppt_base.is_junk_text(text) is junk


def test_save_and_load_base_roundtrip(workdir, monkeypatch):
    folder = workdir / "workspaces" / "u1" / "ppt_templates"
    folder.mkdir(parents=True)
    _make_template(folder / "abc123.pptx")
    monkeypatch.setattr(ppt_base, "DATA_DIR", workdir)
    monkeypatch.setattr(ppt_base, "BUILTIN_DIR", workdir / "builtin")

    assert ppt_base.load_base("abc123") is None
    base = ppt_base.guess_base(folder / "abc123.pptx")
    ppt_base.save_base("abc123", base)
    loaded = ppt_base.load_base("abc123")
    assert loaded is not None
    assert loaded["confirmed"] is True          # 存盘即确认
    assert loaded["pages"] == base["pages"]
    assert (folder / "abc123.base.json").is_file()
    assert ppt_base.load_base("没有这个模板") is None


def test_builtin_base_uses_builtin_folder(workdir, monkeypatch):
    builtin = workdir / "builtin"
    builtin.mkdir()
    _make_template(builtin / "academic-red.pptx")
    monkeypatch.setattr(ppt_base, "BUILTIN_DIR", builtin)
    monkeypatch.setattr(ppt_base, "DATA_DIR", workdir / "nowhere")

    base = ppt_base.guess_base(builtin / "academic-red.pptx")
    ppt_base.save_base("builtin-academic-red", base)
    assert (builtin / "academic-red.base.json").is_file()
    assert ppt_base.load_base("builtin-academic-red")["confirmed"] is True


def test_base_for_online_and_eval_use_the_same_page_notes(workdir):
    path = workdir / "tpl.pptx"
    _make_template(path)
    notes = {"deck": {"suggested_base": {"cover": 1, "section": 2, "content": 1, "end": 4}}}
    template = {"id": "not-saved", "file_path": str(path)}
    online = ppt_base.base_for(template, notes)
    evaluated = ppt_base.base_for(template, notes)
    assert online["pages"] == evaluated["pages"] == {
        "cover": 1, "section": 2, "content": 1, "end": 4,
    }
    assert ppt_base.guess_base(path)["pages"]["content"] != 1


def test_base_for_reads_page_notes_when_omitted(workdir, monkeypatch):
    path = workdir / "tpl.pptx"
    _make_template(path)
    notes = {"deck": {"suggested_base": {"cover": 1, "section": 2, "content": 1, "end": 4}}}
    monkeypatch.setattr("services.ppt.ppt_pages.load_pages", lambda template: notes)
    got = ppt_base.base_for({"id": "not-saved", "file_path": str(path)})
    assert got["pages"] == {"cover": 1, "section": 2, "content": 1, "end": 4}


def test_base_for_keeps_a_saved_base(workdir, monkeypatch):
    folder = workdir / "workspaces" / "u1" / "ppt_templates"
    folder.mkdir(parents=True)
    path = folder / "abc999.pptx"
    _make_template(path)
    monkeypatch.setattr(ppt_base, "DATA_DIR", workdir)
    monkeypatch.setattr(ppt_base, "BUILTIN_DIR", workdir / "builtin")
    saved = ppt_base.guess_base(path)
    ppt_base.save_base("abc999", saved)
    notes = {"deck": {"suggested_base": {"cover": 1, "section": 1, "content": 1, "end": 1}}}
    got = ppt_base.base_for({"id": "abc999", "file_path": str(path)}, notes)
    assert got["confirmed"] is True
    assert got["pages"] == saved["pages"]
    assert got["pages"]["content"] != 1


# 真实模板：本机没有就跳过
REAL_ID = "f52834f22b5a"
REAL_50 = pathlib.Path(__file__).resolve().parents[1] / "data" / "workspaces" / "c46ff668-f8c3-46d9-ba62-7f2c271816d4" / "ppt_templates" / "7acd6eb053b2.pptx"


@pytest.mark.skipif(ppt_base._template_pptx(REAL_ID) is None,
                    reason="本机没有人大模板 53")
def test_real_template_pages_and_canvas():
    path = ppt_base._template_pptx(REAL_ID)
    base = ppt_base.guess_base(path)
    assert base["pages"]["cover"] == 1
    assert base["pages"]["end"] == 31
    assert base["pages"]["section"] in (3, 9, 14, 19, 24)
    assert 1 < base["pages"]["content"] < 31

    prs = Presentation(str(path))
    src = list(prs.slides)[base["pages"]["content"] - 1]
    slide = ppt_base.content_canvas(prs, src, base, "测试标题", 5)
    assert _big_shapes_in_safe(slide, base["safe_area"]) == []


@pytest.mark.skipif(not REAL_50.is_file(), reason="本机没有人大模板 50")
def test_template_50_guess_skips_toc_and_picks_thanks():
    base = ppt_base.guess_base(REAL_50)
    assert base["pages"]["section"] != 2
    assert base["pages"]["end"] == 29
    prs = Presentation(str(REAL_50))
    slide = list(prs.slides)[base["pages"]["content"] - 1]
    _, y0, _, y1 = base["safe_area"]
    count = 0
    for shape in slide.shapes:
        if not shape.has_text_frame or not shape.text_frame.text.strip():
            continue
        center = ppt_base._inches(shape.top) + ppt_base._inches(shape.height) / 2
        if y0 - 0.1 < center < y1 + 0.15:
            count += 1
    assert count <= 2
