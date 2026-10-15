"""整份归纳：换色分组、底板四页、模型失败时的简版讲法。不连网。"""

import pathlib

import pytest
from pptx import Presentation

from services.ppt import ppt_base, ppt_deck_parse, ppt_pages


@pytest.fixture(autouse=True)
def reset_config(monkeypatch):
    sandbox = pathlib.Path(__file__).resolve().parents[3] / ".tmp" / "pytest-env"
    monkeypatch.setenv("ARS_WEB_DB_PATH", str(sandbox / "test_ars.db"))
    monkeypatch.setenv("ARS_SKILLS_PATH", str(sandbox / "fake_skills"))
    monkeypatch.setenv("ARS_WORKSPACE_DIR", str(sandbox / "workspaces"))


def _sample():
    frames = [{"i": 0, "x": 1.0, "y": 0.4, "w": 6.0, "h": 0.8}]
    same = [{"i": 0, "x": 1.2, "y": 0.45, "w": 6.1, "h": 0.9}]
    other = [{"i": 0, "x": 1.0, "y": 3.2, "w": 6.0, "h": 0.8},
             {"i": 1, "x": 1.0, "y": 4.4, "w": 6.0, "h": 1.2}]
    analysis = {"specimens": [
        {"index": 0, "frames": frames},
        {"index": 1, "frames": same},
        {"index": 2, "frames": other},
        {"index": 3, "frames": [{"i": 0, "x": 1.0, "y": 1.0, "w": 4.0, "h": 0.6}]},
        {"index": 4, "frames": [{"i": 0, "x": 1.0, "y": 1.0, "w": 4.0, "h": 0.6}]},
    ]}
    pages = {"pages": [
        {"page": 1, "role": "cover", "slots": [{"name": "title", "frame": 0}], "decor": [],
         "sample_heading": "开题答辩", "narrative_fit": ["hook"]},
        {"page": 2, "role": "content", "slots": [{"name": "title", "frame": 0}], "decor": [],
         "sample_heading": "关键词标题", "narrative_fit": ["background"]},
        {"page": 3, "role": "content", "slots": [
            {"name": "title", "frame": 0}, {"name": "body", "frame": 1},
        ], "decor": [{"action": "clear"}], "sample_heading": "研究发现", "narrative_fit": ["finding"]},
        {"page": 4, "role": "section", "slots": [{"name": "section_title", "frame": 0}], "decor": [],
         "sample_heading": "研究背景", "narrative_fit": ["background"]},
        {"page": 5, "role": "end", "slots": [{"name": "title", "frame": 0}], "decor": [],
         "sample_heading": "谢谢", "narrative_fit": ["takeaway"]},
    ]}
    return pages, analysis


def test_same_signature_groups_together_and_base_stays_in_range():
    pages, analysis = _sample()
    facts = ppt_deck_parse.deck_facts(pages, analysis)
    grouped = [item for item in facts["variant_groups"] if set(item["pages"]) == {1, 2}]
    assert grouped
    assert all(item["pages"] != [1, 3] for item in facts["variant_groups"])
    assert pages["pages"][0]["variant_group"] == pages["pages"][1]["variant_group"]
    assert pages["pages"][0]["variant_group"] != pages["pages"][2]["variant_group"]
    base = facts["suggested_base"]
    assert set(base) == {"cover", "section", "content", "end"}
    assert all(1 <= number <= 5 for number in base.values())
    assert base["content"] == 2


def test_user_summary_is_kept():
    pages, analysis = _sample()
    pages["deck"] = {"summary": "用户写的简介", "user_summary": True}
    facts = ppt_deck_parse.deck_facts(pages, analysis)
    ppt_pages.write_deck(pages, facts, {"summary": "模型写的", "deck_type_hint": "", "template_storyline": []})
    assert pages["deck"]["summary"] == "用户写的简介"


@pytest.mark.asyncio
async def test_illegal_story_json_uses_headings(monkeypatch):
    async def fake(*args, **kwargs):
        return "不是 JSON", "stub", ""

    monkeypatch.setattr("services.ppt.ppt_plan._ask_text", fake)
    pages, _analysis = _sample()
    story = await ppt_deck_parse.deck_story(pages, "stub-model", "user")
    assert story["summary"] == ""
    assert story["deck_type_hint"] == ""
    headings = [item["heading"] for item in story["template_storyline"]]
    assert "开题答辩" in headings
    assert "研究背景" in headings
    assert "关键词标题" not in headings


def test_guess_base_uses_suggested_pages():
    folder = pathlib.Path(__file__).resolve().parents[3] / ".tmp" / "ppt_deck_parse_test"
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / "four.pptx"
    prs = Presentation()
    for _ in range(4):
        prs.slides.add_slide(prs.slide_layouts[0])
    prs.save(str(path))
    doc = {"deck": {"suggested_base": {"cover": 1, "section": 2, "content": 3, "end": 4}}}
    base = ppt_base.guess_base(path, doc)
    assert base["pages"] == {"cover": 1, "section": 2, "content": 3, "end": 4}
    assert base["confirmed"] is False
