"""版式库：只画在安全区里、文本框互不重叠、超上限抛 Overflow。"""

import copy
import pathlib

import pytest
from pptx import Presentation
from pptx.util import Emu, Inches

from services.ppt import ppt_layouts

BASE = {
    "slide_size": [13.33, 7.5],
    "safe_area": [0.74, 1.45, 12.6, 6.75],
    "colors": {"accent": "#9B1B30", "accent_soft": "#F6ECEE", "dark": "#262626",
               "gray": "#595959", "muted": "#8C8C8C"},
}


@pytest.fixture(autouse=True)
def reset_config(monkeypatch):
    """覆盖 conftest 的同名 fixture。

    DSH 沙箱里 pytest 用 0o700 建的 tmp_path 受限令牌进不去；本模块不碰数据库，
    环境变量指到普通目录拼路径即可。
    """
    sandbox = pathlib.Path(__file__).resolve().parents[3] / ".tmp" / "pytest-env"
    monkeypatch.setenv("ARS_WEB_DB_PATH", str(sandbox / "test_ars.db"))
    monkeypatch.setenv("ARS_SKILLS_PATH", str(sandbox / "fake_skills"))
    monkeypatch.setenv("ARS_WORKSPACE_DIR", str(sandbox / "workspaces"))


def _blank():
    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)
    return prs, prs.slides.add_slide(prs.slide_layouts[6])


def _cjk(n: int) -> str:
    return "字" * n


# 每种版式都用"上限内、条数取最大"的数据
# 注意：数组元素一律用推导式新建，用 [x] * n 会让几个元素共用同一个 dict
FULL: dict[str, dict] = {
    "toc": {"layout": "toc", "title": _cjk(20),
            "items": [{"label": _cjk(6), "desc": _cjk(12)} for _ in range(6)]},
    "cards": {"layout": "cards", "title": _cjk(20),
              "cards": [{"key": _cjk(6), "desc": _cjk(28)} for _ in range(4)],
              "takeaway": _cjk(34)},
    "compare": {"layout": "compare", "title": _cjk(20),
                "left": {"head": _cjk(8), "points": [_cjk(18)] * 4},
                "right": {"head": _cjk(8), "points": [_cjk(18)] * 4}},
    "timeline": {"layout": "timeline", "title": _cjk(20),
                 "nodes": [{"year": _cjk(6), "head": _cjk(6), "desc": _cjk(10)}
                           for _ in range(6)]},
    "quote": {"layout": "quote", "title": _cjk(20), "quote": _cjk(40),
              "source": _cjk(20), "note": _cjk(40)},
    "numbered": {"layout": "numbered", "title": _cjk(20),
                 "rows": [{"head": _cjk(9), "desc": _cjk(30)} for _ in range(4)]},
    "bullets": {"layout": "bullets", "title": _cjk(20), "points": [_cjk(30)] * 5},
    "bignum": {"layout": "bignum", "title": _cjk(20),
               "stats": [{"value": "12345", "label": _cjk(10), "desc": _cjk(24)}
                         for _ in range(3)]},
    "image_text": {"layout": "image_text", "title": _cjk(20), "image": 1,
                   "points": [_cjk(24)] * 4},
}


def _boxes(slide) -> list[tuple[float, float, float, float]]:
    """页面上有文字的框（文本框和带字的形状），返回英寸坐标。"""
    out = []
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text_frame.text.strip():
            out.append((Emu(shape.left).inches, Emu(shape.top).inches,
                        Emu(shape.width).inches, Emu(shape.height).inches))
    return out


def _overlap(a, b) -> bool:
    ax, ay, aw, ah = a
    bx, by, bw, bh = b
    gap = 0.01      # 边框相切不算重叠
    return ax < bx + bw - gap and bx < ax + aw - gap and ay < by + bh - gap and by < ay + ah - gap


@pytest.mark.parametrize("layout", list(FULL))
def test_draws_inside_safe_area(layout):
    prs, slide = _blank()
    ppt_layouts.render(slide, BASE, FULL[layout])
    x0, y0, x1, y1 = BASE["safe_area"]
    boxes = _boxes(slide)
    assert boxes, f"{layout} 什么都没画出来"
    for left, top, w, h in boxes:
        assert left >= x0 - 0.02, f"{layout} 左边越界：{left:.2f} < {x0}"
        assert top >= y0 - 0.02, f"{layout} 上边越界：{top:.2f} < {y0}"
        assert left + w <= x1 + 0.02, f"{layout} 右边越界：{left + w:.2f} > {x1}"
        assert top + h <= y1 + 0.02, f"{layout} 下边越界：{top + h:.2f} > {y1}"


@pytest.mark.parametrize("layout", list(FULL))
def test_text_boxes_do_not_overlap(layout):
    prs, slide = _blank()
    ppt_layouts.render(slide, BASE, FULL[layout])
    boxes = _boxes(slide)
    for i in range(len(boxes)):
        for j in range(i + 1, len(boxes)):
            assert not _overlap(boxes[i], boxes[j]), \
                f"{layout} 第 {i} 个和 第 {j} 个文本框重叠：{boxes[i]} / {boxes[j]}"


@pytest.mark.parametrize("layout, patch, field, limit", [
    ("toc", {"items": [{"label": _cjk(9), "desc": _cjk(12)} for _ in range(3)]},
     "items[0].label", 6),
    ("cards", {"cards": [{"key": _cjk(6), "desc": _cjk(42)} for _ in range(2)]},
     "cards[0].desc", 28),
    ("compare", {"left": {"head": _cjk(8), "points": [_cjk(27)] * 2}}, "left.points[0]", 18),
    ("timeline", {"nodes": [{"year": _cjk(6), "head": _cjk(6), "desc": _cjk(15)}
                            for _ in range(3)]}, "nodes[0].desc", 10),
    ("quote", {"quote": _cjk(60)}, "quote", 40),
    ("numbered", {"rows": [{"head": _cjk(9), "desc": _cjk(45)} for _ in range(2)]},
     "rows[0].desc", 30),
    ("bullets", {"points": [_cjk(45)] * 3}, "points[0]", 30),
    ("bignum", {"stats": [{"value": "1234567", "label": _cjk(10), "desc": _cjk(24)}]},
     "stats[0].value", 6),
    ("image_text", {"points": [_cjk(36)] * 2}, "points[0]", 24),
])
def test_over_limit_raises_overflow(layout, patch, field, limit):
    """超上限 1.5 倍的字必须抛 Overflow，并报出是哪个字段。"""
    plan = copy.deepcopy(FULL[layout])
    plan.update(patch)
    prs, slide = _blank()
    with pytest.raises(ppt_layouts.Overflow) as err:
        ppt_layouts.render(slide, BASE, plan)
    assert err.value.field == field
    assert err.value.limit == limit


def test_too_many_items_raises_overflow():
    plan = copy.deepcopy(FULL["bullets"])
    plan["points"] = [_cjk(10)] * 6      # 上限 5 条
    prs, slide = _blank()
    with pytest.raises(ppt_layouts.Overflow) as err:
        ppt_layouts.render(slide, BASE, plan)
    assert err.value.field == "points"


def test_limits_scale_with_narrow_template():
    narrow = dict(BASE, slide_size=[10.0, 7.5])
    limits = ppt_layouts.limits_of("cards", narrow)
    # 28 × 10/13.33 = 21.0，向下取整
    assert limits["cards[].desc"] == 21
    assert limits["cards[].key"] == 4
    assert ppt_layouts.limits_of("cards", BASE)["cards[].desc"] == 28


def test_plain_layouts_draw_nothing():
    """封面、章节页、结尾走底板换字，版式库不画东西。"""
    prs, slide = _blank()
    for layout in ppt_layouts.PLAIN_LAYOUTS:
        assert ppt_layouts.render(slide, BASE, {"layout": layout}) == {}
    assert len(slide.shapes) == 0


def test_text_len_skips_punctuation():
    """计划里给的示例标题连着标点是 22 个字符，按"字"算是 20，正好在上限内。"""
    assert ppt_layouts.text_len("所有权悖论：产权离不开国家，也容易被国家侵蚀") == 20
    assert ppt_layouts.text_len("国家必要") == 4


def test_em_counts_wide_symbols_as_full_width():
    """箭头和百分号在中文里也是全角，按半宽估会算出偏大的字号。"""
    assert ppt_layouts._em("→") == 1.0
    assert ppt_layouts._em("%") == 1.0
    assert ppt_layouts._em("a") == 0.55
    # 4 个数字按半宽 + % 和 → 按全角
    assert ppt_layouts._em("91%→19") == 4.2


def test_truncate_keeps_within_limit():
    long_text = "产权是一种社会工具，其重要性就在于帮助人们形成合理预期。"
    cut = ppt_layouts.truncate(long_text, 10)
    assert ppt_layouts.text_len(cut) <= 10
    assert ppt_layouts.truncate("短句", 10) == "短句"


def test_check_slide_reports_every_issue():
    plan = copy.deepcopy(FULL["cards"])
    plan["cards"][0]["desc"] = _cjk(42)
    plan["takeaway"] = _cjk(50)
    issues = ppt_layouts.check_slide(plan, BASE)
    fields = {item["field"] for item in issues}
    assert fields == {"cards[0].desc", "takeaway"}
