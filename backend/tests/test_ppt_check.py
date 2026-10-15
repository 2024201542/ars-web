"""出片后自检：几何、文字、缩略图、可选的视觉复核。"""

import pathlib
import shutil

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.util import Inches, Pt

from services.ppt import ppt_check, pptx_render

BASE = {
    "slide_size": [13.33, 7.5],
    "safe_area": [0.7, 1.45, 12.63, 6.75],
    "colors": {"accent": "#9B1B30", "accent_soft": "#F5E8EA", "dark": "#262626",
               "gray": "#595959", "muted": "#8C8C8C"},
}


@pytest.fixture(autouse=True)
def reset_config(monkeypatch):
    """覆盖 conftest 的同名 fixture：不用 pytest 的 tmp_path。"""
    sandbox = pathlib.Path(__file__).resolve().parents[3] / ".tmp" / "pytest-env"
    monkeypatch.setenv("ARS_WEB_DB_PATH", str(sandbox / "test_ars.db"))
    monkeypatch.setenv("ARS_SKILLS_PATH", str(sandbox / "fake_skills"))
    monkeypatch.setenv("ARS_WORKSPACE_DIR", str(sandbox / "workspaces"))


@pytest.fixture()
def workdir():
    folder = pathlib.Path(__file__).resolve().parents[3] / ".tmp" / "ppt_check_test"
    shutil.rmtree(folder, ignore_errors=True)
    folder.mkdir(parents=True, exist_ok=True)
    return folder


def _textbox(slide, x, y, w, h, text="", pt=18):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if text:
        run = box.text_frame.paragraphs[0].add_run()
        run.text = text
        run.font.size = Pt(pt)
    return box


def _make_deck(path: pathlib.Path) -> pathlib.Path:
    """每页只制造一个问题，方便逐条断言。"""
    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    clean = prs.slides.add_slide(blank)                  # 1 干净页
    _textbox(clean, 0.7, 0.3, 8.0, 0.6, "干净页的标题")
    _textbox(clean, 1.0, 2.0, 5.0, 1.0, "正常的一段正文。")
    _textbox(clean, 1.0, 4.0, 5.0, 1.0, "另一段正文。")

    off = prs.slides.add_slide(blank)                    # 2 出页面
    _textbox(off, 0.7, 0.3, 8.0, 0.6, "越界页的标题")
    _textbox(off, -0.8, 2.0, 3.0, 0.8, "这个框跑到页面左边外面了")

    over = prs.slides.add_slide(blank)                   # 3 框里放不下
    _textbox(over, 0.7, 0.3, 8.0, 0.6, "放不下页的标题")
    _textbox(over, 1.0, 2.0, 1.0, 0.3, "这是一段很长的正文，硬塞进一个很小的框里，肯定放不下。", 40)

    junk = prs.slides.add_slide(blank)                   # 4 残留示例字
    _textbox(junk, 0.7, 0.3, 8.0, 0.6, "残留页的标题")
    _textbox(junk, 1.0, 2.0, 5.0, 0.6, "请输入你的标题")

    repeat = prs.slides.add_slide(blank)                 # 5 标题与上一页相同
    _textbox(repeat, 0.7, 0.3, 8.0, 0.6, "残留页的标题")
    _textbox(repeat, 1.0, 2.0, 5.0, 1.0, "正常正文。")

    lap = prs.slides.add_slide(blank)                    # 6 两个框重叠
    _textbox(lap, 0.7, 0.3, 8.0, 0.6, "重叠页的标题")
    _textbox(lap, 1.0, 2.0, 4.0, 1.0, "第一个框")
    _textbox(lap, 2.0, 2.5, 4.0, 1.0, "第二个框压在第一个上面")

    empty = prs.slides.add_slide(prs.slide_layouts[5])   # 7 页眉有占位符却没写标题
    _textbox(empty, 1.0, 2.0, 5.0, 1.0, "正文一。")
    _textbox(empty, 1.0, 4.0, 5.0, 1.0, "正文二。")

    prs.save(str(path))
    return path


def _kinds(issues, page=None):
    return {i["kind"] for i in issues if page is None or i["page"] == page}


def test_check_deck_reports_each_kind(workdir):
    path = _make_deck(workdir / "deck.pptx")
    issues = ppt_check.check_deck(path, BASE)
    assert _kinds(issues, 1) == set()                     # 干净页不报
    assert "offslide" in _kinds(issues, 2)
    assert "overflow" in _kinds(issues, 3)
    assert "junk" in _kinds(issues, 4)
    assert "repeat_title" in _kinds(issues, 5)
    assert "overlap" in _kinds(issues, 6)
    assert "empty_title" in _kinds(issues, 7)


def test_check_deck_issue_shape(workdir):
    path = _make_deck(workdir / "deck.pptx")
    issues = ppt_check.check_deck(path, BASE)
    assert issues
    for issue in issues:
        assert set(issue) == {"page", "kind", "detail", "hint"}
        assert issue["hint"]
        assert isinstance(issue["page"], int) and issue["page"] >= 1
        assert issue["detail"]


def test_check_deck_flags_offslide(workdir):
    """框探出页面边界要报出来（安全区不判：模板自己的封面框就跨越安全区）。"""
    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _textbox(slide, 0.7, 0.3, 8.0, 0.6, "标题")
    _textbox(slide, 12.9, 2.0, 1.0, 1.0, "这个框右边探出页面了")   # 12.9 + 1.0 > 13.33
    path = workdir / "safe.pptx"
    prs.save(str(path))
    issues = ppt_check.check_deck(path, BASE)
    assert "offslide" in _kinds(issues, 1)


@pytest.mark.skipif(not pptx_render.available(), reason="这台机器上没有 LibreOffice Kit")
def test_render_thumbs_returns_every_page(workdir):
    path = _make_deck(workdir / "deck.pptx")
    thumbs = ppt_check.render_thumbs(path, dpi=72)
    assert sorted(thumbs) == list(range(1, 8))
    for png in thumbs.values():
        assert png.is_file() and png.stat().st_size > 0


# ───────────── 视觉复核 ─────────────

def test_vision_disabled_by_default(monkeypatch):
    monkeypatch.delenv("ARS_PPT_VLM_CHECK", raising=False)
    assert ppt_check.vision_enabled() is False
    monkeypatch.setenv("ARS_PPT_VLM_CHECK", "1")
    assert ppt_check.vision_enabled() is True
    monkeypatch.setenv("ARS_PPT_VLM_CHECK", "0")
    assert ppt_check.vision_enabled() is False


@pytest.mark.asyncio
async def test_check_visual_skips_when_disabled(monkeypatch, workdir):
    path = _make_deck(workdir / "deck.pptx")
    monkeypatch.delenv("ARS_PPT_VLM_CHECK", raising=False)
    called = {"n": 0}

    async def fake_opinion(user_id, image_path):
        called["n"] += 1
        return "无", ""

    monkeypatch.setattr(ppt_check, "_vision_opinion", fake_opinion)
    assert await ppt_check.check_visual(path, BASE, user_id="u1") == []
    assert called["n"] == 0          # 开关关着，一次都不该调


@pytest.mark.asyncio
async def test_check_visual_reports_at_most_three_pages(monkeypatch, workdir):
    path = _make_deck(workdir / "deck.pptx")
    monkeypatch.setenv("ARS_PPT_VLM_CHECK", "1")
    seen: list[int] = []

    def fake_render(pptx_path, pages, dpi=110):
        seen.extend(pages)
        return {p: workdir / f"fake-{p}.png" for p in pages}

    async def fake_opinion(user_id, image_path):
        return "文字压在一起了", ""

    monkeypatch.setattr(pptx_render, "render_pages", fake_render)
    monkeypatch.setattr(ppt_check, "_vision_opinion", fake_opinion)
    issues = await ppt_check.check_visual(path, BASE, user_id="u1")
    assert 0 < len(seen) <= ppt_check.VISION_MAX_PAGES
    assert len(issues) == len(seen)
    assert all(i["kind"] == "visual" for i in issues)


@pytest.mark.asyncio
async def test_check_visual_ignores_no_problem(monkeypatch, workdir):
    path = _make_deck(workdir / "deck.pptx")
    monkeypatch.setenv("ARS_PPT_VLM_CHECK", "1")

    monkeypatch.setattr(pptx_render, "render_pages",
                        lambda pptx_path, pages, dpi=110: {p: workdir / f"x-{p}.png" for p in pages})

    async def fake_opinion(user_id, image_path):
        return "无。", ""

    monkeypatch.setattr(ppt_check, "_vision_opinion", fake_opinion)
    assert await ppt_check.check_visual(path, BASE, user_id="u1") == []


def test_repair_pulls_offslide_box_inside(workdir):
    path = workdir / "off.pptx"
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(5.62)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = _textbox(slide, 8.4, 5.1, 2.2, 0.8, "这一行探出页面了")
    box.name = "ars:takeaway"
    prs.save(str(path))

    fixed = ppt_check.repair_deck(path)
    assert fixed["pulled"] >= 1
    issues = ppt_check.check_deck(path)
    assert not any(item["kind"] == "offslide" for item in issues)
    again = Presentation(str(path))
    text = "\n".join(shape.text_frame.text for shape in again.slides[0].shapes if shape.has_text_frame)
    assert "探出页面" in text


def test_repair_drops_template_text_sitting_on_ours(workdir):
    path = workdir / "lap.pptx"
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(5.62)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    ours = _textbox(slide, 1.0, 4.4, 6.0, 0.8, "我们的结论句")
    ours.name = "ars:takeaway"
    leftover = _textbox(slide, 3.2, 4.5, 4.2, 0.5, "Measurement of horizontal angle")
    leftover.name = "小红书-苗猫大丸子"
    prs.save(str(path))

    fixed = ppt_check.repair_deck(path)
    assert fixed["removed"] == 1
    again = Presentation(str(path))
    texts = [shape.text_frame.text for shape in again.slides[0].shapes if shape.has_text_frame]
    assert any("我们的结论句" in text for text in texts)
    assert not any("Measurement" in text for text in texts)


def test_repair_drops_template_watermarks_that_overlap_each_other(workdir):
    path = workdir / "marks.pptx"
    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    kept = _textbox(slide, 1.0, 2.0, 6.0, 0.8, "我们的标题")
    kept.name = "ars:title"
    mark_a = _textbox(slide, 10.6, 6.1, 2.3, 0.4, "MIAOMAO")
    mark_a.name = "小红书-装饰-40"
    mark_b = _textbox(slide, 10.6, 6.4, 2.3, 0.4, "TEMPLATE")
    mark_b.name = "小红书-装饰-42"
    prs.save(str(path))

    fixed = ppt_check.repair_deck(path)
    assert fixed["removed"] == 2
    again = Presentation(str(path))
    texts = [shape.text_frame.text for shape in again.slides[0].shapes if shape.has_text_frame]
    assert any("我们的标题" in text for text in texts)
    assert not any("MIAOMAO" in text or "TEMPLATE" in text for text in texts)
    assert not any(item["kind"] == "overlap" for item in ppt_check.check_deck(path))


def test_repair_keeps_bottom_box_inside_the_page(workdir):
    path = workdir / "bottom.pptx"
    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    footer = _textbox(slide, 9.4, 6.95, 3.0, 0.32, "01")
    footer.name = "ars:footer"
    title = _textbox(slide, 0.45, 7.31, 12.0, 0.28, "这一行原本探出页面")
    title.name = "ars:title"
    prs.save(str(path))

    ppt_check.repair_deck(path)
    issues = ppt_check.check_deck(path)
    assert not any(item["kind"] in ("offslide", "overlap") for item in issues)
    again = Presentation(str(path))
    text = "\n".join(shape.text_frame.text for shape in again.slides[0].shapes if shape.has_text_frame)
    assert "探出页面" in text


def test_repair_leaves_a_clean_page_alone(workdir):
    path = workdir / "clean.pptx"
    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = _textbox(slide, 1.0, 1.5, 6.0, 0.8, "正常正文")
    box.name = "ars:body"
    prs.save(str(path))
    before = ppt_check.check_deck(path)

    fixed = ppt_check.repair_deck(path)
    assert fixed == {"pulled": 0, "removed": 0}
    assert ppt_check.check_deck(path) == before


def test_text_that_runs_into_a_graphic_counts_as_overlap(workdir):
    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _textbox(slide, 0.06, 2.53, 8.57, 1.31, "物理学习记忆机制与干预", 52)
    slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.OVAL, Inches(7.85), Inches(-1.56), Inches(2.5), Inches(10.51))
    path = workdir / "on-arc.pptx"
    prs.save(str(path))
    hits = [item for item in ppt_check.check_deck(path) if "压到图片上" in item["detail"]]
    assert len(hits) == 1
    assert hits[0]["kind"] == "overlap"

    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _textbox(slide, 0.06, 2.53, 8.57, 1.31, "感谢您的观看", 72)
    slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.OVAL, Inches(7.85), Inches(-1.56), Inches(2.5), Inches(10.51))
    short = workdir / "clear.pptx"
    prs.save(str(short))
    assert not any("压到图片上" in item["detail"] for item in ppt_check.check_deck(short))

    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _textbox(slide, 0.06, 2.53, 8.57, 1.31, "一二三四五六七八九十" * 4, 18)
    slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.OVAL, Inches(7.85), Inches(-1.56), Inches(2.5), Inches(10.51))
    body = workdir / "body.pptx"
    prs.save(str(body))
    assert not any("压到图片上" in item["detail"] for item in ppt_check.check_deck(body))
