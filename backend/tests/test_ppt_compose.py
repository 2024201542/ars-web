"""按计划出片：页数、页码、备注、残留示例字清理、超限降级。"""

import pathlib
import shutil

import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

from services.ppt import ppt_base, ppt_compose


@pytest.fixture(autouse=True)
def reset_config(monkeypatch):
    """覆盖 conftest 的同名 fixture：不用 pytest 的 tmp_path。"""
    sandbox = pathlib.Path(__file__).resolve().parents[3] / ".tmp" / "pytest-env"
    monkeypatch.setenv("ARS_WEB_DB_PATH", str(sandbox / "test_ars.db"))
    monkeypatch.setenv("ARS_SKILLS_PATH", str(sandbox / "fake_skills"))
    monkeypatch.setenv("ARS_WORKSPACE_DIR", str(sandbox / "workspaces"))


@pytest.fixture()
def workdir():
    folder = pathlib.Path(__file__).resolve().parents[3] / ".tmp" / "ppt_compose_test"
    shutil.rmtree(folder, ignore_errors=True)
    folder.mkdir(parents=True, exist_ok=True)
    return folder


def _textbox(slide, x, y, w, h, text, pt):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = text
    run.font.size = Pt(pt)
    return box


def _rect(slide, x, y, w, h, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y),
                                   Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(color)
    shape.line.fill.background()
    return shape


def _make_template(path: pathlib.Path) -> pathlib.Path:
    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    accent = "9B1B30"

    cover = prs.slides.add_slide(blank)
    _rect(cover, 0.2, 1.0, 0.5, 2.0, accent)
    _textbox(cover, 1.6, 2.4, 8.0, 0.8, "中国人民大学模板", 32)
    _textbox(cover, 1.6, 3.6, 8.0, 0.5,
             "Thesis defense template for graduation project", 14)
    _textbox(cover, 1.6, 4.6, 6.0, 0.4, "答辩学生：苗猫大丸子", 14)
    _textbox(cover, 1.6, 5.0, 6.0, 0.4, "指导老师：苗猫大丸子", 14)

    section = prs.slides.add_slide(blank)
    _textbox(section, 1.0, 2.4, 1.4, 1.0, "01", 44)
    _textbox(section, 2.4, 2.6, 5.0, 0.6, "研究现状", 24)
    _textbox(section, 1.2, 3.8, 4.0, 0.4, "摘要", 12)
    _textbox(section, 1.2, 4.2, 4.0, 0.4, "引言", 12)

    content = prs.slides.add_slide(blank)
    _rect(content, 11.5, 0.3, 1.2, 0.5, accent)
    _textbox(content, 0.74, 0.3, 6.0, 0.6, "请输入你的标题", 20)
    _rect(content, 0.74, 1.13, 12.0, 0.03, accent)
    _textbox(content, 1.5, 2.0, 8.0, 1.5, "点击添加正文内容。", 16)
    _rect(content, 3.0, 3.5, 2.0, 1.2, accent)
    _textbox(content, 0.74, 6.95, 2.0, 0.3, "2025/5/23", 10)
    _textbox(content, 9.5, 6.95, 1.0, 0.3, "5", 10)

    end = prs.slides.add_slide(blank)
    _textbox(end, 1.6, 2.5, 6.0, 0.8, "谢谢", 30)
    _textbox(end, 1.6, 4.6, 6.0, 0.4, "答辩学生：苗猫大丸子", 14)
    _textbox(end, 1.6, 5.0, 6.0, 0.4, "指导老师：苗猫大丸子", 14)
    prs.save(str(path))
    return path


def _all_texts(prs) -> list[str]:
    out = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                out.append(shape.text_frame.text.strip())
    return out


def _footer_digits(prs, base) -> list[list[str]]:
    limit = base["slide_size"][1] - ppt_base.BOTTOM_BAND
    pages = []
    for slide in prs.slides:
        found = []
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            text = shape.text_frame.text.strip()
            if not text or not text.isdigit():
                continue
            center = ppt_base._inches(shape.top) + ppt_base._inches(shape.height) / 2
            if center > limit:
                found.append(text)
        pages.append(found)
    return pages


def _small_plan() -> dict:
    return {"title": "测试", "slides": [
        {"layout": "cover", "title": "国家与土地产权关系的变化",
         "subtitle": "中国农村改革的经济制度变迁史回顾",
         "info": ["汇报人：stu01", "课程：中国经济史专题"], "notes": "开场：交代主题。"},
        {"layout": "bullets", "title": "材料与方法上的说明",
         "points": ["研究对象是 1950 至 1982 年间的制度变迁。",
                    "材料以中央文件和地方志为主。",
                    "统计口径统一按当年产量折算。"], "notes": "讲三条来源。"},
        {"layout": "cards", "title": "三个并列要点",
         "cards": [{"key": "甲", "desc": "第一条说明。"},
                   {"key": "乙", "desc": "第二条说明。"}],
         "takeaway": "一句小结。", "notes": "两张卡片。"},
        {"layout": "end", "info": ["汇报人：stu01", "欢迎批评指正"], "notes": "结束。"},
    ]}


def _compose(workdir, plan):
    tpl = _make_template(workdir / "demo.pptx")
    base = ppt_base.guess_base(tpl)
    out = workdir / "out.pptx"
    result = ppt_compose.compose(base, plan, tpl, out)
    return base, out, result


def test_pages_and_notes(workdir):
    base, out, result = _compose(workdir, _small_plan())
    assert result["pages"] == 4
    assert result["overflow"] == []
    assert result["warnings"] == []
    prs = Presentation(str(out))
    assert len(prs.slides) == 4
    notes = [s.notes_slide.notes_text_frame.text for s in prs.slides]
    assert notes[0] == "开场：交代主题。"
    assert notes[3] == "结束。"
    texts = _all_texts(prs)
    assert "国家与土地产权关系的变化" in texts
    assert "汇报人：stu01" in texts


def test_app_xml_slide_count_matches_the_deck(workdir):
    """模板自带的页数不改的话，PowerPoint 会报内容有问题。"""
    import zipfile
    import xml.etree.ElementTree as ET

    _, out, result = _compose(workdir, _small_plan())
    ep = "http://schemas.openxmlformats.org/officeDocument/2006/extended-properties"
    vt = "http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes"
    with zipfile.ZipFile(out) as zf:
        root = ET.fromstring(zf.read("docProps/app.xml"))
        deck = len(ET.fromstring(zf.read("ppt/presentation.xml")).findall(
            "{http://schemas.openxmlformats.org/presentationml/2006/main}sldIdLst/"
            "{http://schemas.openxmlformats.org/presentationml/2006/main}sldId"))
    assert int(root.find(f"{{{ep}}}Slides").text) == result["pages"] == deck == 4
    assert int(root.find(f"{{{ep}}}Notes").text) == 4
    variants = root.find(f"{{{ep}}}HeadingPairs").find(f"{{{vt}}}vector").findall(f"{{{vt}}}variant")
    title_count = None
    for index in range(0, len(variants) - 1, 2):
        label = variants[index].find(f"{{{vt}}}lpstr")
        if label is not None and label.text == "Slide Titles":
            title_count = int(variants[index + 1].find(f"{{{vt}}}i4").text)
    titles = root.find(f"{{{ep}}}TitlesOfParts").find(f"{{{vt}}}vector")
    assert title_count == 4
    assert int(titles.get("size")) == len(list(titles))


def test_page_numbers_follow_order(workdir):
    base, out, _ = _compose(workdir, _small_plan())
    prs = Presentation(str(out))
    pages = _footer_digits(prs, base)
    assert pages[0] == []          # 封面不显示页码
    assert pages[1] == ["2"]       # 内容页页码＝实际页号
    assert pages[2] == ["3"]
    assert pages[3] == []          # 结尾不显示页码


def test_our_numbers_survive_five_runs(workdir):
    """自己画的序号不能被当成残留字清掉，而且连出 5 次结果一样。"""
    plan = _small_plan()
    plan["slides"].insert(1, {
        "layout": "numbered", "title": "三条建议",
        "rows": [
            {"head": "夯实底座", "desc": "统一口径"},
            {"head": "小步验证", "desc": "从高频报表切入"},
            {"head": "人工复核", "desc": "关键结论留人看"},
        ],
    })
    plan["slides"].insert(1, {"layout": "section", "index": "01", "title": "建议", "subs": ["落地"]})
    texts = []
    for _ in range(5):
        _, out, _ = _compose(workdir, plan)
        prs = Presentation(str(out))
        blob = "\n".join(_all_texts(prs))
        texts.append(blob)
        assert "01" in blob
        assert "02" in blob
        assert "03" in blob
    assert len(set(texts)) == 1


def test_no_junk_text_left(workdir):
    """模板残留的示例字一个都不能剩。

    用 JUNK_PATTERN 而不是 is_junk_text：后者连"只含一两个数字"的框也算残留，
    而我们自己在卡片/编号版式上画的「01」正是这种，属于正常内容。
    """
    base, out, _ = _compose(workdir, _small_plan())
    prs = Presentation(str(out))
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            text = shape.text_frame.text.strip()
            if not text:
                continue
            assert not ppt_base.JUNK_PATTERN.search(text), f"残留示例字：{text!r}"


def test_overflow_is_trimmed_not_fatal(workdir):
    plan = _small_plan()
    plan["slides"][2]["cards"][0]["desc"] = "甲" * 42      # 上限 28 的 1.5 倍
    base, out, result = _compose(workdir, plan)
    assert result["pages"] == 4
    assert len(result["overflow"]) == 1
    assert result["overflow"][0]["field"] == "cards[0].desc"
    assert result["overflow"][0]["limit"] == 28
    texts = _all_texts(Presentation(str(out)))
    assert "三个并列要点" in texts          # 这一页仍然画出来了


def test_unknown_layout_falls_back(workdir):
    plan = _small_plan()
    plan["slides"][1]["layout"] = "sparkle"
    base, out, result = _compose(workdir, plan)
    assert result["pages"] == 4
    assert result["warnings"] and "sparkle" in result["warnings"][0]["detail"]
    assert "材料与方法上的说明" in _all_texts(Presentation(str(out)))


def test_image_text_without_images_warns(workdir):
    plan = _small_plan()
    plan["slides"][1] = {"layout": "image_text", "title": "配图页",
                         "points": ["第一点。", "第二点。"], "image": 1, "notes": ""}
    base, out, result = _compose(workdir, plan)
    assert result["pages"] == 4
    assert any("配图" in w["detail"] for w in result["warnings"])
    # 占位框留着，阶段 4 之前不会凭空多一张图
    prs = Presentation(str(out))
    names = [s.name for s in list(prs.slides)[1].shapes]
    assert ppt_layouts_slot_name() in names


def ppt_layouts_slot_name() -> str:
    from services.ppt import ppt_layouts
    return ppt_layouts.PIC_SLOT_NAME


# ───────────── 真实模板 53 + demo 的 9 页数据 ─────────────

REAL_ID = "f52834f22b5a"

DEMO_PLAN = {"title": "国家与土地产权关系的变化", "slides": [
    {"layout": "cover", "title": "国家与土地产权关系的变化",
     "subtitle": "中国农村改革的经济制度变迁史回顾",
     "info": ["汇报人：stu01", "课程：中国经济史专题"],
     "notes": "开场：用制度变迁的视角回顾改革。"},
    {"layout": "toc", "title": "汇报框架", "notes": "五个部分。", "items": [
        {"label": "理论框架", "desc": "国家与产权的悖论"},
        {"label": "改革前产权", "desc": "集体化的由来与特征"},
        {"label": "公社内变迁", "desc": "收益、费用与政策退却"},
        {"label": "改革与产权", "desc": "包产到户与乡镇企业"},
        {"label": "结论评论", "desc": "国家如何保护有效产权"},
    ]},
    {"layout": "section", "index": "01", "title": "理论框架",
     "subs": ["所有权悖论", "诺斯难题与新假设"], "notes": "进入第一节。"},
    {"layout": "cards", "title": "所有权悖论：产权离不开国家，也容易被国家侵蚀",
     "cards": [
         {"key": "国家必要", "desc": "产权的界定与执行，离不开国家的强制力和法律保护。"},
         {"key": "国家自利", "desc": "国家追求租金最大化，可能抬高索取，甚至直接剥夺产权。"},
         {"key": "所有权残缺", "desc": "Demsetz：国家或他人可以限制所有者对资产的部分权利。"},
     ],
     "takeaway": "悖论：没有国家，产权无从保护；国家一旦介入，又容易造成所有权残缺。",
     "notes": "三张卡片依次讲必要性、自利性、残缺。"},
    {"layout": "compare", "title": "改革前后：谁掌握土地的剩余权", "notes": "两栏对照。",
     "left": {"head": "人民公社体制", "points": [
         "国家用指令计划、统购统销控制农村经济",
         "集体承受控制结果，却没有完整的所有权",
         "监督计量困难，激励不足，效率长期低下"]},
     "right": {"head": "家庭承包制", "points": [
         "农民自发包产到户，国家逐步认可",
         "国家与农户订立新合约，农户获得剩余权",
         "激励恢复，农业产出与收入明显增长"]}},
    {"layout": "timeline", "title": "国家控制费用的两次高峰与政策调整", "notes": "五个节点。",
     "nodes": [
         {"year": "1952", "head": "合作化起步", "desc": "国家开始集中农业剩余"},
         {"year": "1958", "head": "人民公社", "desc": "集体化取代农民私有权"},
         {"year": "1960", "head": "第一次退却", "desc": "承认副业，以队为基础"},
         {"year": "1978", "head": "包产到户", "desc": "费用高峰下自发试点"},
         {"year": "1982", "head": "承包合法化", "desc": "订立新的产权合约"}]},
    {"layout": "quote", "title": "农民表达不满的方式", "notes": "一句话的分量。",
     "quote": "农民留在体制内，并非出于对集体的忠诚，\n而是因为别无选择。",
     "source": "—— 文中对公社制下农民处境的概括",
     "note": "既没有“退出权”，也无权自由“叫喊”，农民只能用怠工表达不满，集体效率随之下降。"},
    {"layout": "numbered", "title": "结论：国家如何保护有效率的产权", "notes": "三条结论。",
     "rows": [
         {"head": "产权保护是增长关键", "desc": "长期经济增长取决于国家能否保护有效率的产权制度。"},
         {"head": "保护不会自动出现", "desc": "国家通常不会主动提供保护，需要来自社会的持续推动。"},
         {"head": "社会与国家的均势", "desc": "农户与新兴产权代理人参与博弈，与国家达成互利交易。"}]},
    {"layout": "end", "info": ["汇报人：stu01", "欢迎批评指正"], "notes": "结束，致谢。"},
]}


def test_content_page_uses_a_template_page_outside_the_four_bases(workdir):
    tpl = _make_template(workdir / "five.pptx")
    prs = Presentation(str(tpl))
    extra = prs.slides.add_slide(prs.slide_layouts[6])
    _textbox(extra, 0.8, 0.4, 8.0, 0.6, "样本标题", 20)
    _textbox(extra, 0.8, 1.6, 8.0, 1.2, "样本正文一", 16)
    _textbox(extra, 0.8, 3.2, 8.0, 1.2, "样本正文二", 16)
    prs.save(str(tpl))
    base = {
        "slide_size": [13.33, 7.5],
        "safe_area": [0.7, 1.2, 12.6, 6.8],
        "pages": {"cover": 1, "section": 2, "content": 3, "end": 4},
        "colors": {"accent": "#9B1B30"},
    }
    analysis = {"specimens": [{
        "index": 4, "tag": "bullets-3", "image_only": False, "capacity_chars": 200,
        "title_i": 0, "body_i": 1, "bodies_i": [1, 2],
    }]}
    plan = {"slides": [{
        "layout": "bullets", "title": "材料与方法",
        "points": ["第一条要点", "第二条要点"], "notes": "讲来源。",
    }]}
    out = workdir / "used.pptx"
    result = ppt_compose.compose(base, plan, tpl, out, None, analysis)
    assert result["template_pages"] == [5]
    texts = _all_texts(Presentation(str(out)))
    assert any("材料与方法" in text for text in texts)
    assert any("第一条要点" in text for text in texts)


def test_bullets_fall_back_to_an_image_text_page(workdir):
    tpl = _make_template(workdir / "imgtext.pptx")
    prs = Presentation(str(tpl))
    extra = prs.slides.add_slide(prs.slide_layouts[6])
    _textbox(extra, 0.8, 0.4, 8.0, 0.6, "图文标题", 20)
    _textbox(extra, 0.8, 1.8, 6.0, 1.4, "图文正文", 16)
    prs.save(str(tpl))
    base = {
        "slide_size": [13.33, 7.5],
        "safe_area": [0.7, 1.2, 12.6, 6.8],
        "pages": {"cover": 1, "section": 2, "content": 3, "end": 4},
        "colors": {"accent": "#9B1B30"},
    }
    analysis = {"specimens": [{
        "index": 4, "tag": "image+text", "image_only": False, "capacity_chars": 80,
        "title_i": 0, "body_i": 1, "bodies_i": [1],
    }]}
    plan = {"slides": [{
        "layout": "bullets", "title": "材料与方法",
        "points": ["第一条要点"], "notes": "讲来源。",
    }]}
    out = workdir / "fallback.pptx"
    result = ppt_compose.compose(base, plan, tpl, out, None, analysis)
    assert result["template_pages"] == [5]
    texts = _all_texts(Presentation(str(out)))
    assert any("第一条要点" in text for text in texts)


@pytest.mark.skipif(ppt_base._template_pptx(REAL_ID) is None, reason="本机没有人大模板 53")
def test_demo_plan_on_real_template(workdir):
    tpl = ppt_base._template_pptx(REAL_ID)
    base = ppt_base.guess_base(tpl)
    out = workdir / "demo9.pptx"
    result = ppt_compose.compose(base, DEMO_PLAN, tpl, out)

    assert result["pages"] == 9
    assert result["overflow"] == []
    prs = Presentation(str(out))
    assert len(prs.slides) == 9

    pages = _footer_digits(prs, base)
    assert pages[0] == [] and pages[8] == []          # 封面、结尾不显示页码
    # 有页码框的页，页号必须等于实际页号；章节页底板本来就没有页码框，那页就没有
    got = {i + 1: p[0] for i, p in enumerate(pages) if p}
    assert got == {2: "2", 4: "4", 5: "5", 6: "6", 7: "7", 8: "8"}

    limit = base["slide_size"][1] - ppt_base.BOTTOM_BAND
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            text = shape.text_frame.text.strip()
            if not text:
                continue
            center = ppt_base._inches(shape.top) + ppt_base._inches(shape.height) / 2
            if text.isdigit() and center > limit:
                continue
            assert not ppt_base.JUNK_PATTERN.search(text), f"残留示例字：{text!r}"


def _card_slide(path: pathlib.Path, title_box=(0.8, 0.35, 10.0, 0.7, "请输入标题", 28)):
    """一页三栏卡片，外加英文装饰框和页脚校名。返回 (路径, 文本→shape_id, shape_id→frame序号)。"""
    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _textbox(slide, *title_box)
    _textbox(slide, 0.6, 1.4, 3.6, 0.5, "卡片一", 18)
    _textbox(slide, 0.6, 2.0, 3.6, 1.4, "卡片一的说明", 14)
    _textbox(slide, 4.8, 1.4, 3.6, 0.5, "卡片二", 18)
    _textbox(slide, 4.8, 2.0, 3.6, 1.4, "卡片二的说明", 14)
    _textbox(slide, 9.0, 1.4, 3.6, 0.5, "卡片三", 18)
    _textbox(slide, 9.0, 2.0, 3.6, 1.4, "卡片三的说明", 14)
    _textbox(slide, 0.6, 4.0, 5.0, 0.4, "Click to add text", 12)
    _textbox(slide, 0.6, 6.85, 3.2, 0.35, "中国人民大学", 12)
    prs.save(str(path))
    opened = Presentation(str(path))
    from services.ppt.pptx_splice import _frames

    by_text = {}
    by_id = {}
    for index, shape in enumerate(_frames(opened.slides[0])):
        text = shape.text_frame.text.strip()
        by_text[text] = shape.shape_id
        by_id[shape.shape_id] = index
    return path, by_text, by_id


def _slot(name, by_text, by_id, sample, required=False):
    sid = by_text[sample]
    return {
        "name": name, "shape_id": sid, "frame": by_id[sid],
        "max_chars": 40, "max_lines": 2, "font_pt": 18, "required": required,
    }


def _card_pages(by_text, by_id, title_lines=2, title_pt=28):
    title = _slot("title", by_text, by_id, "请输入标题", True)
    title["max_lines"] = title_lines
    title["font_pt"] = title_pt
    slots = [
        title,
        _slot("item1.head", by_text, by_id, "卡片一"),
        _slot("item1.body", by_text, by_id, "卡片一的说明"),
        _slot("item2.head", by_text, by_id, "卡片二"),
        _slot("item2.body", by_text, by_id, "卡片二的说明"),
        _slot("item3.head", by_text, by_id, "卡片三"),
        _slot("item3.body", by_text, by_id, "卡片三的说明"),
    ]
    decor = [
        {"shape_id": by_text["Click to add text"], "frame": by_id[by_text["Click to add text"]], "action": "clear"},
        {"shape_id": by_text["中国人民大学"], "frame": by_id[by_text["中国人民大学"]], "action": "keep"},
    ]
    return {"pages": [{"page": 1, "slots": slots, "decor": decor}]}


def _v2_plan(slots, notes="讲这一页。"):
    return {"plan_version": 2, "title": "测试", "slides": [{
        "layout": "cards", "template_page": 1, "slots": slots, "notes": notes,
    }]}


def _shape_map(slide):
    found = {}
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        found[str(getattr(shape, "name", "") or "")] = shape.text_frame.text.strip()
    return found


def test_v2_fills_three_cards_and_keeps_school(workdir):
    path, by_text, by_id = _card_slide(workdir / "cards.pptx")
    pages = _card_pages(by_text, by_id)
    plan = _v2_plan({
        "title": "开篇",
        "item1.head": "先看", "item1.body": "先把材料读完",
        "item2.head": "再做", "item2.body": "把步骤写下来",
        "item3.head": "后查", "item3.body": "对照结果改一版",
    })
    base = ppt_base.guess_base(path)
    out = workdir / "cards-out.pptx"
    result = ppt_compose.compose(base, plan, path, out, pages=pages)
    assert result["template_pages"] == [1]
    prs = Presentation(str(out))
    named = _shape_map(prs.slides[0])
    assert named.get("ars:s1:title") == "开篇"
    assert named.get("ars:s1:item1.head") == "先看"
    assert named.get("ars:s1:item1.body") == "先把材料读完"
    assert named.get("ars:s1:item2.head") == "再做"
    assert named.get("ars:s1:item2.body") == "把步骤写下来"
    assert named.get("ars:s1:item3.head") == "后查"
    assert named.get("ars:s1:item3.body") == "对照结果改一版"
    texts = set(_all_texts(prs))
    assert "Click to add text" not in texts
    assert "中国人民大学" in texts
    assert "卡片一" not in texts

    from services.ppt import ppt_check
    issues = ppt_check.check_deck(out, base)
    assert [item for item in issues if item["kind"] == "offslide"] == []
    assert [item for item in issues if item["kind"] == "overlap"] == []


def test_v2_unused_card_is_cleared(workdir):
    path, by_text, by_id = _card_slide(workdir / "cards.pptx")
    plan = _v2_plan({
        "title": "开篇",
        "item1.head": "先看", "item1.body": "先把材料读完",
        "item2.head": "再做", "item2.body": "把步骤写下来",
    })
    base = ppt_base.guess_base(path)
    out = workdir / "cards-out.pptx"
    ppt_compose.compose(base, plan, path, out, pages=_card_pages(by_text, by_id))
    texts = set(_all_texts(Presentation(str(out))))
    assert "卡片三" not in texts
    assert "卡片三的说明" not in texts
    assert "先看" in texts and "再做" in texts


def test_v2_long_title_shrinks_and_notes_keep_original(workdir):
    path, by_text, by_id = _card_slide(
        workdir / "tiny.pptx", title_box=(0.8, 0.35, 1.0, 0.35, "请输入标题", 32))
    pages = _card_pages(by_text, by_id, title_lines=1, title_pt=32)
    original = "学习方法要从知道走到做到并且能在真实场景里反复练习直到形成习惯"
    plan = _v2_plan({"title": original}, notes="讲标题。")
    base = ppt_base.guess_base(path)
    out = workdir / "tiny-out.pptx"
    result = ppt_compose.compose(base, plan, path, out, pages=pages)
    prs = Presentation(str(out))
    title = None
    for shape in prs.slides[0].shapes:
        if getattr(shape, "name", "") == "ars:s1:title":
            title = shape
    assert title is not None
    sizes = [run.font.size.pt for para in title.text_frame.paragraphs for run in para.runs if run.font.size]
    assert sizes and min(sizes) >= 10
    notes = prs.slides[0].notes_slide.notes_text_frame.text
    assert original in notes
    assert any(item.get("field") == "title" for item in result["overflow"])


def test_clone_keeps_shape_id_and_frame_fallback(workdir):
    path, by_text, _by_id = _card_slide(workdir / "ids.pptx")
    from services.ppt.pptx_splice import _clone_slide, _frames

    prs = Presentation(str(path))
    before = [shape.shape_id for shape in _frames(prs.slides[0])]
    cloned = _clone_slide(prs, prs.slides[0])
    after = [shape.shape_id for shape in _frames(cloned)]
    assert before == after
    wanted = by_text["卡片二的说明"]
    found = ppt_compose._shape_for_slot(cloned, {"shape_id": wanted})
    assert found is not None and found.shape_id == wanted
    fallback = ppt_compose._shape_for_slot(cloned, {"shape_id": 999999, "frame": after.index(wanted)})
    assert fallback is not None and fallback.shape_id == wanted


def _template55() -> pathlib.Path | None:
    root = pathlib.Path(__file__).resolve().parents[1] / "data"
    hits = list(root.rglob("4d7e82ced452.pptx"))
    return hits[0] if hits else None


def _grouped_text(slide) -> str:
    parts = []
    for shape in ppt_compose._iter_shapes(slide):
        if getattr(shape, "has_text_frame", False):
            parts.append(shape.text_frame.text)
    return "\n".join(parts)


def test_v2_clears_sample_inside_groups(workdir):
    src = _template55()
    if src is None:
        pytest.skip("本机没有模板-55")
    pages = {"pages": [{"page": 17, "slots": [{
        "name": "title", "shape_id": 8, "frame": 0,
        "font_pt": 28, "max_lines": 1, "required": True,
    }], "decor": []}]}
    plan = {"plan_version": 2, "title": "测试", "slides": [{
        "layout": "section", "template_page": 17,
        "slots": {"title": "例子"}, "notes": "章节。",
    }]}
    base = ppt_base.load_base("4d7e82ced452") or ppt_base.guess_base(src)
    out = workdir / "group.pptx"
    ppt_compose.compose(base, plan, src, out, pages=pages)
    blob = _grouped_text(Presentation(str(out)).slides[0])
    assert "例子" in blob
    assert "质量方面" not in blob
    assert "请在这里输入" not in blob


def test_old_plan_ignores_pages_catalog(workdir):
    tpl = _make_template(workdir / "demo.pptx")
    base = ppt_base.guess_base(tpl)
    plan = _small_plan()
    out_old = workdir / "old.pptx"
    out_new = workdir / "old-with-pages.pptx"
    pages = {"pages": [{"page": 1, "slots": [{"name": "title", "shape_id": 1, "frame": 0}]}]}
    ppt_compose.compose(base, plan, tpl, out_old)
    ppt_compose.compose(base, plan, tpl, out_new, pages=pages)
    left = Presentation(str(out_old))
    right = Presentation(str(out_new))
    assert _all_texts(left) == _all_texts(right)
    assert [len(slide.shapes) for slide in left.slides] == [len(slide.shapes) for slide in right.slides]


def test_narrow_title_stretches_up_to_the_logo(workdir):
    path = workdir / "grow.pptx"
    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _textbox(slide, 0.8, 0.4, 0.7, 0.5, "缺口", 28)
    _textbox(slide, 11.2, 0.35, 1.5, 0.45, "中国人民大学", 14)
    _textbox(slide, 0.8, 2.0, 6.0, 1.2, "正文说明放在这里", 16)
    prs.save(str(path))
    opened = Presentation(str(path))
    by_text = {
        shape.text_frame.text.strip(): shape.shape_id
        for shape in opened.slides[0].shapes if shape.has_text_frame
    }
    pages = {"pages": [{"page": 1, "slots": [
        {"name": "title", "shape_id": by_text["缺口"], "frame": 0,
         "max_chars": 24, "max_lines": 1, "font_pt": 28, "required": True, "grow_w": 10.32},
        {"name": "body", "shape_id": by_text["正文说明放在这里"], "frame": 2,
         "max_chars": 80, "max_lines": 4, "font_pt": 16, "required": True},
    ], "decor": [
        {"shape_id": by_text["中国人民大学"], "frame": 1, "action": "keep"},
    ]}]}
    base = ppt_base.guess_base(path)

    def _compose(title, out_name):
        out = workdir / out_name
        ppt_compose.compose(base, {"plan_version": 2, "slides": [{
            "layout": "cards", "template_page": 1,
            "slots": {"title": title, "body": "正文说明放在这里"},
            "notes": "讲稿。",
        }]}, path, out, pages=pages)
        done = Presentation(str(out))
        return done, next(shape for shape in done.slides[0].shapes if getattr(shape, "name", "") == "ars:s1:title")

    _short_deck, short = _compose("好", "grow-short.pptx")
    assert ppt_base._inches(short.width) == pytest.approx(0.7, abs=0.08)
    grown_deck, grown = _compose("提取练习比反复阅读记得更牢", "grow-long.pptx")
    right = ppt_base._inches(grown.left) + ppt_base._inches(grown.width)
    assert ppt_base._inches(grown.width) > 2
    assert right < 11.2
    assert right < 13.33
    assert grown.text_frame.text.strip() == "提取练习比反复阅读记得更牢"
    body = next(shape for shape in grown_deck.slides[0].shapes if getattr(shape, "name", "") == "ars:s1:body")
    assert ppt_base._inches(body.width) == pytest.approx(6.0, abs=0.08)


def test_plan_lines_accepts_string_cards():
    lines = ppt_compose._plan_lines({
        "layout": "cards",
        "cards": ["直接一句", {"key": "甲", "desc": "乙"}],
    })
    texts = [text for _level, text in lines]
    assert "直接一句" in texts
    assert "甲：乙" in texts


def _toc_rows(path: pathlib.Path):
    """四行目录：每行一根空条、一个序号、一条标题。"""
    from pptx.enum.shapes import MSO_SHAPE

    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _textbox(slide, 0.8, 0.3, 4.0, 0.4, "目录", 28)
    labels = []
    for index, label in enumerate(("背景", "方法", "结果", "总结"), 1):
        top = 1.2 + (index - 1) * 1.1
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(top), Inches(6.0), Inches(0.5))
        bar.text_frame.text = ""
        number = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.9), Inches(top), Inches(0.45), Inches(0.45))
        number.text_frame.paragraphs[0].text = str(index)
        labels.append(_textbox(slide, 1.6, top, 4.0, 0.45, label, 18))
    prs.save(str(path))
    opened = Presentation(str(path))
    by_text = {}
    for shape in opened.slides[0].shapes:
        if shape.has_text_frame and shape.text_frame.text.strip():
            by_text[shape.text_frame.text.strip()] = shape.shape_id
    return by_text


def test_toc_with_two_chapters_drops_the_empty_rows(workdir):
    path = workdir / "toc-four.pptx"
    by_text = _toc_rows(path)
    slots = []
    for index, label in enumerate(("背景", "方法", "结果", "总结"), 1):
        slots.append({
            "name": f"item{index}.head", "shape_id": by_text[label], "frame": index,
            "max_chars": 12, "max_lines": 1, "font_pt": 18, "required": True,
        })
    pages = {"pages": [{"page": 1, "role": "toc", "slots": slots, "decor": [
        {"shape_id": by_text["目录"], "frame": 0, "action": "keep"},
    ]}]}
    plan = {"plan_version": 2, "slides": [{
        "layout": "toc", "template_page": 1,
        "slots": {"item1.head": "研究背景", "item2.head": "提取练习"},
        "notes": "目录两章。",
    }]}
    out = workdir / "toc-out.pptx"
    ppt_compose.compose(ppt_base.guess_base(path), plan, path, out, pages=pages)
    done = Presentation(str(out))
    texts = [shape.text_frame.text.strip() for shape in done.slides[0].shapes if shape.has_text_frame]
    assert "研究背景" in texts
    assert "提取练习" in texts
    assert "目录" in texts
    assert "结果" not in texts
    assert "总结" not in texts
    assert "3" not in texts
    assert "4" not in texts
    assert "1" in texts
    assert "2" in texts
    assert len(done.slides[0].shapes) == 7


def test_content_page_keeps_an_empty_row(workdir):
    path = workdir / "content-four.pptx"
    by_text = _toc_rows(path)
    slots = [{
        "name": f"item{index}.head", "shape_id": by_text[label], "frame": index,
        "max_chars": 12, "max_lines": 1, "font_pt": 18, "required": True,
    } for index, label in enumerate(("背景", "方法", "结果", "总结"), 1)]
    pages = {"pages": [{"page": 1, "role": "content", "slots": slots, "decor": [
        {"shape_id": by_text["目录"], "frame": 0, "action": "keep"},
    ]}]}
    plan = {"plan_version": 2, "slides": [{
        "layout": "cards", "template_page": 1,
        "slots": {"item1.head": "研究背景", "item2.head": "提取练习"},
        "notes": "内容页。",
    }]}
    out = workdir / "content-out.pptx"
    ppt_compose.compose(ppt_base.guess_base(path), plan, path, out, pages=pages)
    done = Presentation(str(out))
    assert len(done.slides[0].shapes) == 13
    texts = _all_texts(done)
    assert "研究背景" in texts
    assert "提取练习" in texts
