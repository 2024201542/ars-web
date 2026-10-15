"""改过的 PPT 可以撤回到改前的字节。"""

import json
import pathlib
import uuid

import pytest
from pptx import Presentation
from pptx.util import Inches

from services.ppt import ppt_edit, ppt_history


@pytest.fixture(autouse=True)
def reset_config(monkeypatch):
    sandbox = pathlib.Path(__file__).resolve().parents[3] / ".tmp" / "pytest-env"
    monkeypatch.setenv("ARS_WEB_DB_PATH", str(sandbox / "test_ars.db"))
    monkeypatch.setenv("ARS_SKILLS_PATH", str(sandbox / "fake_skills"))
    monkeypatch.setenv("ARS_WORKSPACE_DIR", str(sandbox / "workspaces"))
    monkeypatch.setenv("ARS_PPT_HISTORY_DIR", str(sandbox / "ppt-history" / uuid.uuid4().hex))


USER = "hist-user"
SOURCE = "site"
REL = "demo.pptx"


def _write_deck(path: pathlib.Path, text: str) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(8), Inches(1))
    shape.name = "ars:s1:body"
    shape.text_frame.text = text
    prs.save(str(path))


def _write_plan(path: pathlib.Path, text: str) -> None:
    path.write_text(json.dumps({
        "title": "测试",
        "slides": [{"title": "一页", "slots": {"body": text}, "notes": text}],
    }, ensure_ascii=False), encoding="utf-8")


def _change(pptx: pathlib.Path, plan: pathlib.Path, text: str) -> None:
    ppt_edit.write_slot_texts(pptx, 1, {"body": text})
    _write_plan(plan, text)


def _snap(pptx: pathlib.Path, plan: pathlib.Path, instruction: str, rel: str) -> int:
    return ppt_history.snapshot(
        USER, pptx, plan,
        {"instruction": instruction, "ops": [{"op": "rewrite", "page": 1}], "changed_pages": [1]},
        source=SOURCE, rel=rel, plan_source=SOURCE, plan_rel=plan.name)


def test_undo_restores_pptx_bytes_and_plan():
    folder = pathlib.Path(__file__).resolve().parents[3] / ".tmp" / "pytest-env" / "hist-deck"
    folder.mkdir(parents=True, exist_ok=True)
    pptx = folder / "demo.pptx"
    plan = folder / "demo.json"
    _write_deck(pptx, "原来的正文")
    _write_plan(plan, "原来的正文")
    before_pptx = pptx.read_bytes()
    before_plan = json.loads(plan.read_text(encoding="utf-8"))
    _snap(pptx, plan, "改一次", "one.pptx")
    _change(pptx, plan, "改过的正文")
    assert pptx.read_bytes() != before_pptx
    restored = ppt_history.restore(USER, pptx, plan, None, source=SOURCE, rel="one.pptx")
    assert restored == 1
    assert pptx.read_bytes() == before_pptx
    assert json.loads(plan.read_text(encoding="utf-8")) == before_plan


def test_two_undos_return_to_after_first_edit():
    folder = pathlib.Path(__file__).resolve().parents[3] / ".tmp" / "pytest-env" / "hist-deck-3"
    folder.mkdir(parents=True, exist_ok=True)
    pptx = folder / "demo.pptx"
    plan = folder / "demo.json"
    _write_deck(pptx, "原")
    _write_plan(plan, "原")
    rel = "three.pptx"
    _snap(pptx, plan, "第一次", rel)
    _change(pptx, plan, "第一次之后")
    after_first = pptx.read_bytes()
    after_first_plan = json.loads(plan.read_text(encoding="utf-8"))
    _snap(pptx, plan, "第二次", rel)
    _change(pptx, plan, "第二次之后")
    _snap(pptx, plan, "第三次", rel)
    _change(pptx, plan, "第三次之后")
    ppt_history.restore(USER, pptx, plan, None, source=SOURCE, rel=rel)
    ppt_history.restore(USER, pptx, plan, None, source=SOURCE, rel=rel)
    assert pptx.read_bytes() == after_first
    assert json.loads(plan.read_text(encoding="utf-8")) == after_first_plan


def test_twenty_first_version_drops_the_oldest():
    folder = pathlib.Path(__file__).resolve().parents[3] / ".tmp" / "pytest-env" / "hist-deck-21"
    folder.mkdir(parents=True, exist_ok=True)
    pptx = folder / "demo.pptx"
    plan = folder / "demo.json"
    _write_deck(pptx, "版本")
    _write_plan(plan, "版本")
    rel = "many.pptx"
    for index in range(1, 22):
        _snap(pptx, plan, f"第{index}次", rel)
    rows = ppt_history.history(USER, pptx, source=SOURCE, rel=rel)
    assert len(rows) == 20
    assert rows[0]["version"] == 2
    assert rows[-1]["version"] == 21
    stored = ppt_history.folder_for(USER, SOURCE, rel)
    assert not (stored / "v1.pptx").is_file()
    assert (stored / "v21.pptx").is_file()


async def test_undo_outside_user_dir_is_403():
    from fastapi import HTTPException

    from routers.template_library import UndoRequest, undo_edit

    with pytest.raises(HTTPException) as caught:
        await undo_edit(UndoRequest(path=r"C:\Windows\win.ini", source=""), "user-test")
    assert caught.value.status_code == 403
