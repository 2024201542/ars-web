"""按计划出片：底板页换字、内容页用版式库画、写讲稿备注、清残留示例字。

单页出问题不中断整份：超限的字段截到上限重画一次（记进 `overflow`），
重画还画不出来就留底板（记进 `warnings`）。
"""

from __future__ import annotations

import copy
import logging
import pathlib
import re
import tempfile

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches, Pt

from services.ppt import ppt_base, ppt_image, ppt_layouts, ppt_template
from services.ppt.pptx_splice import _clone_slide, _delete_slides, sync_app_props

logger = logging.getLogger(__name__)


def _safe_edge(base: dict) -> float:
    safe = (base or {}).get("safe_area") or ppt_base.DEFAULT_SAFE
    return float(list(safe)[2])


def _slide_size(base: dict) -> tuple[float, float]:
    size = (base or {}).get("slide_size") or [13.33, 7.5]
    return float(list(size)[0]), float(list(size)[1])


def _slide_height(base: dict) -> float:
    return _slide_size(base)[1]


def _source(originals: list, page, fallback: int):
    try:
        index = int(page) - 1
    except (TypeError, ValueError):
        index = fallback - 1
    if index < 0 or index >= len(originals):
        index = max(0, min(len(originals) - 1, fallback - 1))
    return originals[index]


def _find_shape(slide, want: str):
    """按模板原文找文本框。"""
    want = (want or "").strip()
    if not want:
        return None
    for shape in slide.shapes:
        if not ppt_template.is_text_candidate(shape):
            continue
        if shape.text_frame.text.strip() == want:
            return shape
    return None


def _biggest_shape(slide):
    """兜底：字号最大的文本框，其次最靠上的。"""
    best = None
    for shape in slide.shapes:
        if not ppt_template.is_text_candidate(shape):
            continue
        text = shape.text_frame.text.strip()
        if not text:
            continue
        key = (-(ppt_base._size_of(shape) or 0), ppt_base._inches(shape.top))
        if best is None or key < best[0]:
            best = (key, shape)
    return best[1] if best else None


def _replace(shape, text, base: dict, right_limit: float | None = None,
             shrink: bool = False, widen: bool = True):
    """换字并拉宽到右边界（不换行、左对齐），保住模板原有的字体样式。

    `shrink` 只给标题开：标题继承的字号常是 60–70pt，换成更长的标题会冲出页面；
    副标题/信息行不碰字号——它们的字号也继承自母版，猜错了会把小字放大。
    `widen` 给章节页的序号关掉：序号框和标题框在同一行，拉宽就会压到标题上。
    """
    if shape is None:
        return False
    value = "" if text is None else str(text)
    edge = right_limit if right_limit else _safe_edge(base)
    edge = min(edge, _slide_size(base)[0] - 0.08)
    ppt_base._set_text(shape, value)
    if widen:
        ppt_base._widen(shape, edge)
    shape.text_frame.word_wrap = False
    ppt_base._left(shape)
    if shrink and value:
        _shrink_to_fit(shape, value, edge)
    return True


def _shrink_to_fit(shape, text: str, edge: float) -> None:
    """底板上的文字不换行，所以只按"这一行放得下"缩字号。

    模板标题的字号常常继承自母版（run 上没有显式 size），这时按 44pt 估个上限，
    并把算出来的字号**显式**写上去——否则继承的 60–70pt 会一直冲出一行。
    留 4% 余量，_em 估算和实际渲染宽度不会刚好卡边。
    """
    left = ppt_base._inches(shape.left)
    width = max(1.0, (edge - left) * 0.96)
    original = ppt_base._size_of(shape)
    max_pt = original if original else 44
    pt = ppt_layouts.fit_one_line(text, width, max_pt, ppt_layouts.TITLE_MIN_PT)
    if original is not None and pt >= original:
        return
    for para in shape.text_frame.paragraphs:
        for run in para.runs:
            run.font.size = Pt(pt)


def _drop_offscreen(slide) -> None:
    """删掉跑到页面外的文本框。

    只看文本框：封面的大水印/装饰图常常故意把 top 设成负数做成"出血"效果，
    一起删就会把水印弄没（踩过一次）。
    """
    for shape in list(slide.shapes):
        if not ppt_template.is_text_candidate(shape):
            continue
        if ppt_base._inches(shape.top) < 0:
            ppt_base._drop(shape)


def _clear_page_number(slide, base: dict) -> None:
    """封面和结尾不显示页码。"""
    limit = _slide_height(base) - ppt_base.BOTTOM_BAND
    for shape in list(slide.shapes):
        if not ppt_template.is_text_candidate(shape):
            continue
        text = shape.text_frame.text.strip()
        if not text or not text.isdigit():
            continue
        center = ppt_base._inches(shape.top) + ppt_base._inches(shape.height) / 2
        if center > limit:
            ppt_base._set_text(shape, "")


def _title_edge(base: dict) -> float:
    """标题类文字最多拉到哪儿：档案里内容页槽位记的 title_right_limit。"""
    slot = ((base.get("slots") or {}).get("content") or {})
    try:
        return float(slot.get("title_right_limit") or ppt_base.DEFAULT_TITLE_RIGHT)
    except (TypeError, ValueError):
        return float(ppt_base.DEFAULT_TITLE_RIGHT)


def _fill_cover(slide, base: dict, item: dict) -> list:
    slots = ((base.get("slots") or {}).get("cover") or {})
    _drop_offscreen(slide)
    touched = []
    title_shape = _find_shape(slide, slots.get("title")) or _biggest_shape(slide)
    if _replace(title_shape, item.get("title") or "", base, _title_edge(base), shrink=True):
        touched.append(title_shape)

    subtitle_shape = _find_shape(slide, slots.get("subtitle"))
    if _replace(subtitle_shape, item.get("subtitle") or "", base):
        touched.append(subtitle_shape)

    info = list(item.get("info") or [])
    info_shapes = []
    for i, source in enumerate(slots.get("info") or []):
        shape = _find_shape(slide, source)
        if _replace(shape, info[i] if i < len(info) else "", base):
            touched.append(shape)
            info_shapes.append(shape)
    _separate_cover_info(slide, subtitle_shape, info_shapes, info, touched)
    _clear_page_number(slide, base)
    return touched


def _separate_cover_info(slide, subtitle_shape, info_shapes, info_texts, touched) -> None:
    """副标题和信息行叠在一起时，把信息行挪到副标题下面单独一行。"""
    shapes = [shape for shape in info_shapes if shape is not None]
    if not shapes:
        return
    hit_subtitle = subtitle_shape is not None and any(
        ppt_base._overlaps(subtitle_shape, shape) for shape in shapes)
    hit_each_other = any(
        ppt_base._overlaps(a, b) for i, a in enumerate(shapes) for b in shapes[i + 1:])
    if not hit_subtitle and not hit_each_other:
        return
    merged = "  ".join(str(text) for text in info_texts if text)
    for shape in shapes:
        ppt_base._set_text(shape, "")
    if subtitle_shape is None:
        return
    if not merged:
        return
    top = ppt_base._inches(subtitle_shape.top) + ppt_base._inches(subtitle_shape.height) + 0.08
    left = ppt_base._inches(subtitle_shape.left)
    box = slide.shapes.add_textbox(Inches(left), Inches(min(top, 6.4)), Inches(8), Inches(0.36))
    ppt_base._set_text(box, merged)
    ppt_base._mark_ours(box)
    touched.append(box)


def _fill_section(slide, base: dict, item: dict) -> list:
    slots = ((base.get("slots") or {}).get("section") or {})
    filled = []
    title_shape = _find_shape(slide, slots.get("title")) or _biggest_shape(slide)
    index_shape = _find_shape(slide, slots.get("index"))

    # 序号和标题在同一行：序号拉到"标题左边界前 0.1 英寸"为止——
    # 拉满会盖住标题，完全不拉会把大号数字挤出框外（96pt 的「01」塞不进 1.3 英寸）
    if index_shape is not None:
        limit = None
        if title_shape is not None:
            left = ppt_base._inches(index_shape.left)
            limit = max(left + 0.5, ppt_base._inches(title_shape.left) - 0.1)
        _replace(index_shape, item.get("index") or "", base, limit, widen=limit is not None)
        filled.append(index_shape)
    if _replace(title_shape, item.get("title") or "", base, _title_edge(base), shrink=True):
        filled.append(title_shape)

    subs = list(item.get("subs") or [])
    placed = 0
    for i, source in enumerate(slots.get("subs") or []):
        shape = _find_shape(slide, source)
        if shape is None:
            continue
        _replace(shape, subs[i] if i < len(subs) else "", base)
        filled.append(shape)
        placed += 1
    if index_shape is None and (item.get("index") or ""):
        box = slide.shapes.add_textbox(Inches(0.7), Inches(2.2), Inches(2.0), Inches(0.9))
        ppt_base._set_text(box, str(item.get("index")))
        ppt_base._mark_ours(box)
        filled.append(box)
    if subs and placed < len(subs):
        anchor = title_shape or (filled[-1] if filled else None)
        top = 3.4
        if anchor is not None:
            top = ppt_base._inches(anchor.top) + ppt_base._inches(anchor.height) + 0.15
        box = slide.shapes.add_textbox(Inches(0.7), Inches(min(top, 5.8)), Inches(10), Inches(0.4))
        ppt_base._set_text(box, "  ".join(str(text) for text in subs if text))
        ppt_base._mark_ours(box)
        filled.append(box)
    ppt_base._drop_lines_over(slide, filled)
    return filled


def _fill_end(slide, base: dict, item: dict) -> list:
    slots = ((base.get("slots") or {}).get("end") or {})
    info = list(item.get("info") or [])
    touched = []
    for i, source in enumerate(slots.get("info") or []):
        shape = _find_shape(slide, source)
        if _replace(shape, info[i] if i < len(info) else "", base):
            touched.append(shape)
    _clear_page_number(slide, base)
    return touched


# ───────────── 内容页 ─────────────

def _slot_of(plan: dict, field: str):
    """字段路径解析（规则在版式库里，出片和改计划共用）。"""
    return ppt_layouts.field_slot(plan, field)


def _rollback(slide, before: list) -> None:
    """删掉这次画失败留下的形状，回到画之前。"""
    for shape in list(slide.shapes):
        if shape._element not in before:
            ppt_base._drop(shape)


def _trim_plan(plan_slide: dict, exc, base: dict) -> dict:
    """按报出来的上限把那个字段截短，其它字段不动。"""
    out = copy.deepcopy(plan_slide)
    slot = _slot_of(out, exc.field)
    if not slot:
        return out
    container, key = slot
    original = str(container[key])
    limit = exc.limit
    if limit is None:
        # 字数没超上限却排不下：砍掉两成再试
        limit = max(1, int(ppt_layouts.text_len(original) * 0.8))
    container[key] = ppt_layouts.truncate(original, limit)
    return out


def _place_image(slide, base: dict, info: dict, images, page_no: int, warnings: list) -> None:
    slot = (info or {}).get("image_slot")
    if not slot:
        return
    x, y, w, h = slot
    pick = None
    index = (info or {}).get("image")
    pool = list(images or [])
    if pool:
        pick = pool[index - 1] if isinstance(index, int) and 1 <= index <= len(pool) else pool[0]
    if not pick:
        warnings.append({"page": page_no, "detail": "图文页没有可用配图，保留占位框"})
        return
    try:
        ratio = (w / h) if h else 1.0
        stem = re.sub(r"[^\w\-]+", "_", pathlib.Path(pick).stem)[:20] or "img"
        tmp = pathlib.Path(tempfile.gettempdir()) / f"ppt_slot_{page_no}_{stem}.png"
        fitted = ppt_image.fit_to_slot(pick, ratio, tmp)
        if not fitted:
            raise ValueError("按图片位比例裁剪失败")
        for shape in list(slide.shapes):
            if shape.name == ppt_layouts.PIC_SLOT_NAME:
                ppt_base._drop(shape)
        slide.shapes.add_picture(str(fitted), Inches(x), Inches(y), Inches(w), Inches(h))
    except Exception as exc:
        warnings.append({"page": page_no, "detail": f"配图插不进去：{exc}"})


def _draw_content(slide, base: dict, item: dict, images, page_no: int,
                  overflow: list, warnings: list) -> tuple[dict, list]:
    """画一页内容，返回（附加信息, 我们画上去的形状）。"""
    before = [shape._element for shape in slide.shapes]
    try:
        info = ppt_layouts.render(slide, base, item)
    except ppt_layouts.Overflow as exc:
        overflow.append({"page": page_no, "field": exc.field, "text": exc.text, "limit": exc.limit})
        _rollback(slide, before)
        try:
            info = ppt_layouts.render(slide, base, _trim_plan(item, exc, base))
        except ppt_layouts.Overflow as again:
            warnings.append({"page": page_no, "detail": f"截字后仍放不下：{again.field}"})
            _rollback(slide, before)
            return {}, []

    drawn = [shape for shape in slide.shapes if shape._element not in before]
    ppt_base._drop_lines_over(slide, drawn)
    _place_image(slide, base, info, images, page_no, warnings)
    return info, drawn


# ───────────── 收尾 ─────────────

def _mark_all(shapes) -> None:
    for shape in shapes or []:
        ppt_base._mark_ours(shape)


def _clear_unmarked_text(slide) -> None:
    """写失败的页：没标成我们写的字都是模板原文，清掉。组合里的也算。"""
    for shape in _iter_shapes(slide):
        if not getattr(shape, "has_text_frame", False):
            continue
        if str(getattr(shape, "name", "") or "").startswith("ars:"):
            continue
        if str(shape.text_frame.text or "").strip():
            ppt_base._set_text(shape, "")


def _scrub_junk(prs, base: dict) -> int:
    """整份再扫一遍模板残留的示例字。

    只清"不是我们画的"文本框：卡片编号「01」、页码这类我们自己的数字，
    画的时候加了 ars: 名字，这里按名字跳过。页脚里的纯数字是页码，也不动。
    """
    limit = _slide_height(base) - ppt_base.BOTTOM_BAND
    cleaned = 0
    for slide in prs.slides:
        for shape in list(_iter_shapes(slide)):
            if not ppt_template.is_text_candidate(shape):
                continue
            if str(getattr(shape, "name", "") or "").startswith("ars:"):
                continue
            text = shape.text_frame.text.strip()
            if not text:
                continue
            center = ppt_base._inches(shape.top) + ppt_base._inches(shape.height) / 2
            if text.isdigit() and center > limit:
                continue
            if ppt_base.is_junk_text(text):
                ppt_base._set_text(shape, "")
                cleaned += 1
    return cleaned


def _plan_lines(item: dict) -> list[tuple[int, str]]:
    """把画布计划里的字段收成「一条一行」，好填进模板样本页的正文框。"""
    layout = item.get("layout") or "bullets"
    lines: list[tuple[int, str]] = []

    def add(text: str, level: int = 0) -> None:
        text = str(text or "").strip()
        if text:
            lines.append((level, text))

    if layout == "cards":
        for card in item.get("cards") or []:
            if not isinstance(card, dict):
                add(card)
                continue
            key = str(card.get("key") or "").strip()
            desc = str(card.get("desc") or "").strip()
            add(f"{key}：{desc}" if key and desc else (key or desc))
        add(item.get("takeaway") or "")
    elif layout == "compare":
        for side in (item.get("left"), item.get("right")):
            if not isinstance(side, dict):
                continue
            add(side.get("title") or "")
            for point in side.get("points") or []:
                add(point, 1)
    elif layout == "numbered":
        for row in item.get("rows") or []:
            if not isinstance(row, dict):
                add(row)
                continue
            head = str(row.get("head") or "").strip()
            desc = str(row.get("desc") or "").strip()
            add(f"{head}：{desc}" if head and desc else (head or desc))
    elif layout == "quote":
        add(item.get("quote") or "")
        add(item.get("source") or "")
    elif layout == "timeline":
        for event in item.get("events") or []:
            if not isinstance(event, dict):
                add(event)
                continue
            bits = [str(event.get(key) or "").strip() for key in ("year", "head", "desc")]
            add(" ".join(bit for bit in bits if bit))
    elif layout == "toc":
        for entry in item.get("entries") or []:
            if not isinstance(entry, dict):
                add(entry)
                continue
            add(f"{entry.get('label') or ''} {entry.get('desc') or ''}")
    elif layout == "bignum":
        for stat in item.get("stats") or []:
            if not isinstance(stat, dict):
                add(stat)
                continue
            add(f"{stat.get('value') or ''} {stat.get('label') or ''}")
    else:
        for point in item.get("points") or []:
            add(point)
    if not lines:
        for bullet in item.get("bullets") or []:
            text = bullet.get("text") if isinstance(bullet, dict) else bullet
            add(text)
    return lines


def _reserved_indexes(base: dict) -> set[int]:
    pages = (base or {}).get("pages") or {}
    reserved: set[int] = set()
    for key in ("cover", "section", "content", "end"):
        try:
            number = int(pages.get(key) or 0)
        except (TypeError, ValueError):
            continue
        if number > 0:
            reserved.add(number - 1)
    return reserved


_CONTENT_FALLBACK = (
    ppt_template.TAG_IMAGE_TEXT,
    ppt_template.TAG_CONTENT,
    "cards",
    ppt_template.TAG_CARDS,
    ppt_template.TAG_TWO_COL,
    ppt_template.TAG_TIMELINE,
)


def _usable_specimen(spec: dict, reserved: set[int]) -> bool:
    if spec.get("image_only"):
        return False
    if spec.get("index") in reserved:
        return False
    return spec.get("title_i") is not None or bool(spec.get("bodies_i") or spec.get("body_i") is not None)


def _pick_other_page(analysis: dict | None, item: dict, reserved: set[int], used: set[int]):
    """从底板四页以外挑一页来填字。

    先挑版式对得上的。对不上时，用其余的图文页，章节过渡页不拿来当正文。
    还是没有，就返回 None，改在内容底板的安全区里重画。
    """
    if not analysis:
        return None
    from services.ppt.pptx_splice import _want_tag

    tag = _want_tag(item)
    family = tag.split("-")[0]
    exact = []
    fallback = []
    for spec in analysis.get("specimens") or []:
        if not _usable_specimen(spec, reserved):
            continue
        spec_tag = str(spec.get("tag") or "")
        if spec_tag == tag or spec_tag == family or spec_tag.startswith(family + "-"):
            exact.append(spec)
            continue
        if family in (ppt_template.TAG_COVER, ppt_template.TAG_SECTION, ppt_template.TAG_END):
            continue
        if spec_tag in _CONTENT_FALLBACK or spec_tag.startswith(ppt_template.TAG_BULLETS):
            fallback.append(spec)
    specs = exact or fallback
    if not specs:
        return None
    lines = _plan_lines(item)
    needed = len(str(item.get("title") or "")) + sum(len(text) for _, text in lines)
    fresh = [spec for spec in specs if spec.get("index") not in used] or specs
    enough = [spec for spec in fresh if (spec.get("capacity_chars") or 0) >= max(20, int(needed * 0.4))]
    return (enough or fresh)[0]


def _fill_specimen(slide, spec: dict, item: dict) -> str:
    """把这一页的字填进样本页自己的文本框，并标成我们写的。返回要写进备注的全文。"""
    from services.ppt.pptx_splice import _fill_frame, _frames, _scrub_placeholders, _set_notes

    shapes = _frames(slide)
    frames = []
    for shape in shapes:
        try:
            frames.append(shape.text_frame)
        except Exception:
            frames.append(None)
    written: set[int] = set()
    title = str(item.get("title") or "").strip()
    title_i = spec.get("title_i")
    if title and title_i is not None and title_i < len(frames) and frames[title_i] is not None:
        _fill_frame(frames[title_i], [(0, title)], keep_bold=True, add_prefix=False)
        written.add(title_i)
    lines = _plan_lines(item)
    bodies = [i for i in (spec.get("bodies_i") or []) if i is not None]
    if not bodies and spec.get("body_i") is not None:
        bodies = [spec["body_i"]]
    body2 = spec.get("body2_i")
    if body2 is not None and body2 not in bodies:
        bodies.append(body2)
    for slot, idx in enumerate(bodies):
        if slot >= len(lines) or idx >= len(frames) or frames[idx] is None:
            continue
        _fill_frame(frames[idx], [lines[slot]])
        written.add(idx)
    for idx in written:
        if idx < len(shapes):
            ppt_base._mark_ours(shapes[idx])
    overflow = lines[len(bodies):]
    notes = str(item.get("notes") or "").strip()
    if overflow:
        extra = "\n".join(f"· {text}" for _, text in overflow)
        notes = (notes + "\n\n（这一页框不够，多出来的放在备注）\n" + extra).strip()
    if notes:
        _set_notes(slide, notes)
    _scrub_placeholders(slide, written, force=False)
    return notes


def _iter_shapes(container):
    """页面上的形状，包含组合里的。页说明的编号有时落在组合内部。"""
    for shape in container.shapes:
        yield shape
        try:
            grouped = shape.shape_type == MSO_SHAPE_TYPE.GROUP
        except Exception:
            grouped = False
        if grouped:
            yield from _iter_shapes(shape)


def _shape_for_slot(slide, slot: dict):
    """按页说明里的形状编号找框。克隆没保住编号时，按文本框序号回退。"""
    from services.ppt.pptx_splice import _frames

    wanted = (slot or {}).get("shape_id")
    if wanted is not None:
        try:
            number = int(wanted)
        except (TypeError, ValueError):
            number = None
        if number is not None:
            for shape in _iter_shapes(slide):
                if getattr(shape, "shape_id", None) == number:
                    return shape
    frame = (slot or {}).get("frame")
    frames = _frames(slide)
    try:
        index = int(frame)
    except (TypeError, ValueError):
        return None
    if 0 <= index < len(frames):
        return frames[index]
    return None


def _slot_point_size(shape, text: str, spec: dict) -> int | None:
    """按格子的字号上限和行数，缩到放得进框。最小 10pt，放不下返回 None。"""
    import math

    width = max(0.4, ppt_base._inches(shape.width))
    height = max(0.2, ppt_base._inches(shape.height))
    try:
        max_pt = int(spec.get("font_pt") or 0)
    except (TypeError, ValueError):
        max_pt = 0
    if max_pt <= 0:
        max_pt = int(ppt_base._size_of(shape) or 18)
    max_pt = max(10, min(max_pt, 72))
    try:
        max_lines = int(spec.get("max_lines") or 0)
    except (TypeError, ValueError):
        max_lines = 0
    for pt in range(max_pt, 9, -1):
        per_line = max(1, math.floor(width * 72 * 0.94 / pt))
        lines = sum(
            max(1, math.ceil(ppt_layouts._em(part) / per_line))
            for part in str(text).split("\n")
        )
        if max_lines > 0 and lines > max_lines:
            continue
        if lines * pt * 1.3 <= height * 72:
            return pt
    return None


def _apply_point_size(shape, pt: int) -> None:
    for para in shape.text_frame.paragraphs:
        for run in para.runs:
            run.font.size = Pt(pt)


def _grow_shape(shape, spec: dict, text: str, page_width: float) -> None:
    """标题字比原框宽时，把框拉到 grow_w。左边不动，右边不超出页边 0.5 英寸。"""
    raw = (spec or {}).get("grow_w")
    if raw in (None, ""):
        return
    try:
        grow = float(raw)
    except (TypeError, ValueError):
        return
    current = ppt_base._inches(shape.width)
    if grow <= current + 0.05:
        return
    try:
        pt = int(spec.get("font_pt") or 18) or 18
    except (TypeError, ValueError):
        pt = 18
    if pt <= 0:
        pt = 18
    if ppt_layouts._em(text) * pt / 72 <= current * 0.94:
        return
    left = ppt_base._inches(shape.left)
    if page_width > 0:
        grow = min(grow, max(current, page_width - 0.5 - left))
    if grow <= current + 0.05:
        return
    shape.width = Inches(grow)


def _ancestor_group(shape):
    node = shape._element.getparent()
    while node is not None:
        if str(node.tag).endswith("grpSp"):
            return node
        node = node.getparent()
    return None


def _inside_element(shape, element) -> bool:
    node = shape._element
    while node is not None:
        if node is element:
            return True
        node = node.getparent()
    return False


def _index_or_empty(shape) -> bool:
    """底条没有字。序号圆里通常只有「1」「04」这种编号。"""
    try:
        if not getattr(shape, "has_text_frame", False):
            return True
        text = str(shape.text_frame.text or "").strip()
    except Exception:
        return True
    return text == "" or bool(re.fullmatch(r"\d{1,2}", text))


def _overlap_y(anchor, shape) -> float:
    top = ppt_base._inches(anchor.top)
    height = ppt_base._inches(anchor.height) or 0.01
    other_top = ppt_base._inches(shape.top)
    other_h = ppt_base._inches(shape.height)
    overlap = min(top + height, other_top + other_h) - max(top, other_top)
    return overlap / height


def _drop_unused_toc_rows(slide, page_card: dict, values: dict, written: set, keep: set) -> None:
    """目录页上没写到的条目整组删掉，连同同一行的空底条和序号。"""
    if str((page_card or {}).get("role") or "") != "toc":
        return
    by_item: dict[str, list] = {}
    for spec in page_card.get("slots") or []:
        if not isinstance(spec, dict):
            continue
        match = re.match(r"^item(\d+)\.", str(spec.get("name") or ""))
        if not match:
            continue
        by_item.setdefault(match.group(1), []).append(spec)
    anchors = []
    used_anchors = []
    for specs in by_item.values():
        filled = any(str((values or {}).get(spec.get("name")) or "").strip() for spec in specs)
        found = []
        for spec in specs:
            shape = _shape_for_slot(slide, spec)
            if shape is not None:
                found.append(shape)
        if filled:
            used_anchors.extend(found)
        else:
            anchors.extend(found)
    for shape in list(_iter_shapes(slide)):
        if not _index_or_empty(shape):
            continue
        height = ppt_base._inches(shape.height)
        sid = getattr(shape, "shape_id", None)
        for anchor in used_anchors:
            anchor_h = ppt_base._inches(anchor.height) or 0.01
            if height > anchor_h * 1.8:
                continue
            if sid is not None and _overlap_y(anchor, shape) > 0.5:
                keep.add(sid)
                # 整份清残留时，纯数字会被当成示例字清掉。用过的序号留着。
                try:
                    if re.fullmatch(r"\d{1,2}", str(shape.text_frame.text or "").strip()):
                        shape.name = "ars:toc-index"
                except Exception:
                    pass
                break
    if not anchors:
        return
    anchor_els = {shape._element for shape in anchors}
    group_els = []
    seen: set[int] = set()
    for anchor in anchors:
        group_el = _ancestor_group(anchor)
        if group_el is None or id(group_el) in seen:
            continue
        holds_written = any(
            getattr(shape, "shape_id", None) in written and _inside_element(shape, group_el)
            for shape in _iter_shapes(slide)
        )
        if holds_written:
            continue
        seen.add(id(group_el))
        group_els.append(group_el)
    loose = []
    for shape in list(_iter_shapes(slide)):
        if getattr(shape, "shape_id", None) in written:
            continue
        if any(_inside_element(shape, group_el) for group_el in group_els):
            continue
        element = shape._element
        if element in anchor_els:
            loose.append(shape)
            continue
        if not _index_or_empty(shape):
            continue
        height = ppt_base._inches(shape.height)
        for anchor in anchors:
            anchor_h = ppt_base._inches(anchor.height) or 0.01
            if height > anchor_h * 1.8:
                continue
            if _overlap_y(anchor, shape) > 0.5:
                loose.append(shape)
                break
    for group_el in group_els:
        parent = group_el.getparent()
        if parent is not None:
            parent.remove(group_el)
    for shape in loose:
        try:
            ppt_base._drop(shape)
        except Exception:
            pass


def _clear_unused(slide, written_ids: set, keep_ids: set, originals: dict) -> None:
    """没写过、也不在保留名单里、文字还是模板示例的框，把字清掉。组合里的也算。"""
    for shape in _iter_shapes(slide):
        if not ppt_template.is_text_candidate(shape):
            continue
        sid = getattr(shape, "shape_id", None)
        if sid in written_ids or sid in keep_ids:
            continue
        current = str(shape.text_frame.text or "").strip()
        if current and current == str(originals.get(sid) or "").strip():
            ppt_base._set_text(shape, "")


def _fill_by_slots(slide, page_card: dict, item: dict, base: dict,
                   page_no: int = 1, warnings: list | None = None,
                   overflow: list | None = None) -> str:
    """按格子名把字填进克隆出来的模板页。返回要写进备注的全文。"""
    from services.ppt.pptx_splice import _fill_frame

    originals = {}
    for shape in _iter_shapes(slide):
        if ppt_template.is_text_candidate(shape):
            originals[getattr(shape, "shape_id", None)] = str(shape.text_frame.text or "")
    values = item.get("slots") if isinstance(item.get("slots"), dict) else {}
    deck_title = ""
    written: set = set()
    keep: set = set()
    overflow_notes: list[str] = []
    for spec in page_card.get("slots") or []:
        if not isinstance(spec, dict):
            continue
        name = str(spec.get("name") or "")
        text = str(values.get(name) or "").strip()
        if str(page_card.get("role") or "") == "end" and name == "title":
            from services.ppt import ppt_pages

            deck_title = text or str(item.get("title") or "").strip()
            text = ppt_pages.closing_title(page_card.get("sample_heading"))
        shape = _shape_for_slot(slide, spec)
        if shape is None:
            if warnings is not None:
                warnings.append({"page": page_no, "detail": f"找不到格子 {name}"})
            if text:
                overflow_notes.append(text)
                if overflow is not None:
                    overflow.append({"page": page_no, "field": name, "text": text})
            continue
        sid = getattr(shape, "shape_id", None)
        if not text:
            if not spec.get("required"):
                ppt_base._set_text(shape, "")
                written.add(sid)
            continue
        title_like = name in {"title", "section_title"} or name.endswith(".head")
        page_w, _page_h = _slide_size(base)
        _grow_shape(shape, spec, text, page_w)
        _fill_frame(shape.text_frame, [(0, text)], keep_bold=True if title_like else False, add_prefix=False)
        pt = _slot_point_size(shape, text, spec)
        if pt is None:
            _apply_point_size(shape, 10)
            overflow_notes.append(text)
            if overflow is not None:
                overflow.append({"page": page_no, "field": name, "text": text})
        else:
            _apply_point_size(shape, pt)
        try:
            shape.name = f"ars:s{page_no}:{name}"
        except Exception:
            pass
        written.add(sid)
    _drop_unused_toc_rows(slide, page_card, values, written, keep)
    for decor in page_card.get("decor") or []:
        if not isinstance(decor, dict):
            continue
        shape = _shape_for_slot(slide, decor)
        if shape is None:
            continue
        sid = getattr(shape, "shape_id", None)
        if decor.get("action") == "clear":
            ppt_base._set_text(shape, "")
            written.add(sid)
        elif decor.get("action") == "keep":
            keep.add(sid)
    _clear_unused(slide, written, keep, originals)
    notes = str(item.get("notes") or "").strip()
    if deck_title and deck_title not in notes:
        notes = (notes + "\n" + deck_title).strip()
    if overflow_notes:
        extra = "\n".join(f"· {text}" for text in overflow_notes)
        notes = (notes + "\n\n（这一页框不够，多出来的放在备注）\n" + extra).strip()
    return notes


def _page_card(pages: dict | None, page_no) -> dict:
    try:
        number = int(page_no)
    except (TypeError, ValueError):
        return {}
    for page in (pages or {}).get("pages") or []:
        if isinstance(page, dict) and int(page.get("page") or 0) == number:
            return page
    return {}


def compose(base: dict, plan: dict, template_path, out_path,
            images: list[str] | None = None, analysis: dict | None = None,
            pages: dict | None = None, plan_mark: tuple[str, str] | None = None) -> dict:
    """按计划出片，返回 {"pages", "overflow", "warnings", "template_pages", "repaired"}。"""
    template = pathlib.Path(template_path)
    prs = Presentation(str(template))
    originals = list(prs.slides)
    total = len(originals)
    pages_map = (base or {}).get("pages") or {}

    cover_src = _source(originals, pages_map.get("cover"), 1)
    section_src = _source(originals, pages_map.get("section"), min(2, total))
    content_src = _source(originals, pages_map.get("content"), min(3, total))
    end_src = _source(originals, pages_map.get("end"), total)

    overflow: list = []
    warnings: list = []
    slides_plan = (plan or {}).get("slides") or []
    page_no = 0
    reserved = _reserved_indexes(base)
    used_specimen: set = set()
    template_pages: list[int] = []

    for index, item in enumerate(slides_plan, 1):
        item = item or {}
        layout = item.get("layout") or "bullets"
        notes = str(item.get("notes") or "")
        card = _page_card(pages, item.get("template_page")) if (
            (plan or {}).get("plan_version") == 2 and item.get("template_page")
            and isinstance(item.get("slots"), dict) and pages) else {}
        slot_source = None
        if card.get("slots"):
            try:
                source_index = int(item.get("template_page")) - 1
            except (TypeError, ValueError):
                source_index = -1
            if 0 <= source_index < len(originals):
                slot_source = originals[source_index]
            else:
                warnings.append({"page": index, "detail": "模板页超出范围"})
        if slot_source is not None:
            slide = _clone_slide(prs, slot_source)
            notes = _fill_by_slots(slide, card, item, base, index, warnings, overflow) or notes
            template_pages.append(int(item.get("template_page")))
            touched = []
        elif layout == "cover":
            slide = _clone_slide(prs, cover_src)
            touched = _fill_cover(slide, base, item)
        elif layout == "section":
            spec = None if "section" not in used_specimen else _pick_other_page(
                analysis, item, reserved, used_specimen)
            if spec is not None:
                slide = _clone_slide(prs, originals[int(spec["index"])])
                notes = _fill_specimen(slide, spec, item) or notes
                used_specimen.add(int(spec["index"]))
                template_pages.append(int(spec["index"]) + 1)
                touched = []
            else:
                slide = _clone_slide(prs, section_src)
                touched = _fill_section(slide, base, item)
                used_specimen.add("section")
        elif layout == "end":
            slide = _clone_slide(prs, end_src)
            touched = _fill_end(slide, base, item)
        else:
            spec = _pick_other_page(analysis, item, reserved, used_specimen)
            if spec is not None:
                slide = _clone_slide(prs, originals[int(spec["index"])])
                notes = _fill_specimen(slide, spec, item) or notes
                used_specimen.add(int(spec["index"]))
                template_pages.append(int(spec["index"]) + 1)
                touched = []
            else:
                if layout not in ppt_layouts.LAYOUT_SPECS:
                    warnings.append({"page": index, "detail": f"不认识的版式 {layout}，按要点画"})
                    item = dict(item, layout="bullets")
                slide = ppt_base.content_canvas(prs, content_src, base, item.get("title") or "", index)
                _, touched = _draw_content(slide, base, item, images, index, overflow, warnings)
        page_no = index
        _mark_all(touched)
        if item.get("write_fallback"):
            _clear_unmarked_text(slide)
        ppt_base._notes(slide, notes)

    _delete_slides(prs, total)
    _scrub_junk(prs, base)
    prs.save(str(out_path))
    from services.ppt import ppt_check

    repaired = ppt_check.repair_deck(out_path)
    sync_app_props(out_path)
    if plan_mark and plan_mark[0]:
        from services.ppt import ppt_edit
        ppt_edit.stamp_plan_mark(out_path, plan_mark[0], plan_mark[1] or "")
    return {
        "pages": page_no,
        "overflow": overflow,
        "warnings": warnings,
        "template_pages": template_pages,
        "repaired": repaired,
    }
