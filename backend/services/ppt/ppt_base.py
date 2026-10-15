"""模板底板档案：哪几页当底板、安全区在哪、模板原文长什么样、配色是什么。

模板只提供"底板"（页眉、红条、校徽、页码），内容由 `ppt_layouts` 画进安全区。
档案存成模板旁边的 `<id>.base.json`：自动猜的 `confirmed=False`，用户在界面上确认后 `True`。

`slots` 里记的是**模板原文**，用来在复制出来的页里找到那个文本框——
按文字匹配最稳，找不到再由调用方按字号最大、位置最上兜底。
"""

from __future__ import annotations

import json
import logging
import pathlib
import re
from typing import Optional

from pptx import Presentation
from pptx.enum.dml import MSO_FILL
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches

from services.ppt import ppt_template
from services.ppt.pptx_splice import _clone_slide

logger = logging.getLogger(__name__)

BACKEND_DIR = pathlib.Path(__file__).resolve().parents[2]
DATA_DIR = BACKEND_DIR / "data"
BUILTIN_DIR = BACKEND_DIR / "assets" / "ppt_templates" / "builtin"

BASE_VERSION = 1
EMU_PER_INCH = 914400

TOP_BAND = 1.5          # 页眉区：中心在这条线以上算"顶部元素"
BOTTOM_BAND = 1.0       # 页脚区：中心在页高减这条线以下算"页脚元素"
FOOTER_MARGIN = 0.2     # 安全区下界离页脚顶部留的空
HEADER_MARGIN = 0.3     # 安全区上界离页眉底部留的空
SIDE_MARGIN = 0.7       # 安全区左右各留的空
DEFAULT_SAFE = [0.74, 1.45, 12.6, 6.75]
DEFAULT_TITLE_RIGHT = 10.1   # 内容页标题最多拉到这个横向位置（英寸）
DEFAULT_COLORS = {"accent": "#9B1B30", "accent_soft": "#F6ECEE",
                  "dark": "#262626", "gray": "#595959", "muted": "#8C8C8C"}

# 模板自带的示例文字，命中就清空或删框
JUNK_PATTERN = re.compile(
    r"请输入|请在此|请在这里输入|关键词标题|点击(输入|添加|此处)|Add your text|Type Your|Click here|"
    r"此部分内容|您的内容打在这里|苗猫|HSVmode|XXX|Thesis defense|Thesis template|"
    r"size of the text|appropriate text|Your text|Lorem|2025/5/23",
    re.IGNORECASE,
)
_THANK_RE = re.compile(r"感谢|谢谢|thank|欢迎批评", re.IGNORECASE)
_TOC_NUM_RE = re.compile(r"-?\d{1,2}-")
DATE_PATTERN = re.compile(r"\d{4}\s*/\s*\d{1,2}\s*/\s*\d{1,2}")
INFO_PATTERN = re.compile(r"答辩学生|指导老师|指导教师|汇报人|答辩人|课程")


def is_junk_text(text: str) -> bool:
    """模板残留的示例字。

    纯数字（例如孤零零的「5」）也算残留——**页码占位符除外**，
    判断是不是页码由调用方看位置决定（页脚里的数字走页号逻辑）。
    """
    raw = (text or "").strip()
    if not raw:
        return False
    if JUNK_PATTERN.search(raw):
        return True
    return bool(re.fullmatch(r"\d{1,2}", raw))


# ───────────── 形状读取 ─────────────

def _inches(value) -> float:
    try:
        return round(int(value) / EMU_PER_INCH, 2)
    except Exception:
        return 0.0


def _is_line(shape) -> bool:
    """直线；有些模板的横线是"很扁的矩形"，一并认。"""
    try:
        if shape.shape_type == MSO_SHAPE_TYPE.LINE:
            return True
    except Exception:
        pass
    h = _inches(shape.height)
    w = _inches(shape.width)
    return h <= 0.12 and w >= 2.0


def _size_of(shape) -> Optional[int]:
    try:
        for para in shape.text_frame.paragraphs:
            for run in para.runs:
                if (run.text or "").strip() and run.font.size is not None:
                    return int(run.font.size.pt)
    except Exception:
        pass
    return None


def _frames(slide) -> list[dict]:
    """这一页的文本框（占位符 + 有字的框），带位置和字号。"""
    out: list[dict] = []
    for order, shape in enumerate(slide.shapes):
        if not ppt_template.is_text_candidate(shape):
            continue
        text = shape.text_frame.text.strip()
        out.append({
            "shape": shape, "i": order, "text": text, "size": _size_of(shape),
            "x": _inches(shape.left), "y": _inches(shape.top),
            "w": _inches(shape.width), "h": _inches(shape.height),
        })
    return out


def _texts(slide) -> list:
    return [s for s in slide.shapes if ppt_template.is_text_candidate(s)]


def _fill_hex(shape) -> Optional[str]:
    try:
        if shape.fill.type != MSO_FILL.SOLID:
            return None
        return str(shape.fill.fore_color.rgb).upper()
    except Exception:
        return None


def _is_vivid(hex_value: str) -> bool:
    """去掉近白、近黑、近灰，剩下的才算模板主色。"""
    try:
        r, g, b = (int(hex_value[i:i + 2], 16) for i in (0, 2, 4))
    except Exception:
        return False
    if max(r, g, b) > 240 or max(r, g, b) < 20:
        return False
    return max(r, g, b) - min(r, g, b) >= 24


def _mix_white(hex_value: str, white_ratio: float) -> str:
    parts = []
    for i in (0, 2, 4):
        base = int(hex_value[i:i + 2], 16)
        parts.append(round(base * (1 - white_ratio) + 255 * white_ratio))
    return "".join(f"{v:02X}" for v in parts)


# ───────────── 自动猜 ─────────────

def _pick_colors(prs) -> dict:
    """统计所有页的纯色填充，出现最多的当 accent；accent_soft 是它与白 1:9 混。"""
    counts: dict[str, int] = {}
    for slide in prs.slides:
        for shape in slide.shapes:
            hex_value = _fill_hex(shape)
            if hex_value and _is_vivid(hex_value):
                counts[hex_value] = counts.get(hex_value, 0) + 1
    colors = dict(DEFAULT_COLORS)
    if counts:
        accent = max(counts.items(), key=lambda kv: (kv[1], kv[0]))[0]
        colors["accent"] = "#" + accent
        colors["accent_soft"] = "#" + _mix_white(accent, 0.9)
    return colors


def _looks_like_toc(slide) -> bool:
    """带 3 个以上「-01-」这种编号的是目录页，不能当转场页。"""
    texts = [frame["text"].strip() for frame in _frames(slide)]
    return sum(1 for text in texts if _TOC_NUM_RE.fullmatch(text)) >= 3


def _end_page(slides: list) -> int:
    """结尾页优先认最后几页里写着「感谢 / 谢谢」的那页，找不到再用最后一页。"""
    last = len(slides)
    start = max(2, last - 7)
    for index in range(last, start - 1, -1):
        text = " ".join(frame["text"] for frame in _frames(slides[index - 1]))
        if _THANK_RE.search(text):
            return index
    return last


def _section_page(slides: list) -> int:
    """转场页：有一个 ≥40pt 的「01/02/…」数字框，文本框不超过 6 个。

    设计模板的章节页常写成大号「01」，内置模板的章节页是**一句居中大字、没有横线**，
    所以再加一条次规则兜住后者。目录页（一串 -01- -02-）跳过。
    """
    last = len(slides)
    for index, slide in enumerate(slides, 1):
        if index in (1, last) or _looks_like_toc(slide):
            continue
        frames = _frames(slide)
        if len(frames) > 6:
            continue
        for frame in frames:
            if (frame["size"] or 0) >= 40 and re.fullmatch(r"\d{1,2}", frame["text"]):
                return index
    for index, slide in enumerate(slides, 1):
        if index in (1, last) or _looks_like_toc(slide):
            continue
        frames = _frames(slide)
        if not frames or len(frames) > 2:
            continue
        if any(_is_line(shape) for shape in slide.shapes):
            continue
        if any((frame["size"] or 0) >= 26 for frame in frames):
            return index
    for index, slide in enumerate(slides, 1):
        if index in (1, last) or _looks_like_toc(slide):
            continue
        return index
    return 2 if len(slides) > 1 else 1


def _safe_area(slide, slide_w: float, slide_h: float) -> list[float]:
    """上界＝顶部元素最低底边 +0.3，下界＝页脚最高顶边 −0.2，左右各 0.7。"""
    header_bottom: list[float] = []
    footer_top: list[float] = []
    for shape in slide.shapes:
        y = _inches(shape.top)
        h = _inches(shape.height)
        center = y + h / 2
        if center < TOP_BAND:
            header_bottom.append(y + h)
        elif center > slide_h - BOTTOM_BAND:
            footer_top.append(y)
    y0 = round(max(header_bottom) + HEADER_MARGIN, 2) if header_bottom else DEFAULT_SAFE[1]
    y1 = round(min(footer_top) - FOOTER_MARGIN, 2) if footer_top else round(slide_h - 0.75, 2)
    if y0 >= y1 - 1.0:      # 猜歪了就退回默认值，宁可多留白
        y0, y1 = DEFAULT_SAFE[1], DEFAULT_SAFE[3]
    return [SIDE_MARGIN, y0, round(slide_w - SIDE_MARGIN, 2), y1]


def _content_page(slides: list, safe: list[float]) -> int:
    """内容页：顶部 1.5 英寸内有标题框和一条横线，且安全区里的形状最少。

    目录页也长这个样子，用现成的 `ppt_template.is_toc_title` 排掉。
    """
    _, y0, _, y1 = safe
    best: tuple | None = None
    for index, slide in enumerate(slides, 1):
        if index in (1, len(slides)) or _looks_like_toc(slide):
            continue
        head = [f for f in _frames(slide) if f["y"] < TOP_BAND]
        if not any(f["text"] for f in head):
            continue
        if ppt_template.is_toc_title(next(f["text"] for f in head if f["text"])):
            continue
        if not any(_is_line(shape) and _inches(shape.top) < TOP_BAND for shape in slide.shapes):
            continue
        heavy = 0
        for shape in slide.shapes:
            y = _inches(shape.top)
            h = _inches(shape.height)
            if y0 - 0.1 < y + h / 2 < y1 + 0.15 and h > 0.6:
                heavy += 1
        key = (heavy, index)
        if best is None or key < best[0]:
            best = (key, index)
    def _safe_text_count(slide) -> int:
        count = 0
        for shape in slide.shapes:
            if not ppt_template.is_text_candidate(shape):
                continue
            if not shape.text_frame.text.strip():
                continue
            center = _inches(shape.top) + _inches(shape.height) / 2
            if y0 - 0.1 < center < y1 + 0.15:
                count += 1
        return count

    # 挑中的页安全区里还堆着示例字，就改用字最少的那页
    if best and _safe_text_count(slides[best[1] - 1]) <= 2:
        return best[1]
    sparsest: tuple | None = None
    for index, slide in enumerate(slides, 1):
        if index in (1, len(slides)) or _looks_like_toc(slide):
            continue
        key = (_safe_text_count(slide), index)
        if sparsest is None or key < sparsest[0]:
            sparsest = (key, index)
    if sparsest and sparsest[0][0] <= 2:
        return sparsest[1]
    if best:
        return best[1]
    if sparsest:
        return sparsest[1]
    return min(3, max(1, len(slides) - 1))


def _slots(slides: list, pages: dict) -> dict:
    cover = _guessed_cover(slides[pages["cover"] - 1])
    section = _guessed_section(slides[pages["section"] - 1])
    content = _guessed_content(slides[pages["content"] - 1])
    end = _guessed_end(slides[pages["end"] - 1])
    return {"cover": cover, "section": section, "content": content, "end": end}


def _by_size(frames: list[dict]) -> list[dict]:
    return sorted([f for f in frames if f["text"]], key=lambda f: (-(f["size"] or 0), f["y"]))


def _guessed_cover(slide) -> dict:
    """封面：字号最大的一行当标题，第二行当副标题；「答辩人／指导教师」这类信息行单独收。"""
    frames = _by_size(_frames(slide))
    body = [f for f in frames if not INFO_PATTERN.search(f["text"])]
    title = body[0]["text"] if body else ""
    subtitle = body[1]["text"] if len(body) > 1 else ""
    info = [f["text"] for f in frames if INFO_PATTERN.search(f["text"])][:2]
    return {"title": title, "subtitle": subtitle, "info": info}


def _guessed_section(slide) -> dict:
    frames = _by_size(_frames(slide))
    index = ""
    rest: list[dict] = []
    for frame in frames:
        if not index and re.fullmatch(r"\d{1,2}", frame["text"]):
            index = frame["text"]
        else:
            rest.append(frame)
    title = rest[0]["text"] if rest else ""
    subs = [f["text"] for f in rest[1:] if len(f["text"]) <= 10][:3]
    return {"index": index, "title": title, "subs": subs}


def _guessed_content(slide) -> dict:
    """内容页的占位标题在页眉区找；页脚里的日期虽然也命中残留字样，不能当标题。"""
    frames = _by_size(_frames(slide))
    head = [f for f in frames if f["y"] < TOP_BAND]
    title = ""
    for frame in head:
        if frame["text"] and is_junk_text(frame["text"]):
            title = frame["text"]
            break
    if not title and head:
        title = head[0]["text"]
    if not title and frames:
        title = frames[0]["text"]
    return {"title": title, "title_right_limit": DEFAULT_TITLE_RIGHT}


def _guessed_end(slide) -> dict:
    frames = _by_size(_frames(slide))
    info = [f["text"] for f in frames if INFO_PATTERN.search(f["text"])][:2]
    if not info:
        info = [f["text"] for f in frames[:2]]
    return {"info": info}


def _suggested_pages(pages: dict | None, total: int) -> dict | None:
    """页说明里的底板四项都在页号范围内时采用，缺一项就不用。"""
    suggested = ((pages or {}).get("deck") or {}).get("suggested_base") or {}
    chosen = {}
    for key in ("cover", "section", "content", "end"):
        try:
            number = int(suggested.get(key))
        except (TypeError, ValueError):
            return None
        if number < 1 or number > total:
            return None
        chosen[key] = number
    return chosen


def guess_base(pptx_path, pages: dict | None = None) -> dict:
    """自动猜一份底板档案。页说明里四页都齐时用那四页，否则按几何猜。"""
    path = pathlib.Path(pptx_path)
    prs = Presentation(str(path))
    slides = list(prs.slides)
    if not slides:
        raise ValueError(f"{path.name} 里没有幻灯片")
    slide_w = _inches(prs.slide_width) or DEFAULT_SAFE[2] + SIDE_MARGIN
    slide_h = _inches(prs.slide_height) or 7.5

    chosen = _suggested_pages(pages, len(slides))
    if chosen is None:
        chosen = {"cover": 1, "section": _section_page(slides), "content": 1, "end": _end_page(slides)}
        chosen["content"] = _content_page(slides, DEFAULT_SAFE)
    rough = chosen["content"]
    return {
        "version": BASE_VERSION,
        "slide_size": [round(slide_w, 2), round(slide_h, 2)],
        "pages": chosen,
        "safe_area": _safe_area(slides[rough - 1], slide_w, slide_h),
        "colors": _pick_colors(prs),
        "slots": _slots(slides, chosen),
        "confirmed": False,
    }


# ───────────── 存档 ─────────────

def _template_pptx(tid: str) -> Optional[pathlib.Path]:
    tid = (tid or "").strip()
    if not tid:
        return None
    if tid.startswith("builtin-"):
        candidate = BUILTIN_DIR / f"{tid[len('builtin-'):]}.pptx"
        return candidate if candidate.is_file() else None
    for pptx in sorted(DATA_DIR.glob(f"workspaces/*/ppt_templates/{tid}.pptx")):
        return pptx
    return None


def _base_path(tid: str) -> Optional[pathlib.Path]:
    pptx = _template_pptx(tid)
    return pptx.with_name(pptx.stem + ".base.json") if pptx else None


def base_for(template: dict, pages: dict | None = None) -> dict:
    """先用已经存下的底板。没有就按页说明猜；没把页说明传进来就自己读。"""
    template = template or {}
    tid = str(template.get("id") or "")
    saved = load_base(tid) if tid else None
    if saved:
        return saved
    path = str(template.get("file_path") or "")
    if not path:
        raise ValueError("模板没有文件路径")
    doc = pages
    if doc is None:
        from services.ppt import ppt_pages
        doc = ppt_pages.load_pages(template)
    return guess_base(path, doc)


def load_base(tid: str) -> Optional[dict]:
    path = _base_path(tid)
    if not path or not path.is_file():
        return None
    try:
        base = json.loads(path.read_text(encoding="utf-8"))
    except Exception as exc:
        logger.warning("底板档案读不出来（%s）：%s", path.name, exc)
        return None
    return base if isinstance(base, dict) else None


def save_base(tid: str, base: dict) -> None:
    """存盘即视为用户确认过这份档案。"""
    path = _base_path(tid)
    if not path:
        raise FileNotFoundError(f"找不到模板 {tid} 的 pptx")
    payload = dict(base or {})
    payload["version"] = BASE_VERSION
    payload["confirmed"] = True
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


# ───────────── 形状改写（content_canvas 与阶段 4 共用） ─────────────

def _set_text(shape, text: str) -> None:
    """换字但保留第一段第一个 run 的样式。"""
    tf = shape.text_frame
    first = tf.paragraphs[0]
    runs = first.runs
    if runs:
        runs[0].text = text
        for extra in runs[1:]:
            extra._r.getparent().remove(extra._r)
    else:
        first.text = text
    for extra in tf.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)


def _drop(shape) -> None:
    shape._element.getparent().remove(shape._element)


def _widen(shape, right_limit: float) -> None:
    shape.width = Inches(max(_inches(shape.width), right_limit - _inches(shape.left)))


def _left(shape) -> None:
    for para in shape.text_frame.paragraphs:
        para.alignment = PP_ALIGN.LEFT


def _overlaps(a, b) -> bool:
    ax, ay, aw, ah = (_inches(v) for v in (a.left, a.top, a.width, a.height))
    bx, by, bw, bh = (_inches(v) for v in (b.left, b.top, b.width, b.height))
    return ax < bx + bw and bx < ax + aw and ay < by + bh and by < ay + ah


def _drop_lines_over(slide, boxes) -> None:
    """模板的装饰细线压到我们的文字上就删掉。"""
    for shape in list(slide.shapes):
        if _is_line(shape) and any(_overlaps(shape, box) for box in boxes):
            _drop(shape)


def _notes(slide, text: str) -> None:
    if not text:
        return
    slide.notes_slide.notes_text_frame.text = text


def _is_picture(shape) -> bool:
    try:
        return shape.shape_type == MSO_SHAPE_TYPE.PICTURE
    except Exception:
        return False


def _mark_ours(shape) -> None:
    """给自己改过的形状加名字前缀。收尾清残留字时按名字认，不用 id()。

    图文页的图片占位框名字是给换图用的，不能改。
    """
    if shape is None:
        return
    try:
        name = str(shape.name or "shape")
        if name == "ppt_pic_slot" or name.startswith("ars:"):
            return
        shape.name = "ars:" + name
    except Exception:
        pass


def content_canvas(prs, src_slide, base: dict, title: str, page_no: int):
    """复制内容页底板：清掉安全区里的示例内容，只留页眉、校徽和页脚。"""
    slide = _clone_slide(prs, src_slide)
    safe = (base or {}).get("safe_area") or DEFAULT_SAFE
    _, y0, _, y1 = (float(v) for v in list(safe)[:4])
    content_slot = ((base or {}).get("slots") or {}).get("content") or {}
    title_src = str(content_slot.get("title") or "")
    slide_w = _inches(prs.slide_width) or 13.33
    right_limit = min(
        float(content_slot.get("title_right_limit") or DEFAULT_TITLE_RIGHT),
        slide_w - 0.12,
    )
    footer_line = (_inches(prs.slide_height) or 7.5) - BOTTOM_BAND

    for shape in list(slide.shapes):
        center = _inches(shape.top) + _inches(shape.height) / 2
        if y0 - 0.1 < center < y1 + 0.15 and not _is_picture(shape):
            _drop(shape)

    for shape in _texts(slide):
        text = shape.text_frame.text.strip()
        if not text:
            continue
        center = _inches(shape.top) + _inches(shape.height) / 2
        if title_src and text == title_src:
            _widen(shape, right_limit)
            _set_text(shape, title)
            shape.text_frame.word_wrap = False
            _mark_ours(shape)
        elif DATE_PATTERN.fullmatch(text):
            _set_text(shape, "")
        elif text.isdigit() and center > footer_line:
            _set_text(shape, str(page_no))
            _mark_ours(shape)
        elif text.isdigit() or is_junk_text(text):
            _set_text(shape, "")
    return slide
