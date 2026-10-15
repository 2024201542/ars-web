"""出片后自检：几何、文字，外加可选的视觉复核。

几何和文字问题在本地判定，量字规则和版式库共用一套口径；视觉复核默认关闭
（`ARS_PPT_VLM_CHECK=1` 才跑、最多抽查 3 页），只报告，不自动改片子。
"""

from __future__ import annotations

import base64
import logging
import os
import pathlib

from pptx import Presentation
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_SHAPE_TYPE
from pptx.util import Inches

from services.ppt import ppt_base, ppt_layouts, ppt_template, pptx_render

logger = logging.getLogger(__name__)

TOP_BAND = ppt_base.TOP_BAND          # 页眉区上界（英寸）
VISION_MODEL = "qwen3-vl-plus"
VISION_MAX_PAGES = 3
VISION_PROMPT = (
    "这是一页 PPT 的渲染图。只看一眼就能看出的排版问题：文字被切掉、文字互相压住、"
    "元素跑到页面外、大片空白、颜色刺眼到看不清字。"
    "没有这类问题就只回答一个字：无。"
    "有问题就用一句话说清是什么现象。不要提改进建议，不要评价内容好不好。"
)
VISION_NONE = {"无", "无。", "没有", "没有问题", "none", "None"}


def _inch(value) -> float:
    """精确换算成英寸。

    这里不能用 `ppt_base._inches`：它会四舍五入到两位小数，
    相邻框"边界相接"会被算成重叠 0.01 英寸（时间轴那页就误报过一次）。
    """
    try:
        return int(value) / ppt_base.EMU_PER_INCH
    except Exception:
        return 0.0


def _is_placeholder(shape) -> bool:
    try:
        return bool(shape.is_placeholder)
    except Exception:
        return False


def _page_size(prs) -> tuple[float, float]:
    width = _inch(prs.slide_width) or 13.33
    height = _inch(prs.slide_height) or 7.5
    return width, height


def _frames(slide) -> list[dict]:
    """这一页有文字的框：位置、文字、字号。"""
    out: list[dict] = []
    for shape in slide.shapes:
        if not ppt_template.is_text_candidate(shape):
            continue
        text = shape.text_frame.text.strip()
        if not text:
            continue
        out.append({
            "text": text,
            "x": _inch(shape.left), "y": _inch(shape.top),
            "w": _inch(shape.width), "h": _inch(shape.height),
            "size": ppt_base._size_of(shape),
        })
    return out


def _overlaps(a: dict, b: dict) -> bool:
    gap = 0.01          # 边框相切不算重叠
    return (a["x"] < b["x"] + b["w"] - gap and b["x"] < a["x"] + a["w"] - gap
            and a["y"] < b["y"] + b["h"] - gap and b["y"] < a["y"] + a["h"] - gap)


def _fits(frame: dict) -> bool:
    """按框里实际的字号，这段字放不放得下（和版式库同一套量字规则）。

    只报"极端放不下"：PowerPoint 的文本框不裁剪溢出的文字，框比文字略小时
    渲染出来照常能看（实测差 8%、差一倍都没事，报了全是噪声）。
    所以这里要求文字所需高度超过框高的 3 倍才算问题；真要看有没有被切，
    得看渲染图，或者开 `ARS_PPT_VLM_CHECK=1` 让看图模型复核。
    """
    size = frame.get("size")
    if not size:
        return True
    return ppt_layouts.fit_pt(frame["text"], frame["w"], frame["h"] * 3,
                              size, size) is not None


def check_deck(pptx_path, base: dict | None = None) -> list[dict]:
    """检查成品，返回 [{"page", "kind", "detail"}]。

    kind：offslide（出页面）/ overflow（框里放不下）/ overlap（文字框互压）/
    junk（模板残留示例字）/ empty_title（页眉有占位符却没写标题）/ repeat_title（相邻页标题相同）。

    不检查"出安全区"：模板自带的封面大标题、页脚信息本来就跨越安全区边界，
    按安全区判会把每一页的封面都报成问题；我们画的内容框由版式库保证不越界
    （`test_ppt_layouts` 里逐个版式测过），真被裁掉的是"出页面"这一种。
    """
    prs = Presentation(str(pptx_path))
    width, height = _page_size(prs)

    issues: list[dict] = []
    headers: list[str] = []
    for index, slide in enumerate(prs.slides, 1):
        frames = _frames(slide)
        header = ""
        for frame in frames:
            head = frame["text"][:20]
            if (frame["x"] < -0.02 or frame["y"] < -0.02
                    or frame["x"] + frame["w"] > width + 0.02
                    or frame["y"] + frame["h"] > height + 0.02):
                issues.append({"page": index, "kind": "offslide",
                               "detail": f"「{head}」超出页面（左 {frame['x']:.2f}，上 {frame['y']:.2f}）"})
            if not _fits(frame):
                issues.append({"page": index, "kind": "overflow",
                               "detail": f"「{head}」在 {frame['w']:.1f}×{frame['h']:.1f} 英寸的框里放不下"})
            if ppt_base.JUNK_PATTERN.search(frame["text"]):
                issues.append({"page": index, "kind": "junk",
                               "detail": f"残留示例字：{frame['text'][:24]}"})
            if frame["y"] < TOP_BAND and not header:
                header = frame["text"]

        for i in range(len(frames)):
            for j in range(i + 1, len(frames)):
                if _overlaps(frames[i], frames[j]):
                    issues.append({"page": index, "kind": "overlap",
                                   "detail": f"「{frames[i]['text'][:12]}」与「{frames[j]['text'][:12]}」重叠"})

        graphics = [_box(shape) for shape in slide.shapes if is_blocking_graphic(shape, width)]
        for frame in frames:
            ink = _ink_box(frame)
            # 卡片底图和小装饰上的字不算。只记大标题从左边伸进页面右半的图或弧线。
            try:
                pt = int(frame.get("size") or 0)
            except (TypeError, ValueError):
                pt = 0
            if not ink or pt < 32:
                continue
            hit = False
            for graphic in graphics:
                if graphic["x"] < width * 0.55 or graphic["x"] <= ink["x"] + 1.0:
                    continue
                if ink["x"] + ink["w"] > graphic["x"] + 0.05:
                    hit = True
                    break
            if hit:
                issues.append({"page": index, "kind": "overlap",
                               "detail": f"「{frame['text'][:12]}」压到图片上"})

        headers.append(header)
        # "漏写标题"：页眉区有**占位符**框却一个字都没有。
        # 只认占位符：页顶的装饰空框（很多模板都有）会被误判；标题用普通文本框的模板
        # 这一项抓不到，那种情况由 junk 检查兜底（模板原文"请输入你的标题"会被抓出来）。
        header_shapes = [s for s in slide.shapes
                         if s.has_text_frame and _inch(s.top) < TOP_BAND
                         and (s.text_frame.text.strip() or _is_placeholder(s))]
        if header_shapes and not header:
            issues.append({"page": index, "kind": "empty_title", "detail": "页眉区有框但没有标题"})

    for i in range(1, len(headers)):
        if headers[i] and headers[i] == headers[i - 1]:
            issues.append({"page": i + 1, "kind": "repeat_title",
                           "detail": f"标题与上一页相同：{headers[i][:20]}"})
    hints = {
        "junk": "模板的示例字没清干净。去「设置底板」，换一张更干净的内容页。",
        "overflow": "这一页字太多。在编辑页把这一段改短。",
        "offslide": "有字跑出页面。点「一键修掉」把框拉回页面里。",
        "overlap": "有文字叠在一起。点「一键修掉」会删掉压在正文上的模板字。",
        "empty_title": "这一页没有标题。在编辑页补上。",
        "repeat_title": "和上一页标题相同。在编辑页改掉其中一个。",
    }
    for issue in issues:
        issue["hint"] = hints.get(issue["kind"], "打开计划编辑页看这一页。")
    return issues


def _box(shape) -> dict:
    return {
        "x": _inch(shape.left), "y": _inch(shape.top),
        "w": _inch(shape.width), "h": _inch(shape.height),
    }


def is_blocking_graphic(shape, page_w: float) -> bool:
    """图片、椭圆和自由曲线会挡住标题。满版底图和色块不算。"""
    try:
        kind = shape.shape_type
        width = _inch(shape.width)
    except Exception:
        return False
    if width <= 0 or width >= float(page_w or 0) * 0.85:
        return False
    if kind in (MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.FREEFORM):
        return True
    if kind == MSO_SHAPE_TYPE.AUTO_SHAPE:
        try:
            return shape.auto_shape_type == MSO_AUTO_SHAPE_TYPE.OVAL
        except Exception:
            return False
    return False


def _ink_box(frame: dict) -> dict | None:
    """按字号估算文字实际占的宽度。框比字宽时，右边空出来的部分不算。"""
    try:
        pt = int(frame.get("size") or 0)
    except (TypeError, ValueError):
        pt = 0
    if pt <= 0:
        return None
    width = ppt_layouts._em(frame.get("text") or "") * pt / 72.0
    if width <= 0:
        return None
    return {
        "x": frame["x"], "y": frame["y"],
        "w": min(float(frame["w"]), width), "h": frame["h"],
    }


def _is_picture(shape) -> bool:
    try:
        return shape.shape_type == MSO_SHAPE_TYPE.PICTURE
    except Exception:
        return False


def _clamp_shape(shape, width: float, height: float, margin: float = 0.06) -> bool:
    """把探出页面的框挪回来。挪不开就缩小。返回有没有改。"""
    x, y, w, h = _inch(shape.left), _inch(shape.top), _inch(shape.width), _inch(shape.height)
    if w <= 0 or h <= 0:
        return False
    limit_w = max(margin * 2, width - margin * 2)
    limit_h = max(margin * 2, height - margin * 2)
    nw, nh = min(w, limit_w), min(h, limit_h)
    nx, ny = x, y
    if nx < margin:
        nx = margin
    if ny < margin:
        ny = margin
    if nx + nw > width - margin:
        nx = width - margin - nw
    if ny + nh > height - margin:
        ny = height - margin - nh
    nx = max(margin, nx)
    ny = max(margin, ny)
    if abs(nx - x) < 0.01 and abs(ny - y) < 0.01 and abs(nw - w) < 0.01 and abs(nh - h) < 0.01:
        return False
    shape.left = Inches(nx)
    shape.top = Inches(ny)
    shape.width = Inches(nw)
    shape.height = Inches(nh)
    return True


def repair_deck(pptx_path) -> dict:
    """跑出页面的框拉回来；叠在一起的模板残留字删掉；我们自己的框错开后仍留在页内。

    返回 {"pulled", "removed"}。
    """
    prs = Presentation(str(pptx_path))
    width, height = _page_size(prs)
    pulled = 0
    removed = 0
    for slide in prs.slides:
        removed += _drop_leftover_overlaps(slide)
        for shape in list(slide.shapes):
            if _is_picture(shape):
                continue
            area = _inch(shape.width) * _inch(shape.height)
            if area > width * height * 0.7:
                continue
            name = str(getattr(shape, "name", "") or "")
            has_text = False
            try:
                has_text = bool(shape.has_text_frame and (shape.text_frame.text or "").strip())
            except Exception:
                has_text = False
            if not has_text and not name.startswith("ars:"):
                continue
            if _clamp_shape(shape, width, height):
                pulled += 1
        pulled += _unstick_ours(slide, height)
        for shape in list(slide.shapes):
            if _is_picture(shape):
                continue
            area = _inch(shape.width) * _inch(shape.height)
            if area > width * height * 0.7:
                continue
            name = str(getattr(shape, "name", "") or "")
            has_text = False
            try:
                has_text = bool(shape.has_text_frame and (shape.text_frame.text or "").strip())
            except Exception:
                has_text = False
            if not has_text and not name.startswith("ars:"):
                continue
            if _clamp_shape(shape, width, height):
                pulled += 1
    prs.save(str(pptx_path))
    return {"pulled": pulled, "removed": removed}


def _ours_text(slide) -> list:
    found = []
    for shape in list(slide.shapes):
        name = str(getattr(shape, "name", "") or "")
        if not name.startswith("ars:"):
            continue
        try:
            if not (shape.text_frame.text or "").strip():
                continue
        except Exception:
            continue
        found.append(shape)
    return found


def _drop_leftover_overlaps(slide) -> int:
    """模板残留字叠在别的字上就删掉，成对的水印两行一起删。标了 ars: 的是我们写的，留着。"""
    shapes = []
    for shape in list(slide.shapes):
        if not ppt_template.is_text_candidate(shape):
            continue
        try:
            if not (shape.text_frame.text or "").strip():
                continue
        except Exception:
            continue
        shapes.append(shape)
    drop = []
    for shape in shapes:
        name = str(getattr(shape, "name", "") or "")
        if name.startswith("ars:"):
            continue
        box = _box(shape)
        if any(other is not shape and _overlaps(box, _box(other)) for other in shapes):
            drop.append(shape)
    removed = 0
    for shape in drop:
        parent = shape._element.getparent()
        if parent is None:
            continue
        parent.remove(shape._element)
        removed += 1
    return removed


def _unstick_ours(slide, height: float, margin: float = 0.06) -> int:
    """两个我们自己的框叠在一起时错开。错开之后仍必须留在页面里。"""
    changed = 0
    floor = height - margin
    for _ in range(8):
        shapes = _ours_text(slide)
        shapes.sort(key=lambda shape: (_inch(shape.top), _inch(shape.left)))
        moved = False
        for index, cur in enumerate(shapes):
            current = _box(cur)
            for prev in shapes[:index]:
                upper = _box(prev)
                if not _overlaps(upper, current):
                    continue
                if current["y"] < upper["y"]:
                    continue
                avail = floor - upper["y"]
                upper_h = min(upper["h"], max(0.22, avail * 0.5))
                if upper["y"] + upper_h + 0.04 + 0.18 > floor:
                    upper_h = max(0.18, avail - 0.04 - 0.18)
                prev.height = Inches(upper_h)
                new_y = _inch(prev.top) + _inch(prev.height) + 0.04
                lower_h = min(current["h"], max(0.18, floor - new_y))
                if new_y + lower_h > floor:
                    new_y = max(margin, floor - lower_h)
                cur.top = Inches(new_y)
                cur.height = Inches(lower_h)
                changed += 1
                moved = True
                break
            if moved:
                break
        if not moved:
            break
    return changed


def render_thumbs(pptx_path, dpi: int = 72) -> dict[int, pathlib.Path]:
    """整套缩略图（dpi 72 够界面看），交给阶段 0 的渲染路。"""
    prs = Presentation(str(pptx_path))
    pages = list(range(1, len(prs.slides) + 1))
    return pptx_render.render_pages(pptx_path, pages, dpi=dpi)


# ───────────── 可选的视觉复核 ─────────────

def vision_enabled() -> bool:
    return (os.environ.get("ARS_PPT_VLM_CHECK") or "").strip().lower() in ("1", "true", "yes", "on")


async def _vision_opinion(user_id: str, image_path: pathlib.Path) -> tuple[str, str]:
    """让看图模型说这一页有没有明显排版问题。返回 (意见, 错误)。"""
    import httpx

    from providers import PROVIDERS
    from services.state_tracker import tracker

    provider = PROVIDERS["qwen"]
    key = (await tracker.get_setting(provider["key_setting"], user_id) or "").strip()
    if not key:
        return "", f"{provider['display_name']} 还没配 API Key"
    base = (await tracker.get_setting(provider["base_url_setting"], user_id) or "").strip() \
        or (provider.get("default_base_url") or "")

    data_url = "data:image/png;base64," + base64.b64encode(image_path.read_bytes()).decode()
    body = {
        "model": VISION_MODEL,
        "messages": [{"role": "user", "content": [
            {"type": "text", "text": VISION_PROMPT},
            {"type": "image_url", "image_url": {"url": data_url}},
        ]}],
        "max_tokens": 200,
        "temperature": 0.1,
    }
    try:
        async with httpx.AsyncClient(timeout=120) as client:
            resp = await client.post(f"{base.rstrip('/')}/chat/completions", json=body,
                                     headers={"Authorization": f"Bearer {key}"})
        if resp.status_code != 200:
            return "", f"HTTP {resp.status_code}: {resp.text[:120]}"
        content = resp.json()["choices"][0]["message"].get("content") or ""
        return str(content).strip(), ""
    except Exception as exc:
        return "", f"{type(exc).__name__}: {exc}"


async def check_visual(pptx_path, base: dict | None = None, user_id: str = "",
                       pages: list[int] | None = None, dpi: int = 90) -> list[dict]:
    """抽查最多 3 页让看图模型复核。没开开关直接返回空，不花调用。"""
    if not vision_enabled():
        return []
    prs = Presentation(str(pptx_path))
    total = len(prs.slides)
    if not total:
        return []
    if pages:
        picks = [p for p in pages if 1 <= p <= total][:VISION_MAX_PAGES]
    else:
        body = list(range(2, total)) or list(range(1, total + 1))
        step = max(1, len(body) // VISION_MAX_PAGES)
        picks = body[::step][:VISION_MAX_PAGES]

    thumbs = pptx_render.render_pages(pptx_path, picks, dpi=dpi)
    issues: list[dict] = []
    for page in picks:
        image = thumbs.get(page)
        if not image:
            issues.append({"page": page, "kind": "visual",
                           "detail": f"缩略图渲染失败：{pptx_render.last_error()}"})
            continue
        opinion, err = await _vision_opinion(user_id, image)
        if err:
            issues.append({"page": page, "kind": "visual",
                           "detail": f"视觉复核没跑成：{err}"})
            continue
        text = opinion.strip()
        if text and text not in VISION_NONE:
            issues.append({"page": page, "kind": "visual", "detail": text[:160]})
    return issues
